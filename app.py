import os
import re
import json
import random
import time
import threading
import requests
from flask import Flask, render_template, request, redirect, url_for, jsonify, session, send_from_directory, send_file
from datetime import datetime, timedelta
import gspread
from oauth2client.service_account import ServiceAccountCredentials
import urllib.parse
import pytz
import html
import logging
import secrets
from functools import wraps

app = Flask(__name__)
app.secret_key = secrets.token_hex(16)

# ============================================================
# НАСТРОЙКА
# ============================================================
CREDENTIALS_FILE = "/data/credentials.json"
SHEET_NAME = "Система ремонта ВВ"
START_COORDS = "55.775267, 37.745690"
MASTERS = ['Алексей', 'Сергей', 'Транзит', 'Иван']
PARTS = [
    ('Покрышка', 1200),
    ('Камера', 300),
    ('Обод', 1200),
    ('Ось', 180),
    ('Крыло', 400),
    ('Ручка тормоза с бачком', 550),
    ('Ручка тормоза', 140),
    ('Колодки', 200),
    ('Концевик', 220),
    ('Суппорт', 550),
    ('Подшипник', 50),
    ('Вилка', 2500),
    ('Зеркала для электровелосипеда (2 шт.)', 500),
    ('Хомут для зеркал комплект', 200),
    ('Подножка', 850),
    ('Пружина', 50),
    ('Багажник задний для Monster', 1000),
    ('Багажник передний 39х32см для электровелосипедов', 800),
    ('Сиденье', 700),
    ('Подседельный штырь', 450),
    ('Сигналка', 350),
    ('Ручка газа', 1000),
    ('Блок переключения скоростей /фар/сигнал/поворотники', 450),
    ('Треугольник', 300),
    ('Стойка под руль', 550),
    ('Руль', 550),
    ('Фара', 750),
]
PARTS_PRICE = {name: price for name, price in PARTS}
PARTS_SET = set(PARTS_PRICE)
CACHE_TTL = 300
CACHE_DIR = "/data/cache"
BOT_API_URL = "https://route-bot-dzufear.waw0.amvera.tech"
BOT_HOOK_KEY = os.environ.get("BOT_HOOK_KEY", "velobot-route")
XDEVICE_URL = "https://app.xdevice.ru/api/v1/objects"
XDEVICE_TTL = 3600

# IOT настройки
IOT_SOURCE_SHEET_ID = "1BoZ7GFmL2q56bjqt1CFVV_PeqnKVaeb3anTJkf0s6xA"
IOT_VEHICLES_SHEET_ID = "1s_hXPSWueMAo3W1pgLCG0eHhYKoW0uoIZ6CGWFf55VU"
IOT_REPORT_SHEET_ID = "1s_hXPSWueMAo3W1pgLCG0eHhYKoW0uoIZ6CGWFf55VU"

IOT_BAG_FILE = "iot_bag.json"
IOT_SOURCE_FILE = "iot_source.json"
IOT_VEHICLES_FILE = "iot_vehicles.json"
IOT_HISTORY_FILE = "iot_history.json"
IOT_CURATORS_SHEET_ID = "17lzP6e2-8a9AfxuBEA98h-FcsOljeDW_gjOw5JK7Drg"
IOT_CURATORS_FILE = "iot_curators.json"
# Эти версии в маршрут не попадают. 156, 158 и 159 показываем.
IOT_DONE_VERSIONS = {'3.0.5', '3.0.6', '160'}
IOT_FLASHED_VERSION = '3.0.5'

os.makedirs(CACHE_DIR, exist_ok=True)

MSK = pytz.timezone('Europe/Moscow')
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)
import db
db.configure(PARTS)
db.start()

MASTER_CREDENTIALS = {
    'admin': {'password': 'admin123', 'name': 'Админ', 'role': 'admin'},
    'alexey': {'password': 'alexey0304', 'name': 'Алексей', 'role': 'master'},
    'sergey': {'password': 'sergey1992', 'name': 'Сергей', 'role': 'master'},
    'transit': {'password': 'transit2024', 'name': 'Транзит', 'role': 'master'},
    'ivan': {'password': 'ivan1994', 'name': 'Иван', 'role': 'master'},
    'iot': {'password': 'iot2026', 'name': 'IOT-Мастер', 'role': 'iot'}
}

# ============================================================
# ГЛОБАЛЬНЫЕ ПЕРЕМЕННЫЕ
# ============================================================
uid_index = {}
darks_ref = {}
_darks_loaded_at = 0
queue_lock = threading.Lock()
flush_lock = threading.Lock()

# Кэш в памяти для IOT (мгновенный доступ)
_iot_source_memory = None
_iot_vehicles_memory = None
_iot_curators_memory = None

# Глобальный gspread-клиент (переиспользование)
_gspread_client = None
_gspread_client_lock = threading.Lock()
_workbook = None
_tickets_mem = []
_tickets_mem_at = 0
_sheet_backoff_until = 0
_sheet_read_lock = threading.Lock()
_app_started_at = time.time()

# ============================================================
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ============================================================
def get_msk_now():
    return datetime.now(MSK)

def is_active_status(status):
    return status in ['pending', 'todo', 'processed', 'fail', 'evacuation']

def is_evacuation_ticket(ticket):
    return (ticket or {}).get('status') == 'evacuation'

def is_supply(ticket):
    text = ' '.join([
        str((ticket or {}).get('type') or ''),
        str((ticket or {}).get('bike_type') or ''),
    ]).lower()
    return 'аккумулятор' in text or 'зарядн' in text

def ticket_hours(ticket):
    if not ticket:
        return 0
    return get_hours_since(ticket.get('timer_from') or ticket.get('created') or '')

def extract_numbers(s):
    if not s:
        return ''
    return ''.join(re.findall(r'\d', str(s)))

# ============================================================
# РАБОТА С КЭШЕМ
# ============================================================
def get_cache_path(filename):
    return os.path.join(CACHE_DIR, filename)

def read_cache(filename):
    try:
        path = get_cache_path(filename)
        if os.path.exists(path):
            with open(path, 'r', encoding='utf-8') as f:
                return json.load(f)
    except:
        pass
    return None

def write_cache(filename, data):
    try:
        path = get_cache_path(filename)
        data['updated_at'] = get_msk_now().strftime('%Y-%m-%d %H:%M:%S')
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        return True
    except:
        return False

def get_master_cache(name):
    return read_cache(f"master_{name}.json")

def save_master_cache(name, tickets, sent_date=''):
    active_tickets = [t for t in tickets if is_active_status(t.get('status'))]
    return write_cache(f"master_{name}.json", {
        'master': name,
        'sent_date': sent_date,
        'tickets': active_tickets
    })

def get_admin_cache():
    return read_cache("admin_cache.json")

def save_admin_cache(tickets):
    directions_data = {}
    for t in tickets:
        if not is_active_status(t.get('status')):
            continue
        dir_name = t.get('direction') or 'Без направления'
        if dir_name not in directions_data:
            directions_data[dir_name] = {'tickets': [], 'active_count': 0}
        directions_data[dir_name]['tickets'].append(t)
        directions_data[dir_name]['active_count'] += 1
    
    order = ['Напр 1', 'Напр 2', 'Напр 3', 'Напр 4', 'Без направления']
    sorted_directions = {}
    for key in order:
        if key in directions_data:
            sorted_directions[key] = directions_data[key]
    for key, value in directions_data.items():
        if key not in sorted_directions:
            sorted_directions[key] = value
    
    return write_cache("admin_cache.json", {'directions': sorted_directions})

def build_uid_index(tickets):
    index = {}
    for t in tickets:
        uid = t.get('uid')
        if uid:
            index[uid] = {
                'source': t.get('source'),
                'row_index': t.get('row_index'),
                'ticket': t
            }
    return index

def get_ticket_row_by_uid(uid):
    if uid in uid_index:
        return uid_index[uid]['row_index']
    return None

def update_ticket_in_admin_cache(uid, new_status, note='', display_desc='', extra=None):
    try:
        admin_cache = get_admin_cache() or {}
        updated = False
        mirrored = False
        stamp = time.time()
        for dir_name, dir_data in admin_cache.get('directions', {}).items():
            for t in dir_data.get('tickets', []):
                if t.get('uid') == uid:
                    t['status'] = new_status
                    t['is_done'] = new_status == 'done'
                    t['is_active'] = is_active_status(new_status)
                    t['_local_ts'] = stamp
                    if note:
                        t['note'] = note
                    if display_desc:
                        t['display_desc'] = display_desc
                    if extra:
                        t.update(extra)
                    t['hours_since'] = ticket_hours(t)
                    updated = True
                    break
            if updated:
                break
        for ticket in _tickets_mem:
            if ticket.get('uid') != uid:
                continue
            ticket['status'] = new_status
            ticket['is_done'] = new_status == 'done'
            ticket['is_active'] = is_active_status(new_status)
            ticket['_local_ts'] = stamp
            if note:
                ticket['note'] = note
            if display_desc:
                ticket['display_desc'] = display_desc
            if extra:
                ticket.update(extra)
            ticket['hours_since'] = ticket_hours(ticket)
            db.mirror_ticket(ticket)
            mirrored = True
        if updated and not mirrored:
            for dir_data in admin_cache.get('directions', {}).values():
                for ticket in dir_data.get('tickets', []):
                    if ticket.get('uid') == uid:
                        db.mirror_ticket(ticket)
                        mirrored = True
                        break
                if mirrored:
                    break
        if _tickets_mem:
            store_ticket_memory(_tickets_mem, fresh=False)
        if updated:
            all_tickets = []
            for dir_name, dir_data in admin_cache.get('directions', {}).items():
                all_tickets.extend(dir_data.get('tickets', []))
            save_admin_cache(all_tickets)
            return True
        return False
    except Exception as e:
        logger.error(f"Ошибка обновления admin_cache для {uid}: {e}")
        return False

LOCAL_FIELDS = ('status', 'master', 'note', 'is_done', 'is_active', 'evac_reason', 'timer_from', 'hours_since', 'display_desc')

def apply_pending_queue(tickets):
    """Пока отметка мастера не дошла до таблицы, не возвращаем заявку ему на экран."""
    try:
        tasks = read_queue().get('tasks') or []
    except Exception:
        return tickets
    by_uid = {}
    for task in tasks:
        uid = task.get('uid')
        if uid:
            by_uid[uid] = task
    for ticket in tickets or []:
        task = by_uid.get(ticket.get('uid'))
        if not task:
            continue
        kind = task.get('type')
        data = task.get('data') or {}
        if kind == 'fail':
            ticket['status'] = 'processed'
            ticket['master'] = ''
            ticket['is_done'] = False
            ticket['is_active'] = True
            if data.get('note'):
                ticket['note'] = data.get('note')
        elif kind == 'status_update':
            if data.get('note_only'):
                if data.get('note'):
                    ticket['note'] = data.get('note')
                continue
            new_status = data.get('new_status') or 'pending'
            ticket['status'] = new_status
            ticket['is_done'] = new_status == 'done'
            ticket['is_active'] = is_active_status(new_status)
            if new_status != 'evacuation':
                ticket['evac_reason'] = ''
            if new_status == 'processed':
                ticket['master'] = ''
            if data.get('write_master'):
                ticket['master'] = data.get('master') or ''
            if data.get('note'):
                ticket['note'] = data.get('note')
        elif kind == 'evacuation':
            ticket['status'] = 'evacuation'
            ticket['master'] = ''
            ticket['is_done'] = False
            ticket['is_active'] = True
            if data.get('reason'):
                ticket['evac_reason'] = data.get('reason')
        elif kind in ('done', 'replace_yes', 'evacuation_and_replace', 'transit_replace', 'transit_bulk_close'):
            ticket['status'] = 'done'
            ticket['is_done'] = True
            ticket['is_active'] = False
    return tickets

def merge_local_tickets(previous, fresh):
    """Свежая таблица не должна возвращать заявку, которую мастер только что закрыл."""
    prev = {}
    now = time.time()
    for ticket in previous or []:
        uid = ticket.get('uid')
        stamp = ticket.get('_local_ts') or 0
        if uid and stamp and now - stamp < 900:
            prev[uid] = ticket
    for ticket in fresh or []:
        old = prev.get(ticket.get('uid'))
        if not old:
            continue
        same = all(ticket.get(key) == old.get(key) for key in ('status', 'master'))
        if same:
            continue
        for key in LOCAL_FIELDS:
            if key in old:
                ticket[key] = old[key]
        ticket['_local_ts'] = old.get('_local_ts')
    return fresh

# ============================================================
# GOOGLE SHEETS
# ============================================================
def get_gspread_client():
    """Возвращает переиспользуемый gspread-клиент"""
    global _gspread_client
    with _gspread_client_lock:
        if _gspread_client is None:
            scope = ['https://spreadsheets.google.com/feeds', 'https://www.googleapis.com/auth/drive']
            creds = ServiceAccountCredentials.from_json_keyfile_name(CREDENTIALS_FILE, scope)
            _gspread_client = gspread.authorize(creds)
        return _gspread_client

def get_sheet_client():
    global _workbook
    if _workbook is None:
        _workbook = get_gspread_client().open(SHEET_NAME)
    return _workbook

def get_iot_sheet_by_id(sheet_id, sheet_name=None):
    client = get_gspread_client()
    spreadsheet = client.open_by_key(sheet_id)
    if sheet_name:
        return spreadsheet.worksheet(sheet_name)
    return spreadsheet

def clean_name(value):
    return ' '.join(str(value or '').replace('\xa0', ' ').split()).lower()

def is_transit_master(name):
    return clean_name(name) in ('транзит', 'сергей транзит')

def same_master(left, right):
    if clean_name(left) == clean_name(right):
        return True
    return is_transit_master(left) and is_transit_master(right)

RETIRED_MASTERS = {'антон', 'руслан'}

def queue_clears_master(uid):
    try:
        tasks = read_queue().get('tasks') or []
    except Exception:
        return False
    for task in tasks:
        if task.get('uid') != uid:
            continue
        kind = task.get('type')
        data = task.get('data') or {}
        if kind in ('fail', 'evacuation'):
            return True
        if kind == 'assign' and not (data.get('master') or '').strip():
            return True
        if kind == 'status_update' and data.get('write_master') and not (data.get('master') or '').strip():
            return True
    return False

def forget_retired_caches():
    for name in ('Антон', 'Руслан'):
        path = os.path.join(CACHE_DIR, f'master_{name}.json')
        if os.path.exists(path):
            try:
                os.remove(path)
            except Exception as exc:
                logger.warning(f'Не удалился кэш {name}: {exc}')

def cleared_master_uids():
    data = read_cache('masters_cleared.json') or {}
    return {key for key in data if key and key != 'updated_at'}

def remember_master_cleared(uid):
    uid = str(uid or '').strip()
    if not uid:
        return
    data = read_cache('masters_cleared.json') or {}
    data.pop('updated_at', None)
    data[uid] = int(time.time())
    write_cache('masters_cleared.json', data)

def forget_master_cleared(uid):
    uid = str(uid or '').strip()
    if not uid:
        return
    data = read_cache('masters_cleared.json') or {}
    if uid not in data:
        return
    data.pop(uid, None)
    data.pop('updated_at', None)
    write_cache('masters_cleared.json', data)

def strip_cleared_masters(tickets):
    blocked = cleared_master_uids()
    if not blocked:
        return tickets
    for ticket in tickets or []:
        if str(ticket.get('uid') or '') in blocked:
            ticket['master'] = ''
    return tickets

def drop_retired_masters():
    forget_retired_caches()
    flag = os.path.join(CACHE_DIR, 'retired_masters_done.json')
    if os.path.exists(flag):
        return
    tickets = list(_tickets_mem)
    if not tickets:
        return
    changed = 0
    for ticket in tickets:
        if clean_name(ticket.get('master')) not in RETIRED_MASTERS:
            continue
        ticket['master'] = ''
        ticket['_local_ts'] = time.time()
        remember_master_cleared(ticket.get('uid'))
        try:
            db.mirror_ticket(ticket)
        except Exception as exc:
            logger.warning(f'База не сняла старого мастера: {exc}')
        add_to_queue({
            'uid': ticket.get('uid'),
            'source': ticket.get('source') or 'Заявки',
            'type': 'assign',
            'data': {'master': ''},
        })
        changed += 1
    if changed:
        publish_caches(tickets)
        logger.info(f'Сняты старые мастера Антон и Руслан: {changed}')
    try:
        with open(flag, 'w', encoding='utf-8') as handle:
            handle.write(str(changed))
    except Exception as exc:
        logger.warning(f'Отметка о старых мастерах не сохранилась: {exc}')

def column_index(header, needles, default):
    for index, title in enumerate(header or []):
        text = str(title or '').strip().lower().replace('\xa0', ' ')
        if all(needle in text for needle in needles):
            return index
    return default

def is_quota_error(exc):
    text = str(exc)
    return '429' in text or 'Quota exceeded' in text or 'RATE_LIMIT' in text

def load_darks_reference(force=False):
    global darks_ref, _darks_loaded_at
    if darks_ref and not force and time.time() - _darks_loaded_at < 90:
        return darks_ref
    try:
        sheet_client = get_sheet_client()
        worksheet = sheet_client.worksheet("Дарксторы")
        rows = worksheet.get_all_values()
        fresh = {}
        if len(rows) > 1:
            header = rows[0]
            i_address = column_index(header, ['адрес'], 1)
            i_direction = column_index(header, ['направ'], 2)
            i_coords = column_index(header, ['коорд'], 3)
            i_master = column_index(header, ['мастер'], 7)
            i_tg = column_index(header, ['тг', 'id'], 8)
            for row in rows[1:]:
                if not row or not str(row[0]).strip():
                    continue
                darks_num = str(row[0]).strip()
                def cell(index):
                    return row[index].strip() if len(row) > index and row[index] else ''
                fresh[darks_num] = {
                    'address': cell(i_address),
                    'direction': cell(i_direction),
                    'coords': cell(i_coords),
                    'curators': cell(4),
                    'curator_ids': cell(6),
                    'master': cell(i_master),
                    'master_tg': cell(i_tg),
                }
                ids_text = fresh[darks_num]['curator_ids']
                for token in re.findall(r'\d{6,}', cell(4) + ' ' + cell(6)):
                    if token not in ids_text:
                        ids_text = (ids_text + ', ' + token).strip(', ')
                fresh[darks_num]['curator_ids'] = ids_text
        if fresh:
            darks_ref = fresh
            _darks_loaded_at = time.time()
            db.mirror_darks(darks_ref)
    except Exception as e:
        logger.error(f"Ошибка загрузки Дарксторов: {e}")
        if is_quota_error(e):
            raise
    return darks_ref

def enrich_tickets(tickets):
    if not darks_ref:
        load_darks_reference()
    for ticket in tickets or []:
        info = darks_ref.get(str(ticket.get('darks') or '').strip()) or {}
        if info.get('address'):
            ticket['address'] = info['address']
        if info.get('coords'):
            ticket['coords'] = info['coords']
        if info.get('direction') and not ticket.get('direction'):
            ticket['direction'] = info['direction']
    try:
        attach_iot(tickets)
    except Exception as exc:
        logger.warning(f'IoT к заявкам не подставился: {exc}')
    return tickets

def parse_created_date(date_str):
    if not date_str:
        return None
    date_str = date_str.strip()
    match = re.match(r'(\d{4})-(\d{2})-(\d{2})\s+(\d{2}):(\d{2}):(\d{2})', date_str)
    if match:
        return datetime(int(match.group(1)), int(match.group(2)), int(match.group(3)),
                       int(match.group(4)), int(match.group(5)), int(match.group(6)), tzinfo=MSK)
    match = re.match(r'(\d{2})\.(\d{2})\.(\d{4})\s+(\d{2}):(\d{2})', date_str)
    if match:
        return datetime(int(match.group(3)), int(match.group(2)), int(match.group(1)),
                       int(match.group(4)), int(match.group(5)), 0, tzinfo=MSK)
    return None

def get_hours_since(created_str):
    if not created_str:
        return 0
    created = parse_created_date(created_str)
    if not created:
        return 0
    delta = get_msk_now() - created
    return round(delta.total_seconds() / 3600, 1)

def generate_ticket_id(date_str):
    if not date_str:
        return None
    try:
        date_obj = parse_created_date(date_str)
        if not date_obj:
            return None
        date_part = date_obj.strftime('%d%m%y')
        random_part = str(random.randint(1000, 9999))
        return f"{date_part}-{random_part}"
    except:
        return None

_xdevice_lock = threading.Lock()
_xdevice_points = {}
_xdevice_loaded_at = 0
_xdevice_error = ""
_xdevice_started = False

def xdevice_token():
    token = (os.environ.get("XDEVICE_TOKEN") or "").strip()
    if token:
        return token
    path = "/data/xdevice_token.txt"
    if os.path.exists(path):
        with open(path, encoding="utf-8") as handle:
            return handle.read().strip()
    return ""

def normalize_iot(value):
    text = str(value or '').strip().lower().replace('ё', 'е')
    text = text.replace('–', '-').replace('—', '-').replace('\xa0', ' ')
    return re.sub(r'^iot[\s\-_]*', '', text).strip()

def iot_tokens(value):
    raw = str(value or '').strip().lower().replace('ё', 'е')
    raw = raw.replace('–', '-').replace('—', '-').replace('\xa0', ' ')
    tokens = []
    if raw:
        tokens.append(raw)
    stripped = re.sub(r'^iot[\s\-_]*', '', raw).strip()
    if stripped and stripped not in tokens:
        tokens.append(stripped)
    for match in re.findall(r'\d{2}\s*-\s*\d{3,6}', stripped or raw):
        token = re.sub(r'\s+', '', match)
        if token not in tokens:
            tokens.append(token)
    compact = re.sub(r'[^0-9a-zа-я]', '', stripped or raw)
    if compact and compact not in tokens:
        tokens.append(compact)
    digits = re.sub(r'\D', '', stripped or raw)
    if 4 <= len(digits) <= 12 and digits not in tokens:
        tokens.append(digits)
    return tokens

def refresh_xdevice_points(force=False):
    global _xdevice_points, _xdevice_loaded_at, _xdevice_error
    now_ts = time.time()
    if not force and _xdevice_points and now_ts - _xdevice_loaded_at < XDEVICE_TTL:
        return
    token = xdevice_token()
    if not token:
        _xdevice_error = "В настройках сервера нет ключа XDEVICE_TOKEN"
        return
    if not _xdevice_lock.acquire(timeout=50):
        return
    try:
        if not force and _xdevice_points and time.time() - _xdevice_loaded_at < XDEVICE_TTL:
            return
        response = requests.get(
            XDEVICE_URL,
            params={"limit": 5000},
            headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
            timeout=40,
        )
        response.raise_for_status()
        items = response.json()
        if isinstance(items, dict):
            for key in ('items', 'data', 'objects', 'result', 'rows'):
                if isinstance(items.get(key), list):
                    items = items[key]
                    break
        if not isinstance(items, list):
            raise RuntimeError("Ответ IoT не список")
        points = {}
        for obj in items:
            if not isinstance(obj, dict):
                continue
            state = obj.get("state") or {}
            lat, lon = find_coords(obj)
            point = {
                "lat": lat,
                "lon": lon,
                "online": bool(state.get("online")) if isinstance(state, dict) else False,
                "name": str(obj.get("name") or ""),
                "at": (state.get("time") if isinstance(state, dict) else 0) or 0,
                "object_id": str(obj.get("id") or obj.get("_id") or ""),
                "lock": bool(state.get("lock")) if isinstance(state, dict) else False,
            }
            if not point["object_id"] and lat is None:
                continue
            keys = set()
            for field in (obj.get("name"), obj.get("id"), obj.get("_id"), obj.get("description")):
                keys.update(iot_tokens(field))
            tags = obj.get("tags")
            if isinstance(tags, list):
                for tag in tags:
                    if isinstance(tag, str):
                        keys.update(iot_tokens(tag))
            for key in keys:
                if key:
                    points[key] = point
        _xdevice_points = points
        _xdevice_loaded_at = time.time()
        _xdevice_error = ""
        logger.info(f"Точки IoT обновлены: {len(items)} модулей, с координатами {len(set(id(p) for p in points.values()))}")
    except Exception as exc:
        _xdevice_error = "Точки IoT не обновились"
        logger.warning(f"Точки IoT не обновились: {exc}")
    finally:
        _xdevice_lock.release()

def xdevice_loop():
    while True:
        refresh_xdevice_points(force=True)
        time.sleep(XDEVICE_TTL)

def ensure_xdevice_loop():
    global _xdevice_started
    if _xdevice_started:
        return
    _xdevice_started = True
    threading.Thread(target=xdevice_loop, daemon=True).start()

def coord_number(value):
    try:
        if value is None or value == '':
            return None
        if isinstance(value, str):
            value = value.replace(',', '.').strip()
        number = float(value)
    except (TypeError, ValueError):
        return None
    if abs(number) >= 1000:
        scaled = number / 1000000.0
        if abs(scaled) <= 180:
            number = scaled
    if abs(number) < 1 or abs(number) > 180:
        return None
    return number

_LAT_KEYS = {'lat', 'latitude', 'gps-latitude', 'coord-latitude', 'last_valid_lat'}
_LON_KEYS = {'lon', 'lng', 'longitude', 'gps-longitude', 'coord-longitude', 'last_valid_lon'}

def find_coords(node, depth=0):
    if depth > 6 or node is None:
        return None, None
    if isinstance(node, dict):
        lat = lon = None
        for key, value in node.items():
            name = str(key).lower()
            if name in _LAT_KEYS:
                lat = coord_number(value)
            elif name in _LON_KEYS:
                lon = coord_number(value)
        if lat is not None and lon is not None:
            return lat, lon
        coords = node.get("coordinates")
        if isinstance(coords, (list, tuple)) and len(coords) >= 2 and not isinstance(coords[0], (list, dict)):
            first = coord_number(coords[0])
            second = coord_number(coords[1])
            if first is not None and second is not None:
                return second, first
        for value in node.values():
            lat, lon = find_coords(value, depth + 1)
            if lat is not None and lon is not None:
                return lat, lon
    elif isinstance(node, list):
        for value in node[:40]:
            lat, lon = find_coords(value, depth + 1)
            if lat is not None and lon is not None:
                return lat, lon
    return None, None

def fetch_object_coords(object_id):
    token = xdevice_token()
    if not token or not object_id:
        return None, None
    try:
        response = requests.get(
            XDEVICE_URL.rstrip('/') + '/' + str(object_id),
            headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
            timeout=20,
        )
        response.raise_for_status()
        return find_coords(response.json())
    except Exception as exc:
        logger.warning(f"Карточка IoT {object_id} не открылась: {exc}")
        return None, None

def lookup_bike_point(iot):
    ensure_xdevice_loop()
    refresh_xdevice_points(force=False)
    key = normalize_iot(iot)
    if not key and not iot_tokens(iot):
        return None
    point = None
    for token in iot_tokens(iot):
        point = _xdevice_points.get(token)
        if point:
            break
    if point is None:
        return None
    if point.get('lat') is None or point.get('lon') is None:
        lat, lon = fetch_object_coords(point.get('object_id'))
        if lat is not None and lon is not None:
            point['lat'] = lat
            point['lon'] = lon
    return point

def resolve_ticket_iot(iot, darks, gos):
    if normalize_iot(iot):
        return iot
    key = normalize_gos(gos)
    if not key:
        return ''
    try:
        mapping = bike_lookup()
    except Exception as exc:
        logger.warning(f'База вело для карты: {exc}')
        return ''
    direct = mapping.get((str(darks or '').strip(), key))
    if direct and direct.get('iot'):
        return direct['iot']
    for plate_key, item in mapping.items():
        if plate_key[1] == key and item.get('iot'):
            return item['iot']
    return ''

def normalize_gos(value):
    return re.sub(r'[^0-9A-Za-zА-Яа-я]', '', str(value or '')).upper()

_bike_lookup_cache = {'ts': 0, 'by_dark_gos': {}}
_bike_type_cache = {'ts': 0, 'by_gos': {}, 'supplier': {}}
_bike_rows_cache = {'ts': 0, 'rows': []}

def sheet_col(header, titles, fallback):
    wanted = {title.lower() for title in titles}
    for index, cell in enumerate(header or []):
        if str(cell).strip().lower() in wanted:
            return index
    return fallback

def load_bike_rows():
    now_ts = time.time()
    if _bike_rows_cache['rows'] and now_ts - _bike_rows_cache['ts'] < 600:
        return _bike_rows_cache['rows']
    rows = get_sheet_client().worksheet('База данных вело').get_all_values()
    _bike_rows_cache['rows'] = rows or []
    _bike_rows_cache['ts'] = now_ts
    return _bike_rows_cache['rows']

def bike_lookup():
    now_ts = time.time()
    if _bike_lookup_cache['by_dark_gos'] and now_ts - _bike_lookup_cache['ts'] < 600:
        return _bike_lookup_cache['by_dark_gos']
    mapping = {}
    try:
        rows = load_bike_rows()
        for row in rows[1:]:
            gos = normalize_gos(row[1] if len(row) > 1 else '')
            darks = row[4].strip() if len(row) > 4 else ''
            if not gos or not darks:
                continue
            mapping[(darks, gos)] = {
                'serial': row[0].strip() if row else '',
                'gos': (row[1] if len(row) > 1 else '').strip(),
                'iot': (row[2] if len(row) > 2 else '').strip(),
            }
        _bike_lookup_cache['by_dark_gos'] = mapping
        _bike_lookup_cache['ts'] = now_ts
    except Exception as e:
        logger.error(f'Ошибка чтения базы вело: {e}')
        if is_quota_error(e):
            raise
    return _bike_lookup_cache['by_dark_gos']

def fill_old_bike(tickets):
    mapping = bike_lookup()
    for ticket in tickets or []:
        bike = mapping.get((str(ticket.get('darks') or ''), normalize_gos(ticket.get('gos'))))
        if not bike:
            continue
        ticket['old_serial'] = bike['serial']
        ticket['old_gos'] = bike['gos']
        ticket['old_iot'] = bike['iot']

def ticket_label(ticket):
    gos = ((ticket or {}).get('gos') or '').strip()
    if gos:
        return gos
    return (ticket or {}).get('type') or 'Заявка'

def evacuation_reason(ticket):
    note = ((ticket or {}).get('note') or '').strip()
    for line in reversed(note.splitlines()):
        if 'ЭВАКУАЦИЯ' in line.upper():
            text = line.split(':', 1)[-1].strip()
            if text:
                return text
    return ((ticket or {}).get('desc') or '').strip()

def notify_dark_event(darks, text, uid='', event=''):
    if not darks or not text:
        return
    def run():
        try:
            requests.post(
                f'{BOT_API_URL}/notify_event',
                json={
                    'darks': str(darks),
                    'text': text,
                    'uid': uid or '',
                    'event': event or '',
                },
                timeout=8,
            )
        except Exception as e:
            logger.warning(f'Не ушло сообщение куратору {darks}: {e}')
    threading.Thread(target=run, daemon=True).start()

RATINGS_FILE = os.path.join(CACHE_DIR, "ratings.json")
_rating_lock = threading.Lock()

def load_ratings():
    data = _read_json(RATINGS_FILE, {})
    if not isinstance(data, dict):
        data = {}
    if not isinstance(data.get('asked'), dict):
        data['asked'] = {}
    if not isinstance(data.get('items'), list):
        data['items'] = []
    if not isinstance(data.get('sheet'), list):
        data['sheet'] = []
    return data

def save_ratings(data):
    _write_json(RATINGS_FILE, data)

def master_open_at(master, darks):
    left = []
    admin = get_admin_cache() or {}
    for direction in (admin.get('directions') or {}).values():
        for ticket in direction.get('tickets') or []:
            if str(ticket.get('darks')) != str(darks):
                continue
            if not same_master(ticket.get('master'), master):
                continue
            if is_active_status(ticket.get('status')):
                left.append(ticket.get('uid'))
    return left

def maybe_ask_rating(darks, master):
    if not darks or not master:
        return
    if clean_name(master) in ('iot-мастер', 'iot'):
        return
    left = master_open_at(master, darks)
    if left:
        logger.info(f'Оценка ждёт: у {master} на дарксторе {darks} ещё {len(left)} заявок')
        return
    day = get_msk_now().strftime('%Y-%m-%d')
    key = f'{day}|{darks}|{clean_name(master)}'
    info = (darks_ref or {}).get(str(darks)) or {}
    with _rating_lock:
        data = load_ratings()
        asked = data.get('asked') or {}
        if key in asked:
            return
        token = secrets.token_hex(4)
        asked[token] = {
            'key': key,
            'darks': str(darks),
            'master': master,
            'day': day,
            'address': info.get('address') or '',
        }
        asked[key] = token
        data['asked'] = asked
        save_ratings(data)
    def run():
        ok = False
        for _attempt in range(2):
            try:
                response = requests.post(
                    f'{BOT_API_URL}/ask_rating',
                    json={
                        'token': token,
                        'darks': str(darks),
                        'master': master,
                        'address': info.get('address') or '',
                    },
                    timeout=12,
                )
                body = {}
                try:
                    body = response.json() if response.content else {}
                except Exception:
                    body = {}
                ok = response.status_code == 200 and int(body.get('count') or 0) > 0
                if ok:
                    logger.info(f'Оценка ушла куратору {darks}, мастер {master}, получателей {body.get("count")}')
                    break
                logger.warning(f'Оценка не ушла: {response.status_code} {str(body)[:180]}')
            except Exception as exc:
                logger.warning(f'Оценка не ушла: {exc}')
            time.sleep(2)
        if ok:
            return
        with _rating_lock:
            stored = load_ratings()
            asked_now = stored.get('asked') or {}
            asked_now.pop(key, None)
            asked_now.pop(token, None)
            stored['asked'] = asked_now
            save_ratings(stored)
    threading.Thread(target=run, daemon=True).start()

def flush_rating_sheet():
    with _rating_lock:
        stored = load_ratings()
        pending = list(stored.get('sheet') or [])
    if not pending:
        return
    try:
        book = get_sheet_client()
        try:
            sheet = book.worksheet('Оценки кураторов')
        except Exception:
            sheet = book.add_worksheet('Оценки кураторов', rows=2000, cols=8)
            sheet.update(
                values=[['Дата', 'Время', 'Даркстор', 'Адрес', 'Мастер', 'Оценка', 'Комментарий', 'Куратор']],
                range_name='A1:H1',
            )
        rows = []
        for item in pending:
            day = item.get('day') or ''
            try:
                day = datetime.strptime(day, '%Y-%m-%d').strftime('%d.%m.%Y')
            except ValueError:
                pass
            rows.append([
                day,
                item.get('time') or '',
                item.get('darks') or '',
                item.get('address') or '',
                item.get('master') or '',
                item.get('score') or '',
                item.get('comment') or '',
                item.get('curator') or '',
            ])
        sheet.append_rows(rows, value_input_option='RAW')
    except Exception as exc:
        logger.warning(f'Оценки не записались в таблицу: {exc}')
        return
    with _rating_lock:
        stored = load_ratings()
        left = stored.get('sheet') or []
        stored['sheet'] = left[len(pending):] if len(left) >= len(pending) else []
        save_ratings(stored)

def master_rating_value(name):
    latest = {}
    for item in load_ratings().get('items') or []:
        if not same_master(item.get('master'), name):
            continue
        try:
            score = int(item.get('score'))
        except (TypeError, ValueError):
            continue
        key = (item.get('day'), str(item.get('darks')), item.get('chat_id') or item.get('curator'))
        latest[key] = score
    if not latest:
        return None
    return round(sum(latest.values()) / len(latest), 1)

def rating_same_person(row, token, chat_id, day, darks, master):
    if (row.get('chat_id') or '') != chat_id:
        return False
    if row.get('token') == token:
        return True
    return (
        row.get('day') == day
        and str(row.get('darks')) == str(darks)
        and row.get('master') == master
    )

@app.route('/api/service_rating', methods=['POST'])
def api_service_rating():
    if not bot_hook_ok():
        return jsonify({'success': False, 'error': 'Нет доступа'}), 403
    body = request.json or {}
    token = (body.get('token') or '').strip()
    comment = (body.get('comment') or '').strip()
    comment_only = bool(body.get('comment_only'))
    curator = (body.get('curator') or '').strip()
    chat_id = str(body.get('chat_id') or '').strip()
    try:
        score = int(body.get('score') or 0)
    except (TypeError, ValueError):
        score = 0
    if not token:
        return jsonify({'success': False, 'error': 'Нет оценки'})
    with _rating_lock:
        stored = load_ratings()
        ask = (stored.get('asked') or {}).get(token)
        if not isinstance(ask, dict):
            return jsonify({'success': False, 'error': 'Запрос устарел'})
        now = get_msk_now()
        day = ask.get('day') or now.strftime('%Y-%m-%d')
        darks = ask.get('darks') or ''
        master = ask.get('master') or ''
        items = list(stored.get('items') or [])
        existing = None
        for row in items:
            if rating_same_person(row, token, chat_id, day, darks, master):
                existing = row
                break
        if comment_only:
            if not existing:
                return jsonify({'success': False, 'error': 'Сначала поставьте оценку'})
            if existing.get('comment'):
                return jsonify({'success': False, 'error': 'Комментарий уже записан', 'already': True})
            if len(comment) < 3:
                return jsonify({'success': False, 'error': 'Напишите комментарий'})
            existing['comment'] = comment
            existing['curator'] = curator or existing.get('curator') or chat_id
            stored.setdefault('sheet', []).append(dict(existing))
            stored['sheet'] = stored['sheet'][-200:]
            save_ratings(stored)
            need_comment = False
            score = existing.get('score') or score
        else:
            if score < 1 or score > 5:
                return jsonify({'success': False, 'error': 'Нет оценки'})
            if existing:
                return jsonify({'success': False, 'error': 'Оценка уже принята', 'already': True})
            item = {
                'token': token,
                'day': day,
                'time': now.strftime('%H:%M'),
                'darks': darks,
                'address': ask.get('address') or '',
                'master': master,
                'score': score,
                'comment': comment,
                'curator': curator or chat_id,
                'chat_id': chat_id,
                'date_ru': now.strftime('%d.%m.%Y'),
            }
            items.append(item)
            stored['items'] = items[-500:]
            need_comment = score < 5
            if not need_comment:
                stored.setdefault('sheet', []).append(item)
                stored['sheet'] = stored['sheet'][-200:]
            save_ratings(stored)
    if not need_comment:
        threading.Thread(target=flush_rating_sheet, daemon=True).start()
    return jsonify({'success': True, 'score': score, 'need_comment': need_comment})

def supplier_of(gos_value):
    return (_bike_type_cache.get('supplier') or {}).get(normalize_gos(gos_value), '')

def get_bike_type_by_gos_map():
    now_ts = time.time()
    if _bike_type_cache['by_gos'] and now_ts - _bike_type_cache['ts'] < 600:
        return _bike_type_cache['by_gos']
    mapping = {}
    suppliers = {}
    try:
        rows = load_bike_rows()
        header = rows[0] if rows else []
        gos_i = sheet_col(header, ['гос номер', 'госномер'], 1)
        kind_i = sheet_col(header, ['тип вело'], 5)
        company_i = sheet_col(header, ['компания', 'поставщик'], 8)
        for row in rows[1:]:
            gos = normalize_gos(row[gos_i] if len(row) > gos_i else '')
            kind = row[kind_i].strip() if len(row) > kind_i else ''
            company = row[company_i].strip() if len(row) > company_i else ''
            if gos and kind:
                mapping[gos] = kind
            if gos and company:
                suppliers[gos] = company
        _bike_type_cache['by_gos'] = mapping
        _bike_type_cache['supplier'] = suppliers
        _bike_type_cache['ts'] = now_ts
    except Exception as e:
        logger.error(f"Ошибка чтения типов вело: {e}")
        if is_quota_error(e):
            raise
    return _bike_type_cache['by_gos']

def _read_tickets_from_google():
    global darks_ref
    sheet_client = get_sheet_client()
    tickets = []
    darks_ref = load_darks_reference()
    bike_types = get_bike_type_by_gos_map()
    
    for sheet_name in ('Заявки', 'Заявки бот'):
        try:
            worksheet = sheet_client.worksheet(sheet_name)
            rows = worksheet.get_all_values()
            if len(rows) > 1:
                for idx, row in enumerate(rows[1:], start=2):
                    if len(row) < 6 or not row[0].strip():
                        continue
                    if row[0].strip() in ('Номер дарка', 'Номер_дарка'):
                        continue
                    darks_num = row[0].strip()
                    status_raw = row[7].strip() if len(row) > 7 else ''
                    note = row[9].strip() if len(row) > 9 else ''
                
                    if status_raw in ['Выполнено', '✅ Выполнено', 'done']:
                        status = 'done'
                    elif status_raw in ['🔵 Доделать', 'Доделать']:
                        status = 'todo'
                    elif status_raw in ['⏹️ Обработано', 'Обработано', 'Вело отсутствует']:
                        status = 'processed'
                    elif status_raw in ['🔧 Эвакуация', 'Эвакуация']:
                        status = 'evacuation'
                    else:
                        status = 'pending'
                
                    is_done = status == 'done'
                    created_str = row[5].strip() if len(row) > 5 else ''
                    timer_from = row[14].strip() if len(row) > 14 else ''
                    hours_since = get_hours_since(timer_from or created_str)
                    bike_type = row[1].strip() if len(row) > 1 else ''
                    bike_subtype = row[8].strip() if len(row) > 8 else ''
                    gos_value = row[3].strip() if len(row) > 3 else ''
                    if bike_type == 'Электровелосипед':
                        display_type = bike_types.get(normalize_gos(gos_value)) or 'Неопределен тип вело, ошибка в гос номере'
                    elif bike_subtype:
                        display_type = bike_subtype
                    else:
                        display_type = bike_type or 'Не указан'
                    supplier = supplier_of(gos_value)
                    uid = row[10].strip() if len(row) > 10 else ''
                    if not uid and created_str:
                        uid = generate_ticket_id(created_str)
                        if uid:
                            try:
                                worksheet.update_cell(idx, 11, uid)
                            except:
                                pass
                    direction = row[11].strip() if len(row) > 11 else ''
                    if not direction and darks_num in darks_ref:
                        direction = darks_ref[darks_num].get('direction', '')
                    display_desc = row[2].strip() if len(row) > 2 else ''
                    if status == 'todo' and note and 'ЗАБРАЛИ:' in note:
                        match = re.search(r'ЗАБРАЛИ:\s*(\d+)', note)
                        if match:
                            count = match.group(1)
                            display_desc = f'Вернуть {count} АКБ (забирали на ремонт)'
                
                    evac_reason = ''
                    if status == 'evacuation':
                        if note and not note.startswith('ЭВАКУАЦИЯ:'):
                            note = f'ЭВАКУАЦИЯ: {note}'
                        evac_reason = evacuation_reason({'note': note, 'desc': display_desc})
                
                    tickets.append({
                        'source': sheet_name,
                        'darks': darks_num,
                        'type': bike_type,
                        'bike_type': display_type,
                        'bike_subtype': bike_subtype,
                        'supplier': supplier,
                        'desc': display_desc,
                        'gos': gos_value,
                        'contact': row[4].strip() if len(row) > 4 else '',
                        'created': created_str,
                        'timer_from': timer_from,
                        'hours_since': hours_since,
                        'master': row[6].strip() if len(row) > 6 else '',
                        'status': status,
                        'note': note,
                        'evac_reason': evac_reason,
                        'uid': uid,
                        'row_index': idx,
                        'address': darks_ref.get(darks_num, {}).get('address', ''),
                        'direction': direction,
                        'coords': darks_ref.get(darks_num, {}).get('coords', ''),
                        'sent_date': '',
                        'is_done': is_done,
                        'is_active': not is_done,
                        'parts': row[9].strip() if len(row) > 9 else '',
                        'display_desc': display_desc
                    })
        except Exception as e:
            logger.error(f"Ошибка чтения '{sheet_name}': {e}")
            if is_quota_error(e):
                raise

    
    try:
        worksheet = sheet_client.worksheet("Импорт М4")
        rows = worksheet.get_all_values()
        if len(rows) > 1:
            for idx, row in enumerate(rows[1:], start=2):
                if len(row) < 15:
                    continue
                obj = row[4].strip() if len(row) > 4 else ''
                darks_match = re.search(r'^(\d{4})', obj)
                darks_num = darks_match.group(1) if darks_match else ''
                model = row[8].strip() if len(row) > 8 else ''
                bike_type = 'Не указан'
                if 'Электровелосипед' in model or 'электро' in model:
                    bike_type = 'Электровелосипед'
                elif 'Аккумуляторная батарея' in model or 'аккумуляторная' in model or '🔋' in model:
                    bike_type = 'Аккумуляторная батарея'
                elif 'Зарядное устройство' in model or 'зарядное' in model:
                    bike_type = 'Зарядное устройство'
                elif 'Багажник' in model:
                    bike_type = 'Багажник'
                elif 'Номерной знак' in model:
                    bike_type = 'Номерной знак'
                elif 'IoT' in model:
                    bike_type = 'IoT'
                gos = row[9].strip() if len(row) > 9 else ''
                uid = row[1].strip() if len(row) > 1 else ''
                status_raw = row[3].strip() if len(row) > 3 else ''
                status_l = row[11].strip() if len(row) > 11 else ''
                note = row[12].strip() if len(row) > 12 else ''
                
                if status_raw == 'Решено' or status_l in ('Выполнено', '✅ Выполнено'):
                    status = 'done'
                    is_done = True
                elif status_l in ('Эвакуация', '🔧 Эвакуация'):
                    status = 'evacuation'
                    is_done = False
                    if note and not note.startswith('ЭВАКУАЦИЯ:'):
                        note = f'ЭВАКУАЦИЯ: {note}'
                elif status_l in ('Доделать', '🔵 Доделать'):
                    status = 'todo'
                    is_done = False
                elif status_l in ('Обработано', '⏹️ Обработано', 'Вело отсутствует'):
                    status = 'processed'
                    is_done = False
                else:
                    status = 'pending'
                    is_done = False
                created_str = row[6].strip() if len(row) > 6 else ''
                hours_since = get_hours_since(created_str)
                sent_date = row[13].strip() if len(row) > 13 else ''
                tickets.append({
                    'source': 'Импорт М4',
                    'darks': darks_num,
                    'type': bike_type,
                    'bike_type': bike_type,
                    'supplier': supplier_of(gos),
                    'desc': row[7].strip() if len(row) > 7 else '',
                    'gos': gos,
                    'created': created_str,
                    'hours_since': hours_since,
                    'master': row[13].strip() if len(row) > 13 else '',
                    'status': status,
                    'note': note,
                    'uid': uid,
                    'row_index': idx,
                    'address': darks_ref.get(darks_num, {}).get('address', ''),
                    'direction': darks_ref.get(darks_num, {}).get('direction', ''),
                    'coords': darks_ref.get(darks_num, {}).get('coords', ''),
                    'sent_date': sent_date,
                    'is_done': is_done,
                    'is_active': not is_done,
                    'parts': row[12].strip() if len(row) > 12 else ''
                })
    except Exception as e:
        logger.error(f"Ошибка чтения 'Импорт М4': {e}")
        if is_quota_error(e):
            raise
    
    attach_iot(tickets)
    return tickets

def attach_iot(tickets):
    try:
        mapping = bike_lookup()
    except Exception as e:
        if is_quota_error(e):
            raise
        logger.error(f'Не удалось подставить IoT: {e}')
        return
    for ticket in tickets or []:
        bike = mapping.get((str(ticket.get('darks') or ''), normalize_gos(ticket.get('gos'))))
        if bike and bike.get('iot'):
            ticket['iot'] = bike['iot']

def store_ticket_memory(tickets, fresh=False):
    global _tickets_mem, _tickets_mem_at, uid_index
    _tickets_mem = list(tickets or [])
    uid_index = build_uid_index(_tickets_mem)
    if fresh:
        _tickets_mem_at = time.time()
    return _tickets_mem

def publish_caches(tickets):
    strip_cleared_masters(tickets)
    apply_pending_queue(tickets)
    enrich_tickets(tickets)
    save_admin_cache(tickets)
    for master in MASTERS:
        master_tickets = [t for t in tickets if same_master(t.get('master'), master) and is_active_status(t.get('status'))]
        save_master_cache(master, master_tickets, '')

_last_google_import = 0

def _tickets_from_db(previous):
    rows = db.load_tickets()
    if not rows:
        return None
    for ticket in rows:
        ticket['hours_since'] = ticket_hours(ticket)
        ticket['is_done'] = ticket.get('status') == 'done'
        ticket['is_active'] = is_active_status(ticket.get('status'))
    merge_local_tickets(previous, rows)
    apply_pending_queue(rows)
    strip_cleared_masters(rows)
    enrich_tickets(rows)
    return rows

def get_tickets_from_sheets(force=False):
    """Рабочий список заявок. Сначала база, таблицу открываем только если базы нет."""
    global _sheet_backoff_until
    now = time.time()
    if _tickets_mem and not force and now - _tickets_mem_at < 120:
        return _tickets_mem
    if not _sheet_read_lock.acquire(blocking=False):
        return _tickets_mem
    try:
        now = time.time()
        if _tickets_mem and not force and now - _tickets_mem_at < 120:
            return _tickets_mem
        previous = list(_tickets_mem)
        rows = _tickets_from_db(previous)
        if rows:
            store_ticket_memory(rows, fresh=True)
            logger.info(f'Заявки взяты из базы: {len(rows)}')
            return _tickets_mem
        if now < _sheet_backoff_until:
            logger.info('Google ограничил чтение, берём прошлую копию')
            return _tickets_mem
        tickets = _read_tickets_from_google()
        merge_local_tickets(previous, tickets)
        apply_pending_queue(tickets)
        strip_cleared_masters(tickets)
        store_ticket_memory(tickets, fresh=True)
        db.mirror_tickets(_tickets_mem)
        logger.info(f'База недоступна, таблица прочитана: {len(tickets)} заявок')
        return _tickets_mem
    except Exception as e:
        if is_quota_error(e):
            _sheet_backoff_until = time.time() + 75
            logger.warning('Лимит Google. Пауза 75 секунд, работаем на прошлой копии.')
        else:
            logger.error(f'Ошибка чтения заявок: {e}')
        return _tickets_mem
    finally:
        _sheet_read_lock.release()

def import_new_tickets_from_google():
    """Забрать с таблицы только новые заявки. Старые статусы из таблицы не затирают базу."""
    global _last_google_import, _sheet_backoff_until
    now = time.time()
    if now < _sheet_backoff_until or now - _last_google_import < 60:
        return 0
    if not _tickets_mem:
        return 0
    if not _sheet_read_lock.acquire(blocking=False):
        return 0
    try:
        _last_google_import = time.time()
        fresh = _read_tickets_from_google()
        fresh_by_uid = {t.get('uid'): t for t in fresh if t.get('uid')}
        filled = 0
        for ticket in _tickets_mem:
            src = fresh_by_uid.get(ticket.get('uid'))
            if not src:
                continue
            sheet_master = (src.get('master') or '').strip()
            if sheet_master and not (ticket.get('master') or '').strip():
                if str(ticket.get('uid') or '') in cleared_master_uids():
                    continue
                if ticket.get('status') in ('processed', 'fail', 'evacuation'):
                    continue
                note_low = (ticket.get('note') or '').lower()
                if 'оспорено куратором' in note_low or 'вело отсутствует' in note_low:
                    continue
                stamp = ticket.get('_local_ts') or 0
                if stamp and time.time() - stamp < 900:
                    continue
                if queue_clears_master(ticket.get('uid')):
                    continue
                if clean_name(sheet_master) in RETIRED_MASTERS:
                    continue
                ticket['master'] = sheet_master
                ticket['_local_ts'] = time.time()
                db.mirror_ticket(ticket)
                filled += 1
        known = {t.get('uid') for t in _tickets_mem if t.get('uid')}
        new_ones = [t for t in fresh if t.get('uid') and t.get('uid') not in known]
        for ticket in new_ones:
            db.mirror_ticket(ticket)
        if not new_ones and not filled:
            logger.info('Новых заявок в таблице нет')
            return 0
        merged = list(_tickets_mem) + new_ones
        apply_pending_queue(merged)
        store_ticket_memory(merged, fresh=True)
        publish_caches(merged)
        logger.info(f'Из таблицы добавлены новые заявки: {len(new_ones)}, мастеров дописано: {filled}')
        return len(new_ones)
    except Exception as e:
        if is_quota_error(e):
            _sheet_backoff_until = time.time() + 75
            logger.warning('Лимит Google при заборе новых заявок. Пауза 75 секунд.')
        else:
            logger.error(f'Не удалось забрать новые заявки из таблицы: {e}')
        return 0
    finally:
        _sheet_read_lock.release()

def restore_statuses_from_sheet():
    """Вернуть в базу статус и мастера из таблицы. Запись сразу, не через очередь."""
    global _last_google_import
    fresh = _read_tickets_from_google()
    counts = {}
    for src in fresh:
        key = src.get('status') or 'пусто'
        counts[key] = counts.get(key, 0) + 1
    if not _tickets_mem:
        store_ticket_memory(fresh, fresh=True)
    else:
        by_uid = {t.get('uid'): t for t in fresh if t.get('uid')}
        known = set()
        for ticket in _tickets_mem:
            src = by_uid.get(ticket.get('uid'))
            if not src:
                continue
            known.add(ticket.get('uid'))
            ticket['status'] = src.get('status') or 'pending'
            ticket['master'] = src.get('master') or ''
            ticket['note'] = src.get('note') or ''
            ticket['evac_reason'] = src.get('evac_reason') or ''
            ticket['timer_from'] = src.get('timer_from') or ''
            ticket['is_done'] = ticket['status'] == 'done'
            ticket['is_active'] = is_active_status(ticket['status'])
            ticket['hours_since'] = ticket_hours(ticket)
            ticket['_local_ts'] = time.time()
        for src in fresh:
            if src.get('uid') and src.get('uid') not in known:
                _tickets_mem.append(src)
        store_ticket_memory(_tickets_mem, fresh=True)
    changed = db.write_statuses_now(_tickets_mem)
    publish_caches(_tickets_mem)
    _last_google_import = time.time()
    logger.info(f'Статусы возвращены из таблицы, записано {changed}, по статусам {counts}')
    return {'count': changed, 'from_sheet': counts}

def reconcile_from_google():
    """Кнопка руководителя: сайт и база становятся как Google Таблица.
    Свои несохранённые записи сначала отправляются в таблицу, они важнее."""
    try:
        flush_queue()
    except Exception as exc:
        logger.warning(f'Очередь перед сверкой не ушла: {exc}')
    fresh = _read_tickets_from_google()
    fresh = [ticket for ticket in fresh if ticket.get('uid')]
    if not fresh:
        raise RuntimeError('Таблица не открылась или в ней нет заявок. Базу не менял.')
    apply_pending_queue(fresh)
    for ticket in fresh:
        ticket['is_done'] = ticket.get('status') == 'done'
        ticket['is_active'] = is_active_status(ticket.get('status'))
        ticket['hours_since'] = ticket_hours(ticket)
    closed = db.reconcile_from_sheet(fresh)
    store_ticket_memory(fresh, fresh=True)
    publish_caches(fresh)
    open_count = len([ticket for ticket in fresh if ticket.get('status') != 'done'])
    logger.info(f'Сверка с таблицей: строк {len(fresh)}, открытых {open_count}, закрыто лишних {closed}')
    return {'sheet': len(fresh), 'open': open_count, 'closed': closed}


def refresh_admin_cache():
    try:
        flush_queue()
        tickets = get_tickets_from_sheets()
        save_admin_cache(tickets)
        logger.info(f"✅ Админ кэш обновлен: {len(tickets)} заявок")
    except Exception as e:
        logger.error(f"Ошибка обновления админ кэша: {e}")

def refresh_master_cache(master_name):
    try:
        tickets = get_tickets_from_sheets()
        master_tickets = [t for t in tickets if same_master(t.get('master'), master_name) and is_active_status(t.get('status'))]
        enrich_tickets(master_tickets)
        save_master_cache(master_name, master_tickets, '')
        logger.info(f"✅ Кэш для {master_name} обновлен: {len(master_tickets)} заявок")
        return len(master_tickets)
    except Exception as e:
        logger.error(f"Ошибка обновления кэша {master_name}: {e}")
        return 0

def refresh_all_master_caches():
    try:
        tickets = get_tickets_from_sheets()
        publish_caches(tickets)
        logger.info("✅ Кэши всех мастеров обновлены из копии")
    except Exception as e:
        logger.error(f"Ошибка обновления кэшей мастеров: {e}")

def batch_update_masters(changes):
    """Сразу в базу. В таблицу уходит очередью, страница больше не ждёт Google."""
    updated_count = 0
    for change in changes or []:
        uid = (change or {}).get('uid')
        if not uid:
            continue
        master = change.get('master') or ''
        if str(master).strip():
            forget_master_cleared(uid)
        else:
            remember_master_cleared(uid)
        source = change.get('source') or source_of(uid)
        found = False
        for ticket in _tickets_mem:
            if ticket.get('uid') != uid:
                continue
            ticket['master'] = master
            ticket['_local_ts'] = time.time()
            db.mirror_ticket(ticket)
            source = ticket.get('source') or source
            found = True
        if not found:
            cache = get_admin_cache() or {}
            for dir_data in cache.get('directions', {}).values():
                for ticket in dir_data.get('tickets', []):
                    if ticket.get('uid') != uid:
                        continue
                    ticket['master'] = master
                    ticket['_local_ts'] = time.time()
                    db.mirror_ticket(ticket)
                    source = ticket.get('source') or source
                    found = True
                    break
                if found:
                    break
        add_to_queue({
            'uid': uid,
            'source': source,
            'type': 'assign',
            'data': {'master': master, 'skip_report': True},
        })
        updated_count += 1
    if _tickets_mem:
        store_ticket_memory(_tickets_mem, fresh=False)
        publish_caches(_tickets_mem)
    return updated_count

def clear_masters(master_name=None):
    wanted = clean_name(master_name) if master_name else ''
    tickets = list(_tickets_mem or get_tickets_from_sheets())
    grouped = {}
    cleared = 0
    for ticket in tickets:
        if not ticket.get('master') or not is_active_status(ticket.get('status')):
            continue
        if wanted and clean_name(ticket.get('master')) != wanted:
            continue
        row_idx = ticket.get('row_index')
        source = ticket.get('source') or 'Заявки'
        if not row_idx:
            continue
        column = 'N' if source == 'Импорт М4' else 'G'
        grouped.setdefault(source, []).append({'range': f'{column}{row_idx}', 'values': [['']]})
        ticket['master'] = ''
        ticket['_local_ts'] = time.time()
        remember_master_cleared(ticket.get('uid'))
        db.mirror_ticket(ticket)
        cleared += 1
    if grouped:
        book = get_sheet_client()
        for source, updates in grouped.items():
            book.worksheet(source).batch_update(updates)
            logger.info(f"✅ Снято {len(updates)} мастеров в '{source}'")
    store_ticket_memory(tickets, fresh=False)
    publish_caches(tickets)
    return cleared

def clear_all_masters():
    try:
        return clear_masters()
    except Exception as e:
        logger.error(f"Ошибка снятия всех мастеров: {e}")
        return 0

def clear_one_master(master_name):
    if not clean_name(master_name):
        return 0
    try:
        return clear_masters(master_name)
    except Exception as e:
        logger.error(f"Ошибка снятия мастера {master_name}: {e}")
        return 0

def update_status_in_google_sheets(uid, status, note=''):
    try:
        sheet_client = get_sheet_client()
        
        found = find_uid_in_sheets(uid)
        if not found:
            logger.error(f"UID {uid} не найден в Google Sheets")
            return False
        
        row_idx = found['row_index']
        source = found['source']
        
        if source in ('Заявки', 'Заявки бот'):
            worksheet = sheet_client.worksheet(source)
            worksheet.update_cell(row_idx, 8, status)
            if note:
                worksheet.update_cell(row_idx, 10, note)
            logger.info(f"✅ Обновлён статус заявки {uid} в '{source}': {status}")
            
        elif source == 'Импорт М4':
            worksheet = sheet_client.worksheet("Импорт М4")
            status_map = {
                '🟡 В работе': 'В работе',
                '✅ Выполнено': 'Выполнено',
                '🔵 Доделать': 'Доделать',
                '🔧 Эвакуация': 'Эвакуация'
            }
            new_status = status_map.get(status, 'В работе')
            worksheet.update_cell(row_idx, 12, new_status)
            if note:
                worksheet.update_cell(row_idx, 13, note)
            logger.info(f"✅ Обновлён статус заявки {uid} в 'Импорт М4': {new_status}")
        
        return True
    except Exception as e:
        logger.error(f"Ошибка обновления статуса в Google Sheets: {e}")
        return False

def find_uid_in_sheets(uid):
    if not uid:
        return None
    item = uid_index.get(uid)
    if item and item.get('row_index'):
        ticket = item.get('ticket') or {}
        return {'source': ticket.get('source') or item.get('source') or 'Заявки', 'row_index': item.get('row_index')}
    for ticket in list(_tickets_mem or []):
        if ticket.get('uid') == uid and ticket.get('row_index'):
            return {'source': ticket.get('source') or 'Заявки', 'row_index': ticket.get('row_index')}
    cached = find_cached_ticket(uid)
    if cached and cached.get('row_index'):
        return {'source': cached.get('source') or 'Заявки', 'row_index': cached.get('row_index')}
    return None

# ============================================================
# ОЧЕРЕДЬ ЗАДАЧ
# ============================================================
def get_queue_path():
    return os.path.join(CACHE_DIR, "queue.json")

def read_queue():
    with queue_lock:
        return _load_queue()

def write_queue(queue_data):
    with queue_lock:
        try:
            _save_queue(queue_data)
            return True
        except:
            return False

def _load_queue():
    try:
        path = get_queue_path()
        if os.path.exists(path):
            with open(path, 'r', encoding='utf-8') as f:
                data = json.load(f)
                if isinstance(data, dict) and isinstance(data.get('tasks'), list):
                    return data
    except:
        pass
    return {'tasks': [], 'last_sync': None}

def _save_queue(queue_data):
    path = get_queue_path()
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(queue_data, f, ensure_ascii=False, indent=2)

def add_to_queue(task):
    with queue_lock:
        if task.get('uid') and not task.get('row_index'):
            found = find_uid_in_sheets(task.get('uid'))
            if found:
                task['row_index'] = found.get('row_index')
                if not task.get('source'):
                    task['source'] = found.get('source')
        queue_data = _load_queue()
        uid = task.get('uid')
        task['qid'] = f"{time.time():.6f}-{uid}"
        queue_data['tasks'] = [t for t in queue_data.get('tasks', []) if t.get('uid') != uid]
        queue_data['tasks'].append(task)
        _save_queue(queue_data)
    db.mirror_event(task)
    logger.info(f"✅ Задача добавлена в очередь: {uid}")

def remove_ticket_from_master_cache(master_name, uid):
    if not master_name or not uid:
        return
    cache_data = get_master_cache(master_name)
    if not cache_data:
        return
    tickets = [t for t in cache_data.get('tickets', []) if t.get('uid') != uid]
    save_master_cache(master_name, tickets, cache_data.get('sent_date', ''))

def find_cached_ticket(uid):
    admin_cache = get_admin_cache() or {}
    for dir_data in admin_cache.get('directions', {}).values():
        for t in dir_data.get('tickets', []):
            if t.get('uid') == uid:
                return t
    if uid in uid_index:
        return uid_index[uid].get('ticket')
    return None

def source_of(uid):
    ticket = find_cached_ticket(uid) or {}
    source = ticket.get('source')
    if source in ('Заявки', 'Заявки бот', 'Импорт М4'):
        return source
    if uid in uid_index:
        return uid_index[uid].get('source') or 'Заявки'
    return 'Заявки'

def mark_bike_missing(uid, actor_name):
    now = get_msk_now()
    timer_from = now.strftime('%Y-%m-%d %H:%M:%S')
    line = f"Вело отсутствует {now.strftime('%d.%m %H:%M')}, {actor_name}"
    ticket = find_cached_ticket(uid)
    previous = ''
    master_name = ''
    if ticket:
        previous = (ticket.get('note') or '').strip()
        master_name = ticket.get('master') or ''
    note = f'{previous}\n{line}'.strip() if previous else line
    update_ticket_in_admin_cache(uid, 'processed', note, extra={
        'master': '',
        'timer_from': timer_from,
        'hours_since': 0,
        'is_done': False,
        'is_active': True
    })
    remember_master_cleared(uid)
    remove_ticket_from_master_cache(master_name, uid)
    if actor_name and actor_name != master_name:
        remove_ticket_from_master_cache(actor_name, uid)
    refresh_master_caches_from_admin()
    add_to_queue({
        'uid': uid,
        'source': source_of(uid),
        'type': 'fail',
        'data': {
            'master': actor_name,
            'note': note,
            'detail': line,
            'parts': line,
            'timer_from': timer_from,
            'darks_number': (ticket or {}).get('darks', ''),
            'gos': (ticket or {}).get('gos', ''),
            'desc': (ticket or {}).get('desc', ''),
            'skip_report': False
        }
    })
    notify_dark_event((ticket or {}).get('darks', ''), f'⏹️ {ticket_label(ticket)} нет вело на дарксторе, заявка перенесена', uid, 'missing')
    return note

def apply_queue_to_sheets(tasks):
    skipped = set()
    buckets = {'Заявки': [], 'Заявки бот': []}
    for task in tasks:
        source = task.get('source')
        if source not in buckets:
            continue
        updates = buckets[source]
        uid = task.get('uid')
        row_idx = task.get('row_index') or get_ticket_row_by_uid(uid)
        if not row_idx:
            found = find_uid_in_sheets(uid)
            if not found or found.get('source') != source:
                if task.get('qid'):
                    skipped.add(task.get('qid'))
                logger.warning(f'Очередь не нашла строку заявки {uid}, повторю позже')
                continue
            row_idx = found['row_index']
        data = task.get('data') or {}
        task_type = task.get('type')
        if task_type == 'assign':
            updates.append({'range': f'G{row_idx}', 'values': [[data.get('master') or '']]})
        elif task_type == 'fail':
            updates.append({'range': f'G{row_idx}', 'values': [['']]})
            updates.append({'range': f'H{row_idx}', 'values': [['⏹️ Обработано']]})
            updates.append({'range': f'J{row_idx}', 'values': [[data.get('note', '')]]})
            updates.append({'range': f'O{row_idx}', 'values': [[data.get('timer_from', '')]]})
        elif task_type in ('done', 'replace_yes', 'transit_replace', 'evacuation_and_replace', 'transit_bulk_close'):
            updates.append({'range': f'H{row_idx}', 'values': [['✅ Выполнено']]})
            parts = data.get('parts') or data.get('extra') or ''
            if task_type == 'done':
                parts = data.get('parts') or ''
            elif task_type == 'evacuation_and_replace':
                parts = 'Заменен'
            if parts:
                updates.append({'range': f'J{row_idx}', 'values': [[parts]]})
        elif task_type == 'evacuation':
            updates.append({'range': f'G{row_idx}', 'values': [['']]})
            updates.append({'range': f'H{row_idx}', 'values': [['🔧 Эвакуация']]})
            updates.append({'range': f'J{row_idx}', 'values': [[data.get('reason') or 'Эвакуация']]})
        elif task_type == 'taken_no_replace':
            updates.append({'range': f'H{row_idx}', 'values': [['🔵 Доделать']]})
            updates.append({'range': f'J{row_idx}', 'values': [[f"ЗАБРАЛИ: {data.get('parts', '')} АКБ"]]})
        elif task_type == 'replace_no':
            updates.append({'range': f'H{row_idx}', 'values': [['🔵 Доделать']]})
            updates.append({'range': f'J{row_idx}', 'values': [[data.get('reason') or 'Куратор не предоставил']]})
        elif task_type == 'status_update':
            if not data.get('note_only'):
                updates.append({'range': f'H{row_idx}', 'values': [[data.get('status') or '🟡 В работе']]})
            if data.get('note'):
                updates.append({'range': f'J{row_idx}', 'values': [[data.get('note')]]})
            if data.get('write_master'):
                updates.append({'range': f'G{row_idx}', 'values': [[data.get('master') or '']]})
    import_updates = []
    for task in tasks:
        if task.get('source') != 'Импорт М4':
            continue
        uid = task.get('uid')
        row_idx = task.get('row_index') or get_ticket_row_by_uid(uid)
        if not row_idx:
            found = find_uid_in_sheets(uid)
            if not found or found.get('source') != 'Импорт М4':
                if task.get('qid'):
                    skipped.add(task.get('qid'))
                logger.warning(f'Очередь не нашла строку импорта {uid}, повторю позже')
                continue
            row_idx = found['row_index']
        data = task.get('data') or {}
        task_type = task.get('type')
        status_map_import = {
            '🟡 В работе': 'В работе',
            '✅ Выполнено': 'Выполнено',
            '🔵 Доделать': 'Доделать',
            '⏹️ Обработано': 'Обработано',
            '🔧 Эвакуация': 'Эвакуация'
        }
        clear_master = False
        if task_type == 'assign':
            import_updates.append({'range': f'N{row_idx}', 'values': [[data.get('master') or '']]})
            continue
        if task_type == 'fail':
            sheet_status, sheet_note = 'Обработано', data.get('note') or ''
            clear_master = True
        elif task_type in ('done', 'replace_yes', 'transit_replace', 'evacuation_and_replace', 'transit_bulk_close'):
            sheet_status = 'Выполнено'
            sheet_note = 'Заменен' if task_type == 'evacuation_and_replace' else (data.get('parts') or data.get('extra') or '')
        elif task_type == 'evacuation':
            sheet_status, sheet_note = 'Эвакуация', data.get('reason') or 'Эвакуация'
            clear_master = True
        elif task_type == 'taken_no_replace':
            sheet_status, sheet_note = 'Доделать', f"ЗАБРАЛИ: {data.get('parts', '')} АКБ"
        elif task_type == 'replace_no':
            sheet_status, sheet_note = 'Доделать', data.get('reason') or 'Куратор не предоставил'
        elif task_type == 'status_update':
            sheet_note = data.get('note') or ''
            if data.get('note_only'):
                if sheet_note:
                    import_updates.append({'range': f'M{row_idx}', 'values': [[sheet_note]]})
                continue
            sheet_status = status_map_import.get(data.get('status'), 'В работе')
            clear_master = data.get('status') == '⏹️ Обработано' or (data.get('write_master') and not (data.get('master') or '').strip())
        else:
            continue
        import_updates.append({'range': f'L{row_idx}', 'values': [[sheet_status]]})
        if sheet_note:
            import_updates.append({'range': f'M{row_idx}', 'values': [[sheet_note]]})
        if clear_master:
            import_updates.append({'range': f'N{row_idx}', 'values': [['']]})
    written = 0
    sheet_client = get_sheet_client()
    for sheet_name, sheet_updates in buckets.items():
        if not sheet_updates:
            continue
        sheet_client.worksheet(sheet_name).batch_update(sheet_updates)
        written += len(sheet_updates)
        logger.info(f"✅ Очередь записана в {sheet_name}: {len(sheet_updates)} ячеек")
    if import_updates:
        get_sheet_client().worksheet('Импорт М4').batch_update(import_updates)
        written += len(import_updates)
        logger.info(f"✅ Очередь записана в Импорт М4: {len(import_updates)} ячеек")
    return skipped

def append_master_history(master, action, data):
    if not master:
        return
    labels = {
        'done': 'Выполнено',
        'fail': 'Велосипеда не было',
        'evacuation': 'Эвакуация',
        'evacuation_and_replace': 'Замена велосипеда',
        'transit_replace': 'Замена велосипеда',
        'replace_yes': 'Замена',
        'replace_no': 'Куратор не дал',
        'taken_no_replace': 'Забрали без замены',
    }
    raw = read_cache(f"history_{master}.json")
    if isinstance(raw, dict):
        items = raw.get('items') or raw.get('history') or []
    elif isinstance(raw, list):
        items = raw
    else:
        items = []
    data = data or {}
    detail = data.get('detail') or data.get('parts') or data.get('reason') or ''
    if detail == 'Отказ от эвакуации':
        labels['done'] = 'Отказ от эвакуации'
    hours = None
    ticket = find_cached_ticket(data.get('uid')) if data.get('uid') else None
    if ticket:
        try:
            hours = round(float(ticket_hours(ticket)), 1)
        except Exception:
            hours = None
    items.append({
        'timestamp': get_msk_now().strftime('%Y-%m-%d %H:%M:%S'),
        'hours': hours,
        'action': {
            'action': action,
            'label': 'Отказ от эвакуации' if detail == 'Отказ от эвакуации' else labels.get(action, 'Отметка'),
            'parts': '' if detail == 'Отказ от эвакуации' else detail,
            'detail': detail,
            'darks_number': data.get('darks_number', ''),
            'gos': data.get('gos') or ((data.get('old_data') or {}).get('gos') or ''),
        }
    })
    write_cache(f"history_{master}.json", {'items': items[-300:]})

SHIFT_REPORTS_FILE = 'shift_reports.json'
SHIFT_SEEN_FILE = 'shift_seen.json'
SHIFT_STATUSES = {
    'done': 'Все запланированные выезды выполнены',
    'partial': 'Выполнено частично',
    'failed': 'Не выполнен',
}

def shift_report_items():
    raw = read_cache(SHIFT_REPORTS_FILE) or {}
    items = raw.get('items') if isinstance(raw, dict) else raw
    return items if isinstance(items, list) else []

def save_shift_report_item(item):
    items = [row for row in shift_report_items() if not (row.get('master') == item.get('master') and row.get('date') == item.get('date'))]
    items.append(item)
    write_cache(SHIFT_REPORTS_FILE, {'items': items[-1000:]})

def shift_seen_map():
    raw = read_cache(SHIFT_SEEN_FILE) or {}
    return raw if isinstance(raw, dict) else {}

def mark_shift_seen(name):
    data = shift_seen_map()
    days = [day for day in (data.get(name) or []) if isinstance(day, str)]
    today = get_msk_now().date().isoformat()
    if today not in days:
        days.append(today)
    data[name] = days[-90:]
    write_cache(SHIFT_SEEN_FILE, data)

def missing_shift_day(name):
    today = get_msk_now().date()
    reported = {row.get('date') for row in shift_report_items() if row.get('master') == name}
    missing = []
    for day in shift_seen_map().get(name) or []:
        try:
            parsed = datetime.strptime(day, '%Y-%m-%d').date()
        except ValueError:
            continue
        if parsed < today and day not in reported:
            missing.append(parsed)
    return min(missing) if missing else None

def history_rows(master):
    raw = read_cache(f'history_{master}.json') or {}
    if isinstance(raw, dict):
        raw = raw.get('items') or raw.get('history') or []
    return raw if isinstance(raw, list) else []

def history_stamp(item):
    text = str((item or {}).get('timestamp') or '')[:19]
    try:
        return datetime.strptime(text, '%Y-%m-%d %H:%M:%S')
    except ValueError:
        return None

def closed_on_day(master, day):
    count = 0
    for item in history_rows(master):
        stamp = history_stamp(item)
        if stamp and stamp.date().isoformat() == day:
            count += 1
    return count

def master_name_list():
    names = []
    for item in MASTER_CREDENTIALS.values():
        if item.get('role') == 'master' and item.get('name') not in names:
            names.append(item['name'])
    return names

def history_kind(item):
    action = (item or {}).get('action') or {}
    if isinstance(action, dict):
        code = str(action.get('action') or '')
        label = str(action.get('label') or '').lower()
    else:
        code = ''
        label = str(action).lower()
    if code == 'fail' or 'не было' in label or 'отсутств' in label:
        return 'missing'
    if code == 'evacuation' or 'эвакуа' in label:
        return 'evac'
    if code in ('evacuation_and_replace', 'transit_replace', 'replace_yes') or 'замен' in label:
        return 'replace'
    if code == 'done' or 'выполн' in label:
        return 'done'
    return 'other'

def stats_bundle(only_name=None):
    now = get_msk_now().replace(tzinfo=None)
    week = now - timedelta(days=7)
    month = now - timedelta(days=30)
    try:
        averages = db.close_hours(60)
    except Exception:
        averages = {}
    people = []
    names = [only_name] if only_name else master_name_list()
    local_hours = {}
    for name in names:
        week_count = month_count = evac_count = missing_count = replace_count = 0
        hours = []
        for item in history_rows(name):
            stamp = history_stamp(item)
            if not stamp or stamp < month:
                continue
            month_count += 1
            if stamp >= week:
                week_count += 1
            kind = history_kind(item)
            if kind == 'evac':
                evac_count += 1
            elif kind == 'missing':
                missing_count += 1
            elif kind == 'replace':
                replace_count += 1
            if item.get('hours') not in (None, ''):
                try:
                    value = float(item.get('hours'))
                except (TypeError, ValueError):
                    value = None
                if value is not None and 0 <= value <= 24 * 21:
                    hours.append(value)
        if hours:
            local_hours[name] = round(sum(hours) / len(hours), 1)
        open_count = 0
        try:
            _, tickets, _ = master_groups(name)
            open_count = sum(1 for ticket in tickets if is_active_status(ticket.get('status')))
        except Exception:
            open_count = 0
        avg = averages.get(name)
        if avg is None:
            avg = local_hours.get(name)
        people.append({
            'name': name,
            'open': open_count,
            'week': week_count,
            'month': month_count,
            'evac': evac_count,
            'missing': missing_count,
            'replace': replace_count,
            'avg_hours': avg,
            'rating': master_rating_value(name),
        })
    known = [person['avg_hours'] for person in people if person.get('avg_hours') is not None]
    if not known:
        known = list(averages.values()) or list(local_hours.values())
    overall = round(sum(known) / len(known), 1) if known else None
    return {'people': people, 'overall_hours': overall, 'reports': shift_report_items()}

def append_shift_sheet(item):
    book = get_sheet_client()
    title = 'Ежедневный отчет'
    try:
        sheet = book.worksheet(title)
    except Exception:
        sheet = book.add_worksheet(title, rows=2000, cols=10)
        sheet.update(
            values=[['Дата', 'Мастер', 'Начало', 'Конец', 'Пробег', 'Доп. расходы', 'Статус', 'Комментарий', 'Закрыто заявок', 'Записано']],
            range_name='A1:J1',
        )
    row = [[
        item.get('date_ru') or item.get('date') or '',
        item.get('master') or '',
        item.get('start') or '',
        item.get('end') or '',
        item.get('mileage') or '',
        item.get('expenses') or '',
        item.get('status_label') or '',
        item.get('comment') or '',
        item.get('closed_count') or 0,
        item.get('saved_at') or '',
    ]]
    column = sheet.col_values(1)
    start = len(column) + 1
    for index, value in enumerate(column, start=1):
        if index > 1 and not str(value).strip():
            start = index
            break
    sheet.update(values=row, range_name=f'A{start}:J{start}', value_input_option='USER_ENTERED')

def guard_shift(name):
    missing = missing_shift_day(name)
    if missing:
        return redirect(url_for('master_report', name=name, day=missing.isoformat()))
    mark_shift_seen(name)
    return None

TRUNK_FILE = os.path.join(CACHE_DIR, "trunks.json")
PARTS_QUEUE_FILE = os.path.join(CACHE_DIR, "parts_queue.json")
_trunk_lock = threading.Lock()

def _read_json(path, fallback):
    try:
        if os.path.exists(path):
            with open(path, 'r', encoding='utf-8') as f:
                data = json.load(f)
                return data
    except Exception:
        pass
    return fallback

def _write_json(path, data):
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def load_trunks():
    data = _read_json(TRUNK_FILE, {})
    return data if isinstance(data, dict) else {}

def save_trunks(data):
    _write_json(TRUNK_FILE, data)
    db.mirror_trunks(data)

def trunk_rows(master):
    bag = load_trunks().get(master) or {}
    rows = []
    total_qty = 0
    for name, price in PARTS:
        qty = int(bag.get(name) or 0)
        rows.append({'name': name, 'qty': qty})
        total_qty += qty
    return rows, total_qty

def trunk_qty(master, name):
    return int((load_trunks().get(master) or {}).get(name) or 0)

def change_trunk(master, name, delta):
    if name not in PARTS_SET:
        return False, 0
    with _trunk_lock:
        data = load_trunks()
        bag = data.get(master) or {}
        current = int(bag.get(name) or 0)
        new_qty = current + int(delta)
        if new_qty < 0:
            return False, current
        bag[name] = new_qty
        data[master] = bag
        save_trunks(data)
    return True, new_qty

def queue_part_move(master, action, name, qty, uid='', gos='', darks=''):
    price = PARTS_PRICE.get(name, 0)
    move = {
        'time': get_msk_now().strftime('%Y-%m-%d %H:%M:%S'),
        'master': master,
        'action': action,
        'name': name,
        'qty': int(qty),
        'price': price,
        'sum': price * int(qty),
        'uid': uid or '',
        'gos': gos or '',
        'darks': str(darks or ''),
    }
    with _trunk_lock:
        data = _read_json(PARTS_QUEUE_FILE, {'moves': []})
        if not isinstance(data, dict):
            data = {'moves': []}
        data.setdefault('moves', []).append(move)
        data['log'] = (data.get('log') or [])[-1000:]
        data['log'].append(move)
        _write_json(PARTS_QUEUE_FILE, data)
    db.mirror_move(move)

def parts_log():
    data = _read_json(PARTS_QUEUE_FILE, {})
    log = data.get('log') if isinstance(data, dict) else []
    return log if isinstance(log, list) else []

def flush_parts_queue():
    with _trunk_lock:
        data = _read_json(PARTS_QUEUE_FILE, {'moves': [], 'log': []})
        moves = list(data.get('moves') or []) if isinstance(data, dict) else []
    if not moves:
        return 0
    sheet = get_sheet_client()
    try:
        ws = sheet.worksheet('Движение запчастей 2')
    except Exception:
        ws = sheet.add_worksheet('Движение запчастей 2', 2000, 10)
        ws.update('A1:J1', [['Дата', 'Мастер', 'Действие', 'Запчасть', 'Количество', 'Цена', 'Сумма', 'Заявка', 'Госномер', 'Даркстор']], value_input_option='RAW')
    values = [[m['time'], m['master'], m['action'], m['name'], m['qty'], m.get('price', 0), m.get('sum', 0), m.get('uid', ''), m.get('gos', ''), m.get('darks', '')] for m in moves]
    ws.append_rows(values, value_input_option='RAW')
    write_trunk_sheet(sheet)
    with _trunk_lock:
        data = _read_json(PARTS_QUEUE_FILE, {'moves': [], 'log': []})
        left = data.get('moves') or []
        data['moves'] = left[len(moves):] if len(left) >= len(moves) else []
        _write_json(PARTS_QUEUE_FILE, data)
    logger.info(f"✅ Движение запчастей записано: {len(moves)}")
    return len(moves)

def write_trunk_sheet(sheet):
    try:
        ws = sheet.worksheet('Багажник')
    except Exception:
        ws = sheet.add_worksheet('Багажник', 400, 5)
    rows = [['Мастер', 'Запчасть', 'Количество', 'Цена', 'Сумма']]
    stored = load_trunks()
    for master in MASTERS:
        bag = stored.get(master) or {}
        for name, price in PARTS:
            qty = int(bag.get(name) or 0)
            if qty <= 0:
                continue
            rows.append([master, name, qty, price, qty * price])
    ws.clear()
    ws.update('A1', rows, value_input_option='RAW')

def spend_parts(master, items, uid='', gos='', darks=''):
    clean = []
    for item in items:
        name = (item.get('name') or '').strip()
        qty = int(item.get('qty') or 0)
        if qty <= 0 or name not in PARTS_SET:
            continue
        clean.append((name, qty))
    if not clean:
        return True, ''
    with _trunk_lock:
        data = load_trunks()
        bag = dict(data.get(master) or {})
        for name, qty in clean:
            have = int(bag.get(name) or 0)
            bag[name] = max(0, have - qty)
        data[master] = bag
        save_trunks(data)
    lines = []
    for name, qty in clean:
        queue_part_move(master, 'Поставил', name, qty, uid, gos, darks)
        lines.append(f'{name} x{qty}')
    return True, ', '.join(lines)

def flush_queue():
    if not flush_lock.acquire(blocking=False):
        return 0
    try:
        with queue_lock:
            queue_data = _load_queue()
            tasks = list(queue_data.get('tasks') or [])
            if not tasks:
                return 0
            for i, task in enumerate(tasks):
                if not task.get('qid'):
                    task['qid'] = f"legacy-{i}-{task.get('uid')}"
            _save_queue(queue_data)
            taken = [dict(task) for task in tasks]
            taken_ids = {task['qid'] for task in taken}
        skipped = apply_queue_to_sheets(taken) or set()
        report_tasks = []
        for task in taken:
            if task.get('qid') in skipped:
                continue
            data = task.get('data') or {}
            if data.get('skip_report') or task.get('type') == 'status_update':
                continue
            report_tasks.append(task)
        if report_tasks:
            write_to_report(report_tasks)
            for task in report_tasks:
                data = dict(task.get('data') or {})
                data['uid'] = task.get('uid')
                append_master_history(data.get('master'), task.get('type'), data)
        with queue_lock:
            queue_data = _load_queue()
            queue_data['tasks'] = [task for task in (queue_data.get('tasks') or []) if task.get('qid') not in taken_ids or task.get('qid') in skipped]
            queue_data['last_sync'] = get_msk_now().strftime('%Y-%m-%d %H:%M:%S')
            _save_queue(queue_data)
        logger.info(f"✅ Очередь записана и очищена: {len(taken)}")
        return len(taken)
    finally:
        try:
            flush_parts_queue()
        except Exception as e:
            logger.error(f"Ошибка записи запчастей: {e}")
        flush_lock.release()


def clear_queue():
    write_queue({'tasks': [], 'last_sync': get_msk_now().strftime('%Y-%m-%d %H:%M:%S')})

# ============================================================
# ЗАПИСЬ В "ОТЧЕТ МАСТЕРА"
# ============================================================
def write_to_report(tasks):
    try:
        sheet_client = get_sheet_client()
        now = get_msk_now().strftime('%Y-%m-%d %H:%M:%S')
        
        try:
            report_sheet = sheet_client.worksheet("Отчет мастера")
        except:
            report_sheet = sheet_client.add_worksheet("Отчет мастера", 100, 20)
            headers = ['Дата выполнения', 'Мастер', 'ID заявки', 'Госномер', 'Описание', 
                      'Тип техники', 'Количество', 'Статус', 'Запчасти', 'Комментарий', 
                      'Номер даркстора', 'Время создания заявки', 'Статус обработки']
            for i, h in enumerate(headers, start=1):
                report_sheet.update_cell(1, i, h)
        
        current_rows = report_sheet.get_all_values()
        start_row = len(current_rows) + 1
        
        updates_report = []
        
        for idx, task in enumerate(tasks):
            uid = task.get('uid')
            task_type = task.get('type')
            data = task.get('data', {})
            master = data.get('master', '')
            darks_number = data.get('darks_number', '')
            parts = data.get('parts', '')
            reason = data.get('reason', '')
            extra = data.get('extra', '')
            
            ticket = None
            if uid in uid_index:
                ticket = uid_index[uid]['ticket']
            if not ticket:
                ticket = find_cached_ticket(uid) or {}
            if not ticket:
                logger.warning(f"Заявка {uid} не найдена в кэше, в отчёт пишем то, что пришло от мастера")
            
            status_map = {
                'done': '✅ Выполнено',
                'fail': '⏹️ Обработано',
                'evacuation': '🔧 Эвакуация',
                'replace_yes': '✅ Выполнено',
                'taken_no_replace': '🔵 Доделать',
                'replace_no': '🔵 Доделать',
                'transit_replace': '✅ Выполнено'
            }
            status = status_map.get(task_type, '✅ Выполнено')
            
            if ticket.get('type') in ['Аккумуляторная батарея', 'Зарядное устройство']:
                quantity = parts or extra or '1'
            else:
                quantity = '-'
            
            row_idx = start_row + idx
            
            updates_report.append({'range': f'A{row_idx}', 'values': [[now]]})
            updates_report.append({'range': f'B{row_idx}', 'values': [[master]]})
            updates_report.append({'range': f'C{row_idx}', 'values': [[uid]]})
            updates_report.append({'range': f'D{row_idx}', 'values': [[ticket.get('gos', '')]]})
            updates_report.append({'range': f'E{row_idx}', 'values': [[ticket.get('desc', '')]]})
            updates_report.append({'range': f'F{row_idx}', 'values': [[ticket.get('type', '')]]})
            updates_report.append({'range': f'G{row_idx}', 'values': [[quantity]]})
            updates_report.append({'range': f'H{row_idx}', 'values': [[status]]})
            updates_report.append({'range': f'I{row_idx}', 'values': [[parts or extra or '-']]})
            updates_report.append({'range': f'J{row_idx}', 'values': [[reason or '-']]})
            updates_report.append({'range': f'K{row_idx}', 'values': [[darks_number or ticket.get('darks', '')]]})
            updates_report.append({'range': f'L{row_idx}', 'values': [[ticket.get('created', '')]]})
            updates_report.append({'range': f'M{row_idx}', 'values': [['Новый']]})
        
        if updates_report:
            report_sheet.batch_update(updates_report)
            logger.info(f"✅ Записано {len(updates_report)//13} записей в Отчет мастера")
        
    except Exception as e:
        logger.error(f"❌ Ошибка записи в Отчет мастера: {e}")
        raise

# ============================================================
# ФОНОВЫЙ ПРОЦЕСС ОБРАБОТКИ ОЧЕРЕДИ
# ============================================================
def process_queue_background():
    while True:
        try:
            time.sleep(10)
            drop_retired_masters()
            flush_queue()
            if time.time() - _app_started_at < 30:
                continue
            if time.time() - _last_google_import >= 180:
                import_new_tickets_from_google()
        except Exception as e:
            logger.error(f"❌ Ошибка в фоновом процессе: {e}")

_queue_thread_started = False

def ensure_queue_thread():
    global _queue_thread_started
    if _queue_thread_started:
        return
    _queue_thread_started = True
    threading.Thread(target=process_queue_background, daemon=True).start()

# ============================================================
# ОТПРАВКА УВЕДОМЛЕНИЙ
# ============================================================
def tg_plain(value):
    return html.escape(str(value or ''), quote=False)

def ticket_mark(ticket):
    kind = str((ticket or {}).get('type') or '').lower()
    if 'аккумулятор' in kind:
        return '🔋'
    if 'заряд' in kind:
        return '⚡'
    return '🚲'

def curator_notice_groups():
    tickets = []
    cache = get_admin_cache() or {}
    for dir_data in cache.get('directions', {}).values():
        tickets.extend(dir_data.get('tickets') or [])
    if not tickets:
        tickets = get_tickets_from_sheets()
    groups = {}
    for ticket in tickets:
        if not ticket.get('master') or not is_active_status(ticket.get('status')):
            continue
        groups.setdefault(str(ticket.get('darks') or ''), []).append(ticket)
    payload = []
    for dark, items in groups.items():
        if not dark:
            continue
        address = (items[0].get('address') or '').strip()
        names = []
        for ticket in items:
            master = (ticket.get('master') or '').strip()
            if master and master not in names:
                names.append(master)
        master_title = 'Вам назначен мастер' if len(names) < 2 else 'Вам назначены мастера'
        lines = [
            'Привет!',
            f'{master_title}: {", ".join(names)}',
            '',
            f'<b>{tg_plain(address)} (даркстор {tg_plain(dark)})</b>',
            f'Заявки ({len(items)}):',
        ]
        for ticket in items:
            if ticket.get('type') in ('Аккумуляторная батарея', 'Зарядное устройство'):
                name = ticket.get('type')
            else:
                name = ticket.get('gos') or 'Без номера'
            mark = ' — ЗАМЕНА' if is_evacuation_ticket(ticket) else ''
            lines.append(f'{ticket_mark(ticket)} {tg_plain(name)} | {tg_plain(ticket.get("desc") or "-")}{mark}')
        lines.append('')
        lines.append('Подготовьте, пожалуйста, технику к ремонту')
        lines.append('Хорошего дня!')
        payload.append({'darks': dark, 'text': '\n'.join(lines), 'count': len(items)})
    return payload

def curator_sent_map():
    data = read_cache('curator_sent.json') or {}
    data.pop('updated_at', None)
    return data

def remember_curator_sent(rows):
    current = curator_sent_map()
    now = get_msk_now().strftime('%d.%m %H:%M')
    for row in rows:
        current[str(row.get('darks') or '')] = {
            'at': now,
            'sent': int(row.get('sent') or 0),
            'tickets': int(row.get('tickets') or 0),
        }
    write_cache('curator_sent.json', current)
    return curator_sent_map()

def notify_curators(message):
    try:
        groups = curator_notice_groups()
        if not groups:
            return {'success': False, 'error': 'Нет назначенных заявок'}
        response = requests.post(
            f"{BOT_API_URL}/send_notification",
            json={"groups": groups},
            timeout=20,
        )
        if response.status_code != 200:
            return {'success': False, 'error': 'Бот не принял сообщение'}
        body = {}
        try:
            body = response.json() or {}
        except Exception:
            body = {}
        by_dark = {str(item.get('darks') or ''): int(item.get('sent') or 0) for item in (body.get('results') or [])}
        rows = []
        for group in groups:
            dark = str(group.get('darks') or '')
            sent = by_dark.get(dark, 1 if not body.get('results') else 0)
            rows.append({'darks': dark, 'sent': sent, 'tickets': int(group.get('count') or 0)})
        marks = remember_curator_sent(rows)
        return {'success': True, 'marks': marks, 'count': body.get('count', 0)}
    except Exception as e:
        logger.warning(f'Не отправились сообщения кураторам: {e}')
        return {'success': False, 'error': 'Не отправилось'}

def login_required(f):
    # проверка входа объявлена до маршрутов, иначе сайт не запускается
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('authenticated'):
            if request.path.startswith('/api/'):
                return jsonify({'success': False, 'error': 'Сессия истекла. Обновите страницу и войдите снова.'}), 401
            return redirect(url_for('login_page'))
        return f(*args, **kwargs)
    return decorated_function

def curator_chat_ids():
    found = []
    seen = set()
    for info in (load_darks_reference() or {}).values():
        raw = info.get('curator_ids') or ''
        for part in re.split(r'[,;\s]+', str(raw)):
            token = part.strip()
            if token.isdigit() and token not in seen:
                seen.add(token)
                found.append(token)
    return found

def send_via_master_notify(chat_id, text):
    response = requests.post(
        f'{BOT_API_URL}/notify_master',
        json={'chat_id': chat_id, 'text': text},
        timeout=20,
    )
    if response.status_code != 200:
        return False
    try:
        body = response.json()
    except Exception:
        return False
    return bool(body.get('success'))

@app.route('/api/broadcast_curators', methods=['POST'])
@login_required
def api_broadcast_curators():
    if session.get('role') != 'admin':
        return jsonify({'success': False, 'error': 'Только для руководителя'}), 403
    text = ((request.json or {}).get('text') or '').strip()
    if not text:
        return jsonify({'success': False, 'error': 'Напишите текст'})
    try:
        response = requests.post(
            f'{BOT_API_URL}/broadcast',
            json={'text': text},
            timeout=40,
        )
        body = {}
        try:
            body = response.json() if response.content else {}
        except Exception:
            body = {}
        if response.status_code == 200 and body.get('success'):
            return jsonify({'success': True, 'count': body.get('count', 0)})
    except Exception as exc:
        logger.warning(f'Общая рассылка бота не ответила: {exc}')
    ids = curator_chat_ids()
    if not ids:
        return jsonify({'success': False, 'error': 'В столбце G листа Дарксторы нет Telegram ID кураторов'})
    sent = 0
    for chat_id in ids:
        try:
            if send_via_master_notify(chat_id, text):
                sent += 1
        except Exception as exc:
            logger.warning(f'Не ушло куратору {chat_id}: {exc}')
    if not sent:
        return jsonify({'success': False, 'error': 'Сообщение никому не ушло'})
    return jsonify({'success': True, 'count': sent})

def generate_curator_message(tickets_data):
    now = get_msk_now().strftime('%d.%m.%Y %H:%M')
    message = f"📢 Привет, на связи Vanta Bikes! ({now})\n\n"
    groups = {}
    for t in tickets_data:
        darks = t.get('darks', 'без номера')
        if darks not in groups:
            groups[darks] = {
                'address': t.get('address', 'Адрес не указан'),
                'tickets': []
            }
        groups[darks]['tickets'].append(t)

    for darks, group in groups.items():
        message += f"📍 {group['address']} (даркстор {darks})\n"
        message += f"📋 Заявки ({len(group['tickets'])}):\n"
        for t in group['tickets']:
            if t.get('type') in ['Аккумуляторная батарея', 'Зарядное устройство']:
                identifier = t.get('type', 'Без номера')
            else:
                identifier = t.get('gos', 'Без номера')
            message += f"   {identifier} | {t.get('desc', '-')}\n"
        message += "\n"

    message += "Подготовьте, пожалуйста, технику к ремонту\n"
    message += "Хорошего дня! 🙌"
    return message

# ============================================================
# АВТОРИЗАЦИЯ
# ============================================================
@app.route('/api/bike_location')
@login_required
def api_bike_location():
    if session.get('role') != 'admin':
        return jsonify({'success': False, 'error': 'Только для руководителя'}), 403
    iot = resolve_ticket_iot(request.args.get('iot', ''), request.args.get('darks', ''), request.args.get('gos', ''))
    if not normalize_iot(iot):
        return jsonify({'success': False, 'error': 'У этого номера нет IoT в базе велосипедов'})
    point = lookup_bike_point(iot)
    if not point or point.get('lat') is None or point.get('lon') is None:
        return jsonify({'success': False, 'error': _xdevice_error or 'Точки для этого IoT нет'})
    lat = point['lat']
    lon = point['lon']
    when = ""
    if point.get('at'):
        try:
            when = datetime.fromtimestamp(point['at'] / 1000, MSK).strftime('%d.%m %H:%M')
        except Exception:
            when = ""
    return jsonify({
        'success': True,
        'lat': lat,
        'lon': lon,
        'online': point['online'],
        'at': when,
        'url': f'https://yandex.ru/maps/?ll={lon},{lat}&pt={lon},{lat}&z=16&l=map',
    })

def bot_hook_ok():
    return request.headers.get('X-Bot-Key') == BOT_HOOK_KEY

def dispute_actor(note):
    for line in reversed((note or '').splitlines()):
        if ',' not in line:
            continue
        name = line.rsplit(',', 1)[-1].strip()
        if name in MASTERS:
            return name
    return ''

@app.route('/api/bot_map')
def api_bot_map():
    if not bot_hook_ok():
        return jsonify({'success': False, 'error': 'Нет доступа'}), 403
    iot = resolve_ticket_iot(
        request.args.get('iot', ''),
        request.args.get('darks', ''),
        request.args.get('gos', ''),
    )
    if not normalize_iot(iot):
        return jsonify({'success': False, 'error': 'Нет IoT'})
    point = lookup_bike_point(iot)
    if not point or point.get('lat') is None or point.get('lon') is None:
        return jsonify({'success': False, 'error': 'Нет точки'})
    lat = point['lat']
    lon = point['lon']
    when = ''
    if point.get('at'):
        try:
            when = datetime.fromtimestamp(point['at'] / 1000, MSK).strftime('%d.%m %H:%M')
        except Exception:
            when = ''
    return jsonify({
        'success': True,
        'at': when,
        'url': f'https://yandex.ru/maps/?ll={lon},{lat}&pt={lon},{lat}&z=16&l=map',
    })

def bike_map_url(iot, darks, gos):
    iot = resolve_ticket_iot(iot, darks, gos)
    if not normalize_iot(iot):
        return ''
    ensure_xdevice_loop()
    refresh_xdevice_points(force=False)
    point = None
    for token in iot_tokens(iot):
        point = _xdevice_points.get(token)
        if point:
            break
    if not point or point.get('lat') is None or point.get('lon') is None:
        return ''
    lat = point['lat']
    lon = point['lon']
    return f'https://yandex.ru/maps/?ll={lon},{lat}&pt={lon},{lat}&z=16&l=map'

@app.route('/api/bot_maps', methods=['POST'])
def api_bot_maps():
    if not bot_hook_ok():
        return jsonify({'success': False, 'error': 'Нет доступа'}), 403
    bikes = (request.json or {}).get('bikes') or []
    maps = []
    for item in bikes[:40]:
        if not isinstance(item, dict):
            continue
        maps.append({
            'gos': item.get('gos') or '',
            'url': bike_map_url(item.get('iot') or '', item.get('darks') or '', item.get('gos') or ''),
        })
    return jsonify({'success': True, 'maps': maps})

@app.route('/api/curator_dispute', methods=['POST'])
def api_curator_dispute():
    if not bot_hook_ok():
        return jsonify({'success': False, 'error': 'Нет доступа'}), 403
    data = request.json or {}
    uid = (data.get('uid') or '').strip()
    text = (data.get('text') or '').strip()
    event = (data.get('event') or '').strip()
    darks = str(data.get('darks') or '').strip()
    if not uid or len(text) < 3:
        return jsonify({'success': False, 'error': 'Нет текста'})
    ticket = find_cached_ticket(uid)
    if not ticket:
        return jsonify({'success': False, 'error': 'Заявка не найдена'}), 404
    if darks and str(ticket.get('darks') or '') != darks:
        return jsonify({'success': False, 'error': 'Чужая точка'}), 403
    moment = get_msk_now().strftime('%d.%m %H:%M')
    previous = (ticket.get('note') or '').strip()
    repeat = 'Оспорено куратором' in previous or 'Повторный спор' in previous
    replaced = event == 'replaced'
    label = ticket_label(ticket)
    source = ticket.get('source') or source_of(uid)
    if source not in ('Заявки', 'Заявки бот', 'Импорт М4'):
        source = 'Заявки'
    row_index = ticket.get('row_index')
    if repeat or replaced:
        line = f"Повторный спор {moment}: {text}" if repeat else f"Оспорено куратором {moment}: {text}"
        note = f'{previous}\n{line}'.strip() if previous else line
        kept = {
            'done': '✅ Выполнено',
            'processed': '⏹️ Обработано',
            'evacuation': '🔧 Эвакуация',
            'todo': '🔵 Доделать',
        }.get(ticket.get('status'), '✅ Выполнено')
        update_ticket_in_admin_cache(uid, ticket.get('status') or 'done', note)
        add_to_queue({
            'uid': uid,
            'source': source,
            'row_index': row_index,
            'type': 'status_update',
            'data': {
                'status': kept,
                'new_status': ticket.get('status') or 'done',
                'note': note,
                'skip_report': True,
            },
        })
        return jsonify({'success': True, 'reopened': False})
    line = f"Оспорено куратором {moment}: {text}"
    note = f'{previous}\n{line}'.strip() if previous else line
    master = (ticket.get('master') or '').strip() or dispute_actor(previous)
    update_ticket_in_admin_cache(uid, 'pending', note, extra={
        'master': '',
        'evac_reason': '',
        'is_done': False,
        'is_active': True,
    })
    remember_master_cleared(uid)
    if master:
        remove_ticket_from_master_cache(master, uid)
    refresh_master_caches_from_admin()
    add_to_queue({
        'uid': uid,
        'source': source,
        'row_index': row_index,
        'type': 'status_update',
        'data': {
            'status': '🟡 В работе',
            'new_status': 'pending',
            'note': note,
            'skip_report': True,
            'master': '',
            'write_master': True,
        },
    })
    if master:
        for chat_id in master_chat_ids(master):
            try:
                send_via_master_notify(chat_id, f'Куратор не принял {label}. Заявка снята с вас: {text}')
            except Exception as exc:
                logger.warning(f'Не ушло мастеру о споре: {exc}')
    return jsonify({'success': True, 'reopened': True})

def send_bike_command(object_id, command):
    token = xdevice_token()
    if not token:
        return False, "В настройках сервера нет ключа XDEVICE_TOKEN"
    url = "https://app.xdevice.ru/api/v1/objects/" + str(object_id) + "/commands/" + command
    try:
        response = requests.post(
            url,
            json={},
            headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
            timeout=25,
        )
    except Exception as exc:
        logger.warning(f"Команда вело не ушла: {exc}")
        return False, "Сервер IoT не ответил"
    if response.status_code == 400 and "error_offline" in response.text:
        return False, "Велосипед не на связи. Команда не дошла."
    if response.status_code >= 400:
        logger.warning(f"Команда вело {command} для {object_id}: {response.status_code} {response.text[:180]}")
        return False, "Команда не принята"
    return True, ""

@app.route('/api/bike_power', methods=['POST'])
@login_required
def api_bike_power():
    if session.get('role') != 'admin':
        return jsonify({'success': False, 'error': 'Только для руководителя'}), 403
    data = request.json or {}
    action = (data.get('action') or '').strip()
    if action not in ('on', 'off'):
        return jsonify({'success': False, 'error': 'Нужно включить или выключить'})
    iot = resolve_ticket_iot(data.get('iot', ''), data.get('darks', ''), data.get('gos', ''))
    if not normalize_iot(iot):
        return jsonify({'success': False, 'error': 'У этого номера нет IoT в базе велосипедов'})
    point = lookup_bike_point(iot)
    object_id = (point or {}).get('object_id') or ''
    if not object_id:
        refresh_xdevice_points(force=True)
        point = lookup_bike_point(iot)
        object_id = (point or {}).get('object_id') or ''
    if not object_id:
        return jsonify({'success': False, 'error': 'Модуль в списке IoT не сопоставился'})
    command = 'unlock' if action == 'on' else 'lock'
    ok, error = send_bike_command(object_id, command)
    if not ok:
        return jsonify({'success': False, 'error': error})
    if point is not None:
        point['lock'] = action == 'off'
    return jsonify({'success': True, 'lock': action == 'off', 'message': 'Велосипед включён' if action == 'on' else 'Велосипед выключен'})

def publish_new_ticket(ticket):
    cache = get_admin_cache() or {}
    directions = cache.get('directions') or {}
    direction = ticket.get('direction') or 'Без направления'
    bucket = directions.get(direction) or {'tickets': [], 'active_count': 0}
    bucket.setdefault('tickets', []).append(ticket)
    bucket['active_count'] = sum(1 for item in bucket['tickets'] if is_active_status(item.get('status')))
    directions[direction] = bucket
    order = ['Напр 1', 'Напр 2', 'Напр 3', 'Напр 4', 'Без направления']
    sorted_directions = {}
    for key in order:
        if key in directions:
            sorted_directions[key] = directions[key]
    for key, value in directions.items():
        if key not in sorted_directions:
            sorted_directions[key] = value
    write_cache('admin_cache.json', {'directions': sorted_directions})
    if ticket.get('uid'):
        uid_index[ticket['uid']] = {
            'source': ticket.get('source'),
            'row_index': ticket.get('row_index'),
            'ticket': ticket,
        }
    _tickets_mem.append(ticket)
    try:
        db.mirror_ticket(ticket)
    except Exception as exc:
        logger.warning(f'Заявка в базу не ушла: {exc}')

def append_admin_ticket_row(row):
    worksheet = get_sheet_client().worksheet('Заявки бот')
    first = worksheet.acell('A1').value or ''
    if str(first).lstrip().startswith('='):
        raise RuntimeError('Лист Заявки бот начинается с формулы')
    start = len(worksheet.col_values(1)) + 1
    padded = (list(row) + [''] * 15)[:15]
    worksheet.batch_update([{'range': f'A{start}:O{start}', 'values': [padded]}], value_input_option='RAW')
    return start

@app.route('/api/bike_search')
@login_required
def api_bike_search():
    if session.get('role') != 'admin':
        return jsonify({'success': False, 'error': 'Только для руководителя'}), 403
    query = (request.args.get('q') or '').strip()
    if len(query) < 3:
        return jsonify({'success': False, 'error': 'Введите хотя бы 3 символа'})
    needle_gos = normalize_gos(query)
    needle_iot = normalize_iot(query).lower()
    needle_serial = re.sub(r'\s+', '', query).lower()
    try:
        rows = load_bike_rows()
    except Exception:
        return jsonify({'success': False, 'error': 'База велосипедов сейчас не открылась'})
    if not darks_ref:
        try:
            load_darks_reference()
        except Exception:
            pass
    header = rows[0] if rows else []
    kind_i = sheet_col(header, ['тип вело'], 5)
    company_i = sheet_col(header, ['компания', 'поставщик'], 8)
    hits = []
    for row in rows[1:]:
        serial = (row[0] if row else '').strip()
        gos = (row[1] if len(row) > 1 else '').strip()
        iot = (row[2] if len(row) > 2 else '').strip()
        darks = (row[4] if len(row) > 4 else '').strip()
        if not serial and not gos and not iot:
            continue
        exact = (
            (needle_gos and normalize_gos(gos) == needle_gos)
            or (needle_iot and normalize_iot(iot).lower() == needle_iot)
            or (needle_serial and re.sub(r'\s+', '', serial).lower() == needle_serial)
        )
        if not exact:
            continue
        kind = row[kind_i].strip() if len(row) > kind_i else ''
        company = row[company_i].strip() if len(row) > company_i else ''
        point = lookup_bike_point(iot) if normalize_iot(iot) else None
        if not point and gos:
            point = lookup_bike_point(gos)
        if not point and serial:
            point = lookup_bike_point(serial)
        url = ''
        online = False
        map_error = ''
        if not point and not _xdevice_points:
            refresh_xdevice_points(force=True)
            point = lookup_bike_point(iot) if normalize_iot(iot) else None
            if not point and gos:
                point = lookup_bike_point(gos)
        if not normalize_iot(iot) and not point:
            map_error = 'В базе нет номера IoT, точку не к чему привязать'
        elif not point:
            map_error = _xdevice_error or 'Точки на карте сейчас нет'
        if point and point.get('lat') is not None and point.get('lon') is not None:
            lat = point['lat']
            lon = point['lon']
            url = f'https://yandex.ru/maps/?ll={lon},{lat}&pt={lon},{lat}&z=16&l=map'
            online = bool(point.get('online'))
            map_error = ''
        elif point and not map_error:
            map_error = 'У поставщика нет координат этого модуля'
        info = (darks_ref or {}).get(str(darks)) or {}
        hits.append({
            'serial': serial,
            'gos': gos,
            'iot': iot,
            'darks': darks,
            'address': info.get('address') or '',
            'kind': kind,
            'supplier': company,
            'url': url,
            'online': online,
            'map_error': map_error,
        })
        if len(hits) >= 20:
            break
    return jsonify({'success': True, 'items': hits})

@app.route('/api/admin_ticket', methods=['POST'])
@login_required
def api_admin_ticket():
    if session.get('role') != 'admin':
        return jsonify({'success': False, 'error': 'Только для руководителя'}), 403
    data = request.json or {}
    darks = str(data.get('darks') or '').strip()
    bike_type = str(data.get('type') or '').strip()
    custom = str(data.get('custom_type') or '').strip()
    desc = str(data.get('desc') or '').strip()
    gos = str(data.get('gos') or '').strip().upper()
    qty = str(data.get('qty') or '').strip()
    if bike_type == 'Свой вариант':
        bike_type = custom
    if not darks:
        return jsonify({'success': False, 'error': 'Укажите номер даркстора'})
    if not bike_type:
        return jsonify({'success': False, 'error': 'Укажите тип техники'})
    if not desc:
        return jsonify({'success': False, 'error': 'Опишите проблему'})
    supply = bike_type in ('Аккумуляторная батарея', 'Зарядное устройство')
    if supply:
        if not qty.isdigit() or int(qty) < 1:
            return jsonify({'success': False, 'error': 'Укажите количество'})
        gos = ''
        desc = f'{int(qty)} шт. {desc}'.strip()
    elif bike_type == 'Электровелосипед' and not gos:
        return jsonify({'success': False, 'error': 'Укажите госномер'})
    if not darks_ref:
        try:
            load_darks_reference()
        except Exception:
            pass
    info = (darks_ref or {}).get(darks) or {}
    if darks_ref and not info:
        return jsonify({'success': False, 'error': 'Такого даркстора нет в списке'})
    created = get_msk_now().strftime('%Y-%m-%d %H:%M:%S')
    uid = generate_ticket_id(created)
    direction = info.get('direction') or 'Без направления'
    subtype = ''
    if bike_type == 'Электровелосипед' and gos:
        subtype = (get_bike_type_by_gos_map() or {}).get(normalize_gos(gos)) or ''
    row = [''] * 15
    row[0] = darks
    row[1] = bike_type
    row[2] = desc
    row[3] = gos
    row[4] = 'Админ'
    row[5] = created
    row[7] = '🟡 В работе'
    row[8] = subtype
    row[10] = uid
    row[11] = direction
    try:
        row_index = append_admin_ticket_row(row)
    except Exception as exc:
        logger.warning(f'Заявка админа в таблицу не записалась: {exc}')
        return jsonify({'success': False, 'error': 'Не удалось записать в лист «Заявки бот». Заявка не создана.'})
    ticket = {
        'source': 'Заявки бот',
        'darks': darks,
        'type': bike_type,
        'bike_type': subtype or bike_type,
        'bike_subtype': subtype,
        'supplier': supplier_of(gos) if gos else '',
        'desc': desc,
        'display_desc': desc,
        'gos': gos,
        'contact': 'Админ',
        'created': created,
        'timer_from': '',
        'hours_since': 0,
        'master': '',
        'status': 'pending',
        'note': '',
        'evac_reason': '',
        'uid': uid,
        'row_index': row_index,
        'address': info.get('address') or '',
        'direction': direction,
        'coords': info.get('coords') or '',
        'sent_date': '',
        'is_done': False,
        'is_active': True,
        'parts': '',
        'iot': '',
    }
    try:
        attach_iot([ticket])
    except Exception:
        pass
    publish_new_ticket(ticket)
    return jsonify({'success': True, 'ticket': ticket})

DOC_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'templates', 'docs')
DOC_KINDS = {
    'postavka': ('postavka_template.docx', 'ПОСТАВКА'),
    'vyvoz': ('vyvoz_template.docx', 'ВЫВОЗ'),
    'sverka': ('sverka_template.docx', 'СВЕРКА'),
    'zamena': ('zamena_template.docx', 'ЗАМЕНА'),
}

def fill_doc_text(doc, mapping):
    def apply(paragraph):
        text = paragraph.text
        updated = text
        for key, value in mapping.items():
            updated = updated.replace(key, '' if value is None else str(value))
        if updated != text and paragraph.runs:
            paragraph.runs[0].text = updated
            for run in paragraph.runs[1:]:
                run.text = ''
    for paragraph in doc.paragraphs:
        apply(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    apply(paragraph)

def bike_catalog():
    rows = load_bike_rows()
    items = []
    for row in rows[1:]:
        frame = (row[0] if row else '').strip()
        gos = (row[1] if len(row) > 1 else '').strip()
        iot = (row[2] if len(row) > 2 else '').strip()
        if frame or gos or iot:
            items.append({'frame': frame, 'gos': gos, 'iot': iot})
    return items

def iot_tail(value):
    text = normalize_iot(value).lower().replace(' ', '')
    if '-' in text:
        text = text.split('-', 1)[1]
    text = re.sub(r'\D', '', text).lstrip('0')
    return text or '0'

def classify_bike_token(token):
    compact = re.sub(r'\s+', '', token or '')
    if re.fullmatch(r'\d{2}-\d{3,6}', compact):
        return 'iot', compact
    if re.fullmatch(r'\d{10,20}', compact):
        return 'frame', compact
    if re.fullmatch(r'\d{3,5}', compact):
        return 'iot_short', compact
    if re.search(r'[A-Za-zА-Яа-я]', compact) and re.search(r'\d', compact) and len(compact) <= 12:
        return 'gos', compact
    return '', ''

def parse_line_fields(line):
    fields = {'frame': '', 'gos': '', 'iot': '', 'short': False}
    for part in re.split(r'[\s,;]+', (line or '').strip()):
        if not part:
            continue
        kind, value = classify_bike_token(part)
        if kind == 'frame' and not fields['frame']:
            fields['frame'] = value
        elif kind == 'gos' and not fields['gos']:
            fields['gos'] = value
        elif kind == 'iot' and not fields['iot']:
            fields['iot'] = value
            fields['short'] = False
        elif kind == 'iot_short' and not fields['iot']:
            fields['iot'] = '25-' + value
            fields['short'] = True
    return fields

def catalog_by_iot(items, iot):
    tail = iot_tail(iot)
    exact = []
    hits = []
    wanted = normalize_iot(iot).lower()
    for item in items:
        if not item.get('iot'):
            continue
        if normalize_iot(item['iot']).lower() == wanted:
            exact.append(item)
        elif iot_tail(item['iot']) == tail:
            hits.append(item)
    return exact or hits

def catalog_by_field(items, field, token):
    if field == 'gos':
        key = normalize_gos(token)
        return [item for item in items if key and normalize_gos(item.get('gos')) == key]
    key = re.sub(r'\s+', '', token or '').lower()
    return [item for item in items if key and re.sub(r'\s+', '', item.get('frame') or '').lower() == key]

def resolve_bike_line(line, items):
    fields = parse_line_fields(line)
    if not fields['frame'] and not fields['gos'] and not fields['iot']:
        return None, line
    match = None
    if fields['iot']:
        hits = catalog_by_iot(items, fields['iot'])
        if len(hits) > 1:
            return None, line + ' — несколько велосипедов с таким IoT'
        if len(hits) == 1:
            match = hits[0]
    if not match and fields['gos']:
        hits = catalog_by_field(items, 'gos', fields['gos'])
        if len(hits) == 1:
            match = hits[0]
    if not match and fields['frame']:
        hits = catalog_by_field(items, 'frame', fields['frame'])
        if len(hits) == 1:
            match = hits[0]
    if fields['short'] and not match and not fields['frame'] and not fields['gos']:
        return None, line + ' — такого IoT нет в базе'
    bike = {
        'frame': fields['frame'] or ((match or {}).get('frame') or ''),
        'gos': fields['gos'] or ((match or {}).get('gos') or ''),
        'iot': (match or {}).get('iot') if fields['short'] and match else (fields['iot'] or ((match or {}).get('iot') or '')),
    }
    if not bike['frame'] and not bike['gos'] and not bike['iot']:
        return None, line
    return bike, None

def parse_bike_lines(text, items):
    found = []
    missing = []
    seen = set()
    for raw in (text or '').splitlines():
        line = raw.strip()
        if not line:
            continue
        bike, error = resolve_bike_line(line, items)
        if error:
            missing.append(error)
            continue
        key = '|'.join([
            (bike.get('frame') or '').strip().lower(),
            normalize_gos(bike.get('gos')),
            (bike.get('iot') or '').strip().lower(),
        ])
        if key in seen:
            continue
        seen.add(key)
        found.append(bike)
    return found, missing

def bikes_from_request(data, list_key, lines_key, catalog):
    raw = data.get(list_key)
    if isinstance(raw, list) and raw:
        bikes = []
        seen = set()
        for item in raw:
            if not isinstance(item, dict):
                continue
            bike = {
                'frame': str(item.get('frame') or '').strip(),
                'gos': str(item.get('gos') or '').strip(),
                'iot': str(item.get('iot') or '').strip(),
            }
            if not bike['frame'] and not bike['gos'] and not bike['iot']:
                continue
            key = '|'.join([bike['frame'].lower(), normalize_gos(bike['gos']), bike['iot'].lower()])
            if key in seen:
                continue
            seen.add(key)
            bikes.append(bike)
        if bikes:
            return bikes, []
    return parse_bike_lines(data.get(lines_key) or '', catalog)

def equipment_names(company):
    if company == 'ВелоСтеп':
        return 'Электровелосипед\nВелоСтеп', 'АКБ ВелоСтеп'
    return 'Электровелосипед\nVanta Bike Monster', 'АКБ Vanta'

def find_catalog_bike(items, token):
    token = (token or '').strip()
    if len(token) < 3:
        return None
    key_gos = normalize_gos(token)
    key_iot = normalize_iot(token).lower()
    key_frame = re.sub(r'\s+', '', token).lower()
    for item in items:
        if key_gos and normalize_gos(item['gos']) == key_gos:
            return item
        if key_iot and normalize_iot(item['iot']).lower() == key_iot:
            return item
        if key_frame and re.sub(r'\s+', '', item['frame']).lower() == key_frame:
            return item
    return None

def write_bike_row(table, values):
    row = table.add_row()
    for index, value in enumerate(values):
        if index < len(row.cells):
            row.cells[index].text = '' if value is None else str(value)

def add_bike_table(table, bikes, with_state, akb_count, company):
    bike_name, akb_name = equipment_names(company)
    number = 1
    for bike in bikes:
        values = [number, bike_name, bike.get('frame') or '-', bike.get('gos') or '-', bike.get('iot') or '-']
        if with_state:
            values.append('БУ')
        write_bike_row(table, values)
        number += 1
    for _ in range(max(0, akb_count)):
        values = [number, akb_name, '-', '-', '-']
        if with_state:
            values.append('БУ')
        write_bike_row(table, values)
        number += 1

def append_supply_log(rows):
    book = get_gspread_client().open_by_key(IOT_VEHICLES_SHEET_ID)
    sheet = book.worksheet('ПОСТАВКА/ВЫВОЗ вело')
    column = sheet.col_values(1)
    start = len(column) + 1
    for index, value in enumerate(column, start=1):
        if index > 1 and not str(value).strip():
            start = index
            break
    finish = start + len(rows) - 1
    sheet.update(f'A{start}:G{finish}', rows, value_input_option='USER_ENTERED')

@app.route('/api/resolve_bikes', methods=['POST'])
@login_required
def api_resolve_bikes():
    if session.get('role') != 'admin':
        return jsonify({'success': False, 'error': 'Только для руководителя'}), 403
    data = request.json or {}
    try:
        catalog = bike_catalog()
    except Exception:
        return jsonify({'success': False, 'error': 'База велосипедов сейчас не открылась'})
    bikes, missing = parse_bike_lines(data.get('lines') or '', catalog)
    if not bikes and missing:
        return jsonify({'success': False, 'error': 'Не разобрал: ' + ', '.join(missing[:8])})
    return jsonify({'success': True, 'bikes': bikes, 'missing': missing})

@app.route('/api/make_document', methods=['POST'])
@login_required
def api_make_document():
    if session.get('role') != 'admin':
        return jsonify({'success': False, 'error': 'Только для руководителя'}), 403
    data = request.json or {}
    kind = (data.get('kind') or '').strip()
    if kind not in DOC_KINDS:
        return jsonify({'success': False, 'error': 'Выберите тип документа'})
    date = (data.get('date') or '').strip()
    darks = str(data.get('darks') or '').strip()
    company = (data.get('company') or '').strip()
    if company not in ('Vanta', 'ВелоСтеп'):
        return jsonify({'success': False, 'error': 'Выберите компанию'})
    if not date:
        return jsonify({'success': False, 'error': 'Укажите дату'})
    if not darks:
        return jsonify({'success': False, 'error': 'Укажите номер даркстора'})
    if not darks_ref:
        try:
            load_darks_reference()
        except Exception:
            pass
    info = (darks_ref or {}).get(darks) or {}
    if darks_ref and not info:
        return jsonify({'success': False, 'error': 'Такого даркстора нет в списке'})
    address = info.get('address') or ''
    try:
        catalog = bike_catalog()
    except Exception:
        return jsonify({'success': False, 'error': 'База велосипедов сейчас не открылась'})
    full_kit = bool(data.get('full_kit', True))
    try:
        akb_count = int(data.get('akb_count') or 0)
    except (TypeError, ValueError):
        akb_count = 0
    if akb_count < 0 or akb_count > 500:
        return jsonify({'success': False, 'error': 'Количество АКБ указано неверно'})
    if kind == 'zamena':
        take, missing_take = bikes_from_request(data, 'take_bikes', 'take_lines', catalog)
        give, missing_give = bikes_from_request(data, 'give_bikes', 'give_lines', catalog)
    else:
        take, missing_take = bikes_from_request(data, 'bikes', 'lines', catalog)
        give, missing_give = [], []
    missing = missing_take + missing_give
    if missing:
        return jsonify({'success': False, 'error': 'Не разобрал: ' + ', '.join(missing[:8])})
    if kind == 'zamena':
        if not take or not give:
            return jsonify({'success': False, 'error': 'Для замены нужны оба списка: что забрали и что поставили'})
        bikes = take
        extra_akb = 0
    else:
        if not take:
            return jsonify({'success': False, 'error': 'Добавьте хотя бы один велосипед'})
        bikes = take
        extra_akb = len(bikes) * 2 if kind == 'postavka' or (kind == 'vyvoz' and full_kit) else (akb_count if kind == 'vyvoz' else 0)
    template_name, operation = DOC_KINDS[kind]
    if company == 'ВелоСтеп':
        template_name = template_name.replace('_template.docx', '_velostep.docx')
    path = os.path.join(DOC_DIR, template_name)
    if not os.path.exists(path):
        return jsonify({'success': False, 'error': 'Шаблон документа не найден на сервере'})
    try:
        from docx import Document
        from io import BytesIO
        doc = Document(path)
        akb_total = len(bikes) * 2 if kind != 'vyvoz' else extra_akb
        if kind == 'zamena':
            akb_total = 0
        fill_doc_text(doc, {
            '{ДАТА}': date,
            '{ДС}': darks,
            '{АДРЕС}': address,
            '{ЗАБОР}': len(take),
            '{ПОСТАВКА}': len(give),
            '{ВЕЛО_ИТОГО}': len(bikes),
            '{АКБ_ИТОГО}': akb_total,
            '{ВСЕГО}': len(bikes) + akb_total,
            '{таблица}': '',
            '{КОМПАНИЯ}': company,
        })
        if kind == 'zamena':
            add_bike_table(doc.tables[0], take, False, 0, company)
            add_bike_table(doc.tables[1], give, False, 0, company)
        elif kind == 'sverka':
            add_bike_table(doc.tables[0], bikes, False, 0, company)
        else:
            add_bike_table(doc.tables[0], bikes, True, extra_akb, company)
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
    except Exception as exc:
        logger.warning(f'Документ не собрался: {exc}')
        return jsonify({'success': False, 'error': 'Не удалось собрать документ'})
    log_note = 'ok'
    log_rows = []
    if kind == 'zamena':
        for bike in take:
            log_rows.append([date, 'Вывоз', company, darks, bike.get('frame') or '', bike.get('gos') or '', bike.get('iot') or ''])
        for bike in give:
            log_rows.append([date, 'Поставка', company, darks, bike.get('frame') or '', bike.get('gos') or '', bike.get('iot') or ''])
    else:
        title = {'postavka': 'Поставка', 'vyvoz': 'Вывоз', 'sverka': 'Сверка'}.get(kind, operation)
        for bike in bikes:
            log_rows.append([date, title, company, darks, bike.get('frame') or '', bike.get('gos') or '', bike.get('iot') or ''])
    try:
        if log_rows:
            append_supply_log(log_rows)
    except Exception as exc:
        logger.warning(f'Журнал поставки не записался: {exc}')
        log_note = 'sheet-failed'
    filename = f'{darks}_{operation}_{date.replace(".", "")}.docx'
    response = send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    response.headers['X-Log-Note'] = log_note
    return response

# ============================================================
# МАРШРУТЫ
# ============================================================
@app.route('/')
def index():
    return redirect(url_for('login_page'))

@app.route('/api/db_status')
@login_required
def api_db_status():
    if session.get('role') != 'admin':
        return jsonify({'success': False, 'error': 'Только для руководителя'}), 403
    info = db.status()
    info['success'] = True
    return jsonify(info)

@app.route('/login')
def login_page():
    return render_template('login.html')

@app.route('/api/login', methods=['POST'])
def api_login():
    data = request.json
    login = data.get('login', '').strip().lower()
    password = data.get('password', '').strip()
    if login in MASTER_CREDENTIALS and MASTER_CREDENTIALS[login]['password'] == password:
        role = MASTER_CREDENTIALS[login].get('role', 'master')
        return jsonify({
            'success': True,
            'master': MASTER_CREDENTIALS[login]['name'],
            'login': login,
            'role': role
        })
    return jsonify({'success': False, 'error': 'Неверный логин или пароль'})

@app.route('/auto_login/<login>')
def auto_login(login):
    if login in MASTER_CREDENTIALS:
        master_name = MASTER_CREDENTIALS[login]['name']
        role = MASTER_CREDENTIALS[login].get('role', 'master')
        session['master_login'] = login
        session['master_name'] = master_name
        session['authenticated'] = True
        session['role'] = role

        if role == 'admin':
            return redirect(url_for('admin_panel'))
        elif role == 'iot':
            return redirect(url_for('iot_main'))
        else:
            return redirect(url_for('master_overview', name=master_name))
    return redirect(url_for('login_page'))

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login_page'))

@app.after_request
def no_store_master_pages(response):
    if request.path.startswith('/master'):
        response.headers['Cache-Control'] = 'no-store'
    return response

# ============================================================
# СТАТИКА
# ============================================================
@app.route('/manifest.json')
def web_manifest():
    response = send_from_directory('static', 'manifest.json')
    response.headers['Content-Type'] = 'application/manifest+json'
    return response

@app.route('/.well-known/assetlinks.json')
def assetlinks():
    return send_from_directory('static/.well-known', 'assetlinks.json')

@app.route('/static/<path:filename>')
def static_files(filename):
    return send_from_directory('static', filename)

@app.route('/sw.js')
def service_worker():
    response = send_from_directory('static', 'sw.js')
    response.headers['Cache-Control'] = 'no-cache'
    return response

@app.route('/data/<path:filename>')
def data_files(filename):
    return send_from_directory('/data', filename)

# ============================================================
# АДМИН-ПАНЕЛЬ
# ============================================================
@app.route('/admin')
@login_required
def admin_panel():
    if session.get('role') != 'admin':
        return redirect(url_for('master_overview', name=session.get('master_name')))
    cache_data = get_admin_cache()
    if cache_data and 'directions' in cache_data:
        directions = cache_data.get('directions', {})
    else:
        tickets = get_tickets_from_sheets()
        save_admin_cache(tickets)
        cache_data = get_admin_cache()
        directions = cache_data.get('directions', {})
    load_darks_reference(force=True)
    darks_people = {}
    for number, info in (darks_ref or {}).items():
        names = ((info or {}).get('curators') or '').strip()
        ids = ((info or {}).get('curator_ids') or '').strip()
        shown = names
        if ids and ids not in names:
            shown = (names + ', ' + ids).strip(', ')
        darks_people[str(number)] = {
            'curators': shown,
            'address': (info or {}).get('address') or '',
        }
    state = db.status()
    active = db.count_active_tickets() if state.get('configured') else None
    state = db.status()
    if not state.get('configured'):
        db_line = 'База не настроена. Сайт работает через таблицу.'
    elif state.get('last_error'):
        db_line = 'База не отвечает: ' + str(state.get('last_error'))[:180]
    elif state.get('connected'):
        if active is None:
            active_text = 'активные заявки не посчитались'
        else:
            active_text = f'Активных заявок: {active} (все, кроме «Выполнено»)'
        db_line = f"База на связи. {active_text}. Последняя запись: {state.get('last_sync') or 'ещё не было'}."
    else:
        db_line = 'База ещё не ответила. Если так и останется, на Amvera проверьте, что база запущена.'
    return render_template('admin.html',
                          directions=directions,
                          masters=MASTERS,
                          darks_people=darks_people,
                          curator_sent=curator_sent_map(),
                          db_line=db_line,
                          now=get_msk_now().strftime('%H:%M:%S'))

# ============================================================
# API ДЛЯ АДМИНКИ
# ============================================================
@app.route('/api/sync')
@login_required
def api_sync():
    try:
        flush_queue()
        added = import_new_tickets_from_google()
        tickets = get_tickets_from_sheets(force=False) or []
        if tickets:
            publish_caches(tickets)
        logger.info(f"Синхронизация: в работе {len(tickets)}, новых из таблицы {added}")
        return jsonify({
            'success': True,
            'cached': added == 0,
            'count': len(tickets),
            'new': added,
            'message': f'Новых заявок с сайта: {added}. Остальное читается из базы.',
        })
    except Exception as e:
        logger.error(f"Ошибка синхронизации: {e}")
        return jsonify({'success': False, 'error': str(e)})

@app.route('/api/restore_from_sheet')
@login_required
def api_restore_from_sheet():
    if session.get('role') != 'admin':
        return jsonify({'success': False, 'error': 'Только для руководителя'}), 403
    try:
        result = restore_statuses_from_sheet()
        return jsonify({
            'success': True,
            'count': result.get('count'),
            'from_sheet': result.get('from_sheet'),
            'message': 'Статусы записаны в базу сразу. Обновите таблицу tickets в DBeaver.',
        })
    except Exception as e:
        logger.error(f'Не удалось вернуть статусы: {e}')
        return jsonify({'success': False, 'error': str(e)})

@app.route('/api/reconcile_sheet', methods=['POST'])
@login_required
def api_reconcile_sheet():
    if session.get('role') != 'admin':
        return jsonify({'success': False, 'error': 'Только для руководителя'}), 403
    try:
        result = reconcile_from_google()
        return jsonify({
            'success': True,
            'sheet': result['sheet'],
            'open': result['open'],
            'closed': result['closed'],
            'message': f"В таблице {result['sheet']} заявок, открытых {result['open']}. Лишние в базе закрыты: {result['closed']}.",
        })
    except Exception as e:
        logger.error(f'Сверка с таблицей не удалась: {e}')
        return jsonify({'success': False, 'error': str(e)})


@app.route('/api/queue_status')
@login_required
def api_queue_status():
    tasks = read_queue().get('tasks') or []
    uids = []
    for task in tasks:
        uid = task.get('uid')
        if uid and uid not in uids:
            uids.append(uid)
    return jsonify({'success': True, 'count': len(tasks), 'uids': uids})

@app.route('/api/flush_now', methods=['POST'])
@login_required
def api_flush_now():
    try:
        flushed = flush_queue()
        return jsonify({'success': True, 'flushed': flushed})
    except Exception as e:
        logger.error(f"Ошибка записи очереди: {e}")
        return jsonify({'success': False, 'error': str(e)})

def master_chat_ids(master_name):
    load_darks_reference()
    wanted = clean_name(master_name)
    found = []
    for info in darks_ref.values():
        if clean_name(info.get('master')) != wanted:
            continue
        digits = ''.join(ch for ch in str(info.get('master_tg') or '') if ch.isdigit())
        if digits and digits not in found:
            found.append(digits)
    logger.info(f'ТГ мастера {master_name}: {found or "не найден"}')
    return found

def master_route_text(master_name):
    cache = get_master_cache(master_name) or {}
    tickets = [t for t in cache.get('tickets') or [] if is_active_status(t.get('status'))]
    enrich_tickets(tickets)
    lines = [
        tg_plain(master_name),
        get_msk_now().strftime('%d.%m.%Y'),
        '',
        f'Вам назначены заявки: {len(tickets)}',
    ]
    if not tickets:
        return 'На вас сейчас нет открытых заявок.'
    by_dark = {}
    for ticket in tickets:
        by_dark.setdefault(str(ticket.get('darks') or '—'), []).append(ticket)
    for dark, group in by_dark.items():
        address = (group[0].get('address') or '').strip()
        lines.append('')
        lines.append(f'<b>{tg_plain(address)} (даркстор {tg_plain(dark)})</b>')
        for ticket in group[:20]:
            label = ticket_label(ticket)
            desc = (ticket.get('desc') or '').replace('\n', ' ').strip()
            prefix = 'ЗАМЕНА. ' if is_evacuation_ticket(ticket) else ''
            lines.append(f'• {tg_plain(label)} — {tg_plain(prefix + desc[:90])}')
        extra = len(group) - 20
        if extra > 0:
            lines.append(f'и ещё {extra}')
    return '\n'.join(lines)

def notify_master_route(master_name):
    chat_ids = master_chat_ids(master_name)
    if not chat_ids:
        logger.info(f'У мастера {master_name} нет ТГ ID в столбце I листа Дарксторы')
        return {'sent': 0, 'reason': 'no_id'}
    text = master_route_text(master_name)
    sent = 0
    last_error = ''
    for chat_id in chat_ids:
        try:
            response = requests.post(
                f'{BOT_API_URL}/notify_master',
                json={'chat_id': chat_id, 'text': text},
                timeout=25,
            )
            if response.status_code == 200 and (response.json() or {}).get('success'):
                sent += 1
            else:
                last_error = response.text[:200]
                logger.warning(f'Бот не принял сообщение мастеру {master_name}: {last_error}')
        except Exception as e:
            last_error = str(e)
            logger.warning(f'Не отправилось сообщение мастеру {master_name}: {e}')
    if sent:
        return {'sent': sent, 'reason': 'ok'}
    return {'sent': 0, 'reason': 'bot', 'error': last_error}

@app.route('/api/batch_update', methods=['POST'])
@login_required
def api_batch_update():
    data = request.json
    changes = data.get('changes', [])
    if not changes:
        return jsonify({'success': False, 'error': 'Нет изменений'})
    updated = batch_update_masters(changes)
    if updated:
        names = []
        for change in changes:
            master = (change.get('master') or '').strip()
            if master and master not in names:
                names.append(master)
        for master in names:
            refresh_master_cache(master)
    return jsonify({'success': True, 'updated': updated})

@app.route('/api/send_route', methods=['POST'])
@login_required
def api_send_route():
    data = request.json
    master = data.get('master')
    if not master:
        return jsonify({'success': False, 'error': 'Не указан мастер'})
    count = refresh_master_cache(master)
    info = notify_master_route(master)
    return jsonify({
        'success': True,
        'message': f'Кэш мастера {master} обновлен',
        'notified': info.get('sent', 0),
        'notify_reason': info.get('reason', ''),
        'tickets': count,
    })

@app.route('/api/send_route_all', methods=['POST'])
@login_required
def api_send_route_all():
    refresh_all_master_caches()
    for master in MASTERS:
        notify_master_route(master)
    return jsonify({'success': True, 'message': 'Кэши всех мастеров обновляются'})

@app.route('/api/clear_dates', methods=['POST'])
@login_required
def api_clear_dates():
    cleared = clear_all_masters()
    return jsonify({'success': True, 'cleared': cleared})

@app.route('/api/clear_master', methods=['POST'])
@login_required
def api_clear_master():
    data = request.json or {}
    master = (data.get('master') or '').strip()
    if not master:
        return jsonify({'success': False, 'error': 'Не указан мастер'})
    cleared = clear_one_master(master)
    return jsonify({'success': True, 'cleared': cleared})

@app.route('/api/notify_curators', methods=['POST'])
@login_required
def api_notify_curators():
    data = request.json
    message = data.get('message', '')
    if not message:
        return jsonify({'success': False, 'error': 'Нет сообщения'})
    result = notify_curators(message)
    return jsonify(result)

def refresh_master_caches_from_admin():
    admin_cache = get_admin_cache() or {}
    all_tickets = []
    for dir_data in admin_cache.get('directions', {}).values():
        all_tickets.extend(dir_data.get('tickets', []))
    apply_pending_queue(all_tickets)
    for master in MASTERS:
        master_tickets = [t for t in all_tickets if same_master(t.get('master'), master) and is_active_status(t.get('status'))]
        save_master_cache(master, master_tickets, '')

def queue_admin_status(uid, status_display, note='', skip_report=True, actor='Админ'):
    now = get_msk_now().strftime('%d.%m %H:%M')
    actor = actor or 'Админ'
    if status_display == '⏹️ Обработано':
        moment = get_msk_now()
        timer_from = moment.strftime('%Y-%m-%d %H:%M:%S')
        line = f"Вело отсутствует {moment.strftime('%d.%m %H:%M')}, {actor}"
        ticket = find_cached_ticket(uid)
        previous = (ticket.get('note') or '').strip() if ticket else ''
        note_text = f'{previous}\n{line}'.strip() if previous else line
        update_ticket_in_admin_cache(uid, 'processed', note_text, extra={
            'master': '',
            'timer_from': timer_from,
            'hours_since': 0,
            'evac_reason': '',
            'is_done': False,
            'is_active': True,
        })
        remember_master_cleared(uid)
        remove_ticket_from_master_cache((ticket or {}).get('master') or '', uid)
        add_to_queue({
            'uid': uid,
            'source': source_of(uid),
            'type': 'status_update',
            'data': {
                'status': '⏹️ Обработано',
                'new_status': 'processed',
                'note': note_text,
                'skip_report': True,
                'master': '',
                'write_master': True,
                'timer_from': timer_from,
            }
        })
        return
    ticket = find_cached_ticket(uid) or {}
    previous = (ticket.get('note') or '').strip()
    extra = None
    if status_display == '🔧 Эвакуация':
        if is_supply(ticket):
            logger.info(f'Эвакуация запрещена для {uid}: это аккумулятор или зарядка')
            return
        new_status = 'evacuation'
        plain = (note or 'Эвакуация').replace('ЭВАКУАЦИЯ: ', '').strip() or 'Эвакуация'
        line = 'ЭВАКУАЦИЯ: ' + plain
        sheet_note = f'{previous}\n{line}'.strip() if previous else line
        cache_note = sheet_note
        extra = {'evac_reason': plain, 'master': ''}
    elif status_display == '🔵 Доделать':
        new_status = 'todo'
        text = (note or '').strip() or 'Доделать'
        line = f'{now}, {actor}. Доделать: {text}'
        sheet_note = f'{previous}\n{line}'.strip() if previous else line
        cache_note = sheet_note
    else:
        status_map = {
            '🟡 В работе': 'pending',
            '✅ Выполнено': 'done',
            '🔵 Доделать': 'todo'
        }
        new_status = status_map.get(status_display, 'pending')
        cache_note = note
        sheet_note = note
        extra = {'evac_reason': ''}
    source = 'Заявки'
    if uid in uid_index:
        source = (uid_index[uid].get('ticket') or {}).get('source') or 'Заявки'
    update_ticket_in_admin_cache(uid, new_status, cache_note, extra=extra)
    data = {
        'status': status_display,
        'new_status': new_status,
        'note': sheet_note,
        'skip_report': skip_report,
    }
    if new_status == 'evacuation':
        data['master'] = ''
        data['write_master'] = True
        remember_master_cleared(uid)
        remove_ticket_from_master_cache(ticket.get('master') or '', uid)
    add_to_queue({
        'uid': uid,
        'source': source,
        'type': 'status_update',
        'data': data,
    })

@app.route('/api/update_status', methods=['POST'])
@login_required
def api_update_status():
    data = request.json
    uid = data.get('uid')
    status_display = data.get('status')
    note = data.get('note', '')
    skip_report = data.get('skip_report', True)
    if not uid or not status_display:
        return jsonify({'success': False, 'error': 'Недостаточно данных'})
    ticket = find_cached_ticket(uid) or {}
    if status_display == '🔧 Эвакуация' and is_supply(ticket):
        return jsonify({'success': False, 'error': 'Для аккумулятора и зарядки эвакуацию ставить нельзя'})
    try:
        queue_admin_status(uid, status_display, note, skip_report, 'Админ' if session.get('role') == 'admin' else (session.get('master_name') or 'Админ'))
        refresh_master_caches_from_admin()
        return jsonify({'success': True})
    except Exception as e:
        logger.error(f"Ошибка обновления статуса: {e}")
        return jsonify({'success': False, 'error': str(e)})

@app.route('/api/bulk_update_status', methods=['POST'])
@login_required
def api_bulk_update_status():
    try:
        data = request.json
        uids = data.get('uids', [])
        status_display = data.get('status', '')
        comment = data.get('comment', '')
        skip_report = data.get('skip_report', True)
        if not uids or not status_display:
            return jsonify({'success': False, 'error': 'Не указаны UID или статус'}), 400
        actor = 'Админ' if session.get('role') == 'admin' else (session.get('master_name') or 'Админ')
        for uid in uids:
            queue_admin_status(uid, status_display, comment, skip_report, actor)
        refresh_master_caches_from_admin()
        return jsonify({'success': True, 'updated': len(uids), 'message': f'Обновлено {len(uids)} заявок'})
    except Exception as e:
        logger.error(f'Ошибка при массовом обновлении статусов: {e}')
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/admin_action', methods=['POST'])
@login_required
def api_admin_action():
    data = request.json
    uid = data.get('uid')
    action = data.get('action')
    extra = data.get('extra', '')
    skip_report = data.get('skip_report', True)
    if not uid or not action:
        return jsonify({'success': False, 'error': 'Недостаточно данных'})
    try:
        actor = session.get('master_name') or 'Админ'
        if action == 'done':
            queue_admin_status(uid, '✅ Выполнено', extra or 'Выполнено админом', True, actor)
        elif action == 'evacuation':
            queue_admin_status(uid, '🔧 Эвакуация', extra or 'Эвакуация', True, actor)
        elif action == 'fail':
            mark_bike_missing(uid, actor)
        elif action == 'todo':
            queue_admin_status(uid, '🔵 Доделать', extra or 'Доделать', True, actor)
        elif action == 'taken':
            note = f'ЗАБРАЛИ: {extra} АКБ'
            update_ticket_in_admin_cache(uid, 'todo', note)
            add_to_queue({
                'uid': uid,
                'source': source_of(uid),
                'type': 'taken_no_replace',
                'data': {'parts': extra, 'note': note, 'master': actor}
            })
        else:
            return jsonify({'success': False, 'error': 'Неизвестное действие'})
        refresh_master_caches_from_admin()
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})

@app.route('/api/build_route_for_master')
@login_required
def api_build_route_for_master():
    master = request.args.get('master', '')
    if not master:
        return 'Нет мастера', 400
    if session.get('role') != 'admin' and session.get('master_name') == master:
        blocked = guard_shift(master)
        if blocked:
            return blocked
    
    darks_ref = load_darks_reference()
    
    cache_data = get_master_cache(master)
    if cache_data:
        master_tickets = cache_data.get('tickets', [])
    else:
        tickets = get_tickets_from_sheets()
        master_tickets = [t for t in tickets if t.get('master') == master and is_active_status(t.get('status'))]
    
    ordered = []
    seen = set()
    for ticket in master_tickets:
        darks_num = str(ticket.get('darks') or '')
        if not darks_num or darks_num in seen or darks_num not in darks_ref:
            continue
        coords = (darks_ref[darks_num].get('coords') or '').strip()
        if coords:
            seen.add(darks_num)
            ordered.append(coords)
    coords_list = ordered
    if not coords_list:
        return 'Нет координат', 400

    def lat_lon(value):
        parts = [part.strip() for part in str(value).replace(';', ',').split(',') if part.strip()]
        if len(parts) < 2:
            return None
        try:
            first = float(parts[0])
            second = float(parts[1])
        except ValueError:
            return None
        if abs(first) > 50 and abs(second) < 50:
            return first, second
        if abs(second) > 50 and abs(first) < 50:
            return second, first
        return first, second

    points = []
    for value in [START_COORDS] + coords_list:
        point = lat_lon(value)
        if point and point not in points:
            points.append(point)
    if len(points) < 2:
        return 'Не хватает координат для маршрута', 400
    user_agent = request.headers.get('User-Agent', '').lower()
    is_mobile = any(x in user_agent for x in ['android', 'iphone', 'ipad', 'mobile'])
    if is_mobile:
        start = points[0]
        finish = points[-1]
        vias = points[1:-1][:20]
        query = [
            f'lat_from={start[0]}',
            f'lon_from={start[1]}',
            f'lat_to={finish[0]}',
            f'lon_to={finish[1]}',
        ]
        for index, point in enumerate(vias):
            query.append(f'lat_via_{index}={point[0]}')
            query.append(f'lon_via_{index}={point[1]}')
        url = 'yandexnavi://build_route_on_map?' + '&'.join(query)
    else:
        url = 'https://yandex.ru/maps/?rtext=' + '~'.join(f'{lat},{lon}' for lat, lon in points) + '&rtt=auto'
    return redirect(url)

# ============================================================
# МАРШРУТЫ МАСТЕРОВ
# ============================================================
def master_groups(name):
    cache_data = get_master_cache(name) or {}
    master_tickets = list(cache_data.get('tickets') or [])
    if not master_tickets and not cache_data:
        tickets = get_tickets_from_sheets()
        master_tickets = [t for t in tickets if same_master(t.get('master'), name) and is_active_status(t.get('status'))]
        enrich_tickets(master_tickets)
        save_master_cache(name, master_tickets, '')
    else:
        enrich_tickets(master_tickets)
    groups = []
    if master_tickets:
        groups_dict = {}
        for t in master_tickets:
            darks_num = str(t.get('darks') or 'без номера')
            if darks_num not in groups_dict:
                groups_dict[darks_num] = {
                    'darks_number': darks_num,
                    'address': t.get('address') or 'Адрес не указан',
                    'contact': t.get('contact', ''),
                    'tickets': [],
                    'pending': 0,
                    'done': 0
                }
            groups_dict[darks_num]['tickets'].append(t)
            if t.get('status') in ['pending', 'todo', 'processed', 'fail', 'evacuation']:
                groups_dict[darks_num]['pending'] += 1
            else:
                groups_dict[darks_num]['done'] += 1
        for darks_num in sorted(groups_dict.keys(), key=lambda x: int(x) if str(x).isdigit() else 999999):
            item = groups_dict[darks_num]
            groups.append({
                'darks_number': item['darks_number'],
                'address': item['address'],
                'contact': item['contact'],
                'pending': item['pending'],
                'done': item['done'],
                'tickets': item['tickets'],
            })
    return groups, master_tickets, cache_data.get('updated_at', '')

@app.route('/admin/stats')
@login_required
def admin_stats():
    if session.get('role') != 'admin':
        return redirect(url_for('login_page'))
    bundle = stats_bundle()
    reports = sorted(bundle['reports'], key=lambda row: (row.get('date') or '', row.get('master') or ''), reverse=True)
    expenses = [row for row in reports if (row.get('expenses') or '').strip()]
    return render_template(
        'admin_stats.html',
        people=bundle['people'],
        overall=bundle['overall_hours'],
        reports=reports[:80],
        expenses=expenses[:80],
        ratings=list(reversed(load_ratings().get('items') or []))[:40],
        now=get_msk_now().strftime('%d.%m.%Y %H:%M'),
    )

@app.route('/master/<name>/stats')
@login_required
def master_stats(name):
    if session.get('master_name') != name:
        return redirect(url_for('login_page'))
    blocked = guard_shift(name)
    if blocked:
        return blocked
    bundle = stats_bundle(name)
    person = bundle['people'][0] if bundle['people'] else {'name': name, 'open': 0, 'week': 0, 'month': 0, 'avg_hours': None, 'rating': None}
    return render_template('master_stats.html', name=name, person=person, overall=bundle['overall_hours'])

@app.route('/master/<name>/report')
@login_required
def master_report(name):
    if session.get('master_name') != name:
        return redirect(url_for('login_page'))
    day_text = (request.args.get('day') or '').strip()
    missing = missing_shift_day(name)
    if missing and day_text != missing.isoformat():
        return redirect(url_for('master_report', name=name, day=missing.isoformat()))
    if day_text:
        try:
            day = datetime.strptime(day_text, '%Y-%m-%d').date()
        except ValueError:
            day = get_msk_now().date()
    else:
        day = get_msk_now().date()
    forced = bool(missing and missing.isoformat() == day.isoformat())
    return render_template(
        'master_report.html',
        name=name,
        day=day.isoformat(),
        day_ru=day.strftime('%d.%m.%Y'),
        forced=forced,
        closed_count=closed_on_day(name, day.isoformat()),
    )

@app.route('/master/<name>/report', methods=['POST'])
@login_required
def master_report_save(name):
    if session.get('master_name') != name:
        return jsonify({'success': False, 'error': 'Доступ запрещён'}), 403
    data = request.json or {}
    day_text = (data.get('date') or '').strip()
    try:
        day = datetime.strptime(day_text, '%Y-%m-%d').date()
    except ValueError:
        return jsonify({'success': False, 'error': 'Не понял дату смены'})
    missing = missing_shift_day(name)
    if missing and day != missing:
        return jsonify({'success': False, 'error': 'Сначала закройте отчёт за ' + missing.strftime('%d.%m.%Y')})
    start = (data.get('start') or '').strip()
    end = (data.get('end') or '').strip()
    status = (data.get('status') or '').strip()
    comment = (data.get('comment') or '').strip()
    mileage = (data.get('mileage') or '').strip()
    expenses = (data.get('expenses') or '').strip()
    if not re.fullmatch(r'\d{2}:\d{2}', start) or not re.fullmatch(r'\d{2}:\d{2}', end):
        return jsonify({'success': False, 'error': 'Укажите время начала и конца смены'})
    if status not in SHIFT_STATUSES:
        return jsonify({'success': False, 'error': 'Выберите статус смены'})
    if status != 'done' and not comment:
        return jsonify({'success': False, 'error': 'Напишите комментарий, почему смена закрыта не полностью'})
    if not mileage or mileage.lower() == 'нет':
        mileage = 'Нет'
    item = {
        'date': day.isoformat(),
        'date_ru': day.strftime('%d.%m.%Y'),
        'master': name,
        'start': start,
        'end': end,
        'mileage': mileage,
        'expenses': expenses,
        'status': status,
        'status_label': SHIFT_STATUSES[status],
        'comment': comment,
        'closed_count': closed_on_day(name, day.isoformat()),
        'saved_at': get_msk_now().strftime('%Y-%m-%d %H:%M:%S'),
    }
    save_shift_report_item(item)
    db_ok = False
    try:
        db_ok = db.save_shift_report(item)
    except Exception as exc:
        logger.warning(f'Отчёт в базу: {exc}')
    sheet_ok = True
    try:
        append_shift_sheet(item)
    except Exception as exc:
        sheet_ok = False
        logger.warning(f'Отчёт в таблицу: {exc}')
    note = ''
    if not sheet_ok:
        note = 'Отчёт принят. В таблицу Excel сейчас не записался, копия в базе сайта есть.'
    elif not db_ok:
        note = 'Отчёт принят и записан в таблицу.'
    return jsonify({'success': True, 'message': note or 'Отчёт сохранён', 'next': url_for('master_overview', name=name)})

@app.route('/master/<name>')
@login_required
def master_overview(name):
    if session.get('master_name') != name:
        return redirect(url_for('login_page'))
    blocked = guard_shift(name)
    if blocked:
        return blocked
    groups, master_tickets, updated_at = master_groups(name)
    for t in master_tickets:
        t['hours_since'] = ticket_hours(t)
        t['is_today_done'] = t.get('status') == 'done'
    route_url = f'/api/build_route_for_master?master={name}'
    return render_template('master_home.html',
                          name=name,
                          groups=groups,
                          total_tickets=len(master_tickets),
                          route_url=route_url,
                          updated_at=updated_at,
                          now=get_msk_now().strftime('%H:%M:%S'))

@app.route('/api/master/<name>/home')
@login_required
def api_master_home(name):
    if session.get('master_name') != name and session.get('role') != 'admin':
        return jsonify({'success': False, 'error': 'Доступ запрещён'}), 403
    groups, master_tickets, updated_at = master_groups(name)
    return jsonify({
        'success': True,
        'total': len(master_tickets),
        'updated_at': updated_at,
        'groups': [{
            'darks_number': g['darks_number'],
            'address': g['address'],
            'pending': g['pending'],
        } for g in groups],
    })

@app.route('/master/<name>/darks/<darks_number>')
@login_required
def master_darks(name, darks_number):
    if session.get('master_name') != name:
        return redirect(url_for('login_page'))
    blocked = guard_shift(name)
    if blocked:
        return blocked
    cache_data = get_master_cache(name)
    if cache_data:
        all_tickets = cache_data.get('tickets', [])
        filtered = [t for t in all_tickets if str(t.get('darks')) == str(darks_number)]
    else:
        tickets = get_tickets_from_sheets()
        filtered = [t for t in tickets if same_master(t.get('master'), name) and str(t.get('darks')) == str(darks_number) and is_active_status(t.get('status'))]
        enrich_tickets(filtered)
        save_master_cache(name, filtered, '')
    enrich_tickets(filtered)
    
    for t in filtered:
        t['hours_since'] = ticket_hours(t)
    fill_old_bike(filtered)
    
    if is_transit_master(name):
        gos_counts = {}
        for t in filtered:
            gos = t.get('gos')
            if gos:
                if gos not in gos_counts:
                    gos_counts[gos] = 0
                gos_counts[gos] += 1
        for t in filtered:
            gos = t.get('gos')
            if gos and gos in gos_counts:
                t['duplicate_count'] = gos_counts[gos]
    
    address = filtered[0].get('address', 'Адрес не указан') if filtered else ''
    contact = filtered[0].get('contact', '') if filtered else ''
    return render_template('master_darks.html',
                          name=name,
                          darks_number=darks_number,
                          address=address,
                          contact=contact,
                          tickets=filtered,
                          now=get_msk_now().strftime('%H:%M:%S'))

@app.route('/master/<name>/text_plan')
@login_required
def master_text_plan(name):
    if session.get('master_name') != name:
        return redirect(url_for('login_page'))
    blocked = guard_shift(name)
    if blocked:
        return blocked
    cache_data = get_master_cache(name)
    if cache_data:
        master_tickets = cache_data.get('tickets', [])
    else:
        tickets = get_tickets_from_sheets()
        master_tickets = [t for t in tickets if t.get('master') == name and is_active_status(t.get('status'))]
    for t in master_tickets:
        t['hours_since'] = ticket_hours(t)
    pending_count = len([t for t in master_tickets if t.get('status') in ('pending', 'processed', 'fail')])
    todo_count = len([t for t in master_tickets if t.get('status') == 'todo'])
    return render_template('text_plan.html',
                          name=name,
                          tickets=master_tickets,
                          pending_count=pending_count,
                          todo_count=todo_count,
                          now=get_msk_now().strftime('%H:%M:%S'))

@app.route('/master/<name>/history')
@login_required
def master_history(name):
    if session.get('master_name') != name:
        return redirect(url_for('login_page'))
    blocked = guard_shift(name)
    if blocked:
        return blocked
    return render_template('master_history.html', name=name, now=get_msk_now().strftime('%H:%M:%S'))

@app.route('/api/master_history/<name>')
@login_required
def api_master_history(name):
    if session.get('master_name') != name and session.get('role') != 'admin':
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    history = read_cache(f"history_{name}.json") or []
    if isinstance(history, dict):
        history = history.get('items') or history.get('history') or []
    if not isinstance(history, list):
        history = []
    return jsonify({'success': True, 'history': history})

@app.route('/api/master/<name>/trunk')
@login_required
def api_master_trunk(name):
    if session.get('master_name') != name and session.get('role') != 'admin':
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    rows, total_qty = trunk_rows(name)
    return jsonify({'success': True, 'parts': rows, 'total_qty': total_qty})

@app.route('/master/<name>/trunk', methods=['GET', 'POST'])
@login_required
def master_trunk(name):
    if session.get('master_name') != name:
        return redirect(url_for('login_page'))
    blocked = guard_shift(name)
    if blocked:
        return blocked
    notice = ''
    error = ''
    if request.method == 'POST':
        names = request.form.getlist('part')
        qtys = request.form.getlist('qty')
        chosen = set(request.form.getlist('take'))
        loaded = []
        for part_name, qty_raw in zip(names, qtys):
            if part_name not in chosen or part_name not in PARTS_SET:
                continue
            try:
                qty = int(qty_raw or 0)
            except Exception:
                qty = 0
            if qty <= 0:
                continue
            change_trunk(name, part_name, qty)
            queue_part_move(name, 'Взял', part_name, qty)
            loaded.append(f'{part_name} × {qty}')
        if not loaded:
            error = 'Отметьте запчасти и укажите количество'
        else:
            notice = 'Загружено: ' + ', '.join(loaded)
    rows, total_qty = trunk_rows(name)
    stock = [row for row in rows if row['qty'] > 0]
    return render_template(
        'trunk.html',
        name=name,
        parts=rows,
        stock=stock,
        total_qty=total_qty,
        error=error,
        notice=notice,
        now=get_msk_now().strftime('%H:%M:%S'),
    )

@app.route('/admin/trunks')
@login_required
def admin_trunks():
    if session.get('role') != 'admin':
        return redirect(url_for('login_page'))
    masters = []
    for master in MASTERS:
        rows, total_qty = trunk_rows(master)
        stock = [row for row in rows if row['qty'] > 0]
        masters.append({'name': master, 'rows': stock, 'total_qty': total_qty})
    today = get_msk_now().date()
    month_start = today.replace(day=1)
    week_start = today - timedelta(days=6)
    def period_day(value, fallback):
        try:
            return datetime.strptime(value or '', '%Y-%m-%d').date()
        except ValueError:
            return fallback
    start = period_day(request.args.get('start'), month_start)
    end = period_day(request.args.get('end'), today)
    if end < start:
        start, end = end, start
    by_master = {name: {} for name in MASTERS}
    total = {}
    for move in parts_log():
        if move.get('action') != 'Поставил':
            continue
        try:
            day = datetime.strptime(str(move.get('time') or '')[:10], '%Y-%m-%d').date()
        except ValueError:
            continue
        if day < start or day > end:
            continue
        part_name = move.get('name') or ''
        qty = int(move.get('qty') or 0)
        bucket = None
        for known in MASTERS:
            if same_master(known, move.get('master')):
                bucket = by_master[known]
                break
        if bucket is None:
            continue
        row = bucket.setdefault(part_name, {'name': part_name, 'qty': 0})
        row['qty'] += qty
        total_row = total.setdefault(part_name, {'name': part_name, 'qty': 0})
        total_row['qty'] += qty
    spent_masters = []
    for master in MASTERS:
        rows = sorted(by_master[master].values(), key=lambda item: item['qty'], reverse=True)
        spent_masters.append({'name': master, 'rows': rows, 'qty': sum(item['qty'] for item in rows)})
    spent_rows = sorted(total.values(), key=lambda item: item['qty'], reverse=True)
    return render_template(
        'admin_trunks.html',
        masters=masters,
        spent=spent_rows,
        spent_masters=spent_masters,
        start=start.isoformat(),
        end=end.isoformat(),
        week_start=week_start.isoformat(),
        month_start=month_start.isoformat(),
        today=today.isoformat(),
        now=get_msk_now().strftime('%H:%M:%S'),
    )

# ============================================================
# API ДЛЯ МАСТЕРОВ
# ============================================================
@app.route('/api/parts_catalog')
@login_required
def api_parts_catalog():
    return jsonify({'parts': [name for name, price in PARTS]})

@app.route('/master/<name>/darks/<darks_number>/done/<uid>', methods=['POST'])
@login_required
def master_done(name, darks_number, uid):
    logger.info(f"📝 master_done вызван: name={name}, uid={uid}")
    if session.get('master_name') != name:
        logger.warning(f"❌ Сессия не совпадает: {session.get('master_name')} != {name}")
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    parts_raw = request.form.get('parts_json', '').strip()
    comment = request.form.get('comment', '').strip()
    legacy = request.form.get('parts', '').strip()
    selected = []
    if parts_raw:
        try:
            selected = json.loads(parts_raw)
        except Exception:
            return jsonify({'success': False, 'error': 'Не понял список запчастей'})
        if not isinstance(selected, list):
            selected = []
    ticket = find_cached_ticket(uid) or {}
    raw = (legacy or comment or '').strip()
    if raw.lower() == 'нет' and not selected:
        parts = 'нет'
    elif raw == 'Отказ от эвакуации' and not selected:
        parts = raw
    else:
        ok, parts_text = spend_parts(
            name,
            selected,
            uid,
            ticket.get('gos', ''),
            darks_number,
        )
        if not ok:
            return jsonify({'success': False, 'error': parts_text})
        if not parts_text:
            return jsonify({'success': False, 'error': 'Выберите запчасть или напишите «нет»'})
        parts = parts_text
    event = parts
    update_ticket_in_admin_cache(uid, 'done', event)
    refresh_master_caches_from_admin()
    
    add_to_queue({
        'uid': uid,
        'source': source_of(uid),
        'type': 'done',
        'data': {'parts': event, 'detail': event, 'note': event, 'master': name, 'darks_number': darks_number, 'gos': ticket.get('gos', '')}
    })
    label = ticket_label(ticket)
    if parts == 'Отказ от эвакуации':
        notify_dark_event(darks_number, f'{label} отказ от эвакуации', uid, 'refuse')
    else:
        notify_dark_event(darks_number, f'✅ {label} заявка выполнена', uid, 'done')
    maybe_ask_rating(darks_number, name)
    return jsonify({'success': True})

@app.route('/master/<name>/darks/<darks_number>/fail/<uid>')
@login_required
def master_fail(name, darks_number, uid):
    if session.get('master_name') != name:
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    mark_bike_missing(uid, name)
    maybe_ask_rating(darks_number, name)
    return jsonify({'success': True})

@app.route('/master/<name>/darks/<darks_number>/evacuation/<uid>', methods=['POST'])
@login_required
def master_evacuation(name, darks_number, uid):
    if session.get('master_name') != name:
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    reason = request.form.get('reason', '')
    if not reason.strip():
        return jsonify({'success': False, 'error': 'Не указана причина'})
    ticket = find_cached_ticket(uid) or {}
    if is_supply(ticket):
        return jsonify({'success': False, 'error': 'Для аккумулятора и зарядки эвакуацию ставить нельзя'})
    reason = reason.strip()
    previous = (ticket.get('note') or '').strip()
    line = 'ЭВАКУАЦИЯ: ' + reason
    note = f'{previous}\n{line}'.strip() if previous else line
    update_ticket_in_admin_cache(uid, 'evacuation', note, extra={'evac_reason': reason, 'master': ''})
    remember_master_cleared(uid)
    remove_ticket_from_master_cache(name, uid)
    refresh_master_caches_from_admin()
    add_to_queue({
        'uid': uid,
        'source': source_of(uid),
        'type': 'evacuation',
        'data': {'reason': reason, 'note': note, 'master': name, 'darks_number': darks_number, 'gos': ticket.get('gos', '')}
    })
    label = ticket_label(find_cached_ticket(uid) or {'gos': request.form.get('gos', '')})
    notify_dark_event(darks_number, f'🔧 {label} нужно эвакуировать: {reason.strip()}', uid, 'evacuation')
    maybe_ask_rating(darks_number, name)
    return jsonify({'success': True})

@app.route('/master/<name>/darks/<darks_number>/taken_no_replace/<uid>', methods=['POST'])
@login_required
def master_taken_no_replace(name, darks_number, uid):
    if session.get('master_name') != name:
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    parts = request.form.get('parts', '')
    if not parts or int(parts) <= 0:
        return jsonify({'success': False, 'error': 'Укажите количество'})
    note = f'ЗАБРАЛИ: {parts} АКБ'
    update_ticket_in_admin_cache(uid, 'todo', note)
    refresh_master_caches_from_admin()
    add_to_queue({
        'uid': uid,
        'source': source_of(uid),
        'type': 'taken_no_replace',
        'data': {'parts': parts, 'note': note, 'master': name, 'darks_number': darks_number}
    })
    return jsonify({'success': True})

@app.route('/master/<name>/darks/<darks_number>/reset/<uid>')
@login_required
def master_reset(name, darks_number, uid):
    if session.get('master_name') != name:
        return redirect(url_for('login_page'))
    return redirect(url_for('master_darks', name=name, darks_number=darks_number))

@app.route('/master/<name>/darks/<darks_number>/replace_yes/<uid>', methods=['POST'])
@login_required
def master_replace_yes(name, darks_number, uid):
    if session.get('master_name') != name:
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    parts = request.form.get('parts', '')
    if not parts or int(parts) <= 0:
        return jsonify({'success': False, 'error': 'Укажите количество'})
    update_ticket_in_admin_cache(uid, 'done', parts)
    refresh_master_caches_from_admin()
    add_to_queue({
        'uid': uid,
        'source': source_of(uid),
        'type': 'replace_yes',
        'data': {'parts': parts, 'master': name, 'darks_number': darks_number}
    })
    maybe_ask_rating(darks_number, name)
    return jsonify({'success': True})

@app.route('/master/<name>/darks/<darks_number>/replace_no/<uid>', methods=['POST'])
@login_required
def master_replace_no(name, darks_number, uid):
    if session.get('master_name') != name:
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    reason = (request.form.get('reason') or '').strip() or 'Куратор не предоставил'
    update_ticket_in_admin_cache(uid, 'todo', reason)
    refresh_master_caches_from_admin()
    add_to_queue({
        'uid': uid,
        'source': source_of(uid),
        'type': 'replace_no',
        'data': {'reason': reason, 'master': name, 'darks_number': darks_number}
    })
    return jsonify({'success': True})

# ============================================================
# ТРАНЗИТ - ЗАМЕНА ВЕЛОСИПЕДА
# ============================================================
REPLACEMENT_BOOK_ID = "1s_hXPSWueMAo3W1pgLCG0eHhYKoW0uoIZ6CGWFf55VU"
REPLACEMENT_GID = 1296698244

def append_bike_swap(address, darks_number, old_data, new_data, master='', uid='', kind='Эвакуация'):
    try:
        book = get_gspread_client().open_by_key(REPLACEMENT_BOOK_ID)
    except Exception as e:
        text = str(e).lower()
        if '403' in text or 'permission' in text:
            raise RuntimeError('Нет доступа к таблице замен. Откройте её для route-cache@telegramsenderbot.iam.gserviceaccount.com') from e
        raise
    sheet = None
    for item in book.worksheets():
        title = (item.title or '').strip().lower()
        try:
            gid = int(item.id)
        except Exception:
            gid = item.id
        if title == 'замены вело' or gid == REPLACEMENT_GID:
            sheet = item
            break
    if sheet is None:
        names = ', '.join(item.title for item in book.worksheets())
        raise RuntimeError(f'Лист «Замены вело» не найден. В файле есть: {names}')
    filled = sheet.get('A:I')
    last = 1
    for index, row in enumerate(filled, start=1):
        if any(str(cell).strip() for cell in row):
            last = index
    target = last + 1
    sheet.update(
        f'A{target}:I{target}',
        [[
            get_msk_now().strftime('%d.%m.%Y %H:%M:%S'),
            address or '',
            str(darks_number or ''),
            (old_data or {}).get('serial', ''),
            (old_data or {}).get('gos', ''),
            (old_data or {}).get('iot', ''),
            (new_data or {}).get('serial', ''),
            (new_data or {}).get('gos', ''),
            (new_data or {}).get('iot', ''),
        ]],
        value_input_option='USER_ENTERED',
    )
    try:
        if sheet.col_count < 20:
            sheet.add_cols(20 - sheet.col_count)
        if not (sheet.acell('R1').value or '').strip():
            sheet.update('R1:T1', [['Мастер', 'ID заявки', 'Тип']], value_input_option='RAW')
        sheet.update(f'R{target}:T{target}', [[master or '', uid or '', kind or 'Эвакуация']], value_input_option='USER_ENTERED')
    except Exception as e:
        logger.warning(f'Строка {target} записана, мастер и номер заявки в конец не встали: {e}')
    logger.info(f"✅ Замена записана в «{sheet.title}», строка {target}, даркстор {darks_number}")
    db.mirror_replacement({
        'uid': uid,
        'master': master,
        'darks': darks_number,
        'address': address,
        'kind': kind,
        'old': old_data or {},
        'new': new_data or {},
    })

def write_transit_replacement(uid, master_name, darks_number, address, old_data, new_data):
    try:
        append_bike_swap(address, darks_number, old_data, new_data, master_name, uid, 'Замена')
        sheet_client = get_sheet_client()
        now = get_msk_now().strftime('%Y-%m-%d %H:%M:%S')
        try:
            report_sheet = sheet_client.worksheet("Эвакуация Транзит")
        except Exception:
            report_sheet = sheet_client.add_worksheet("Эвакуация Транзит", 100, 20)
            headers = ['Отметка времени', 'Адрес даркстора', 'Номер даркстора',
                      'ЗАБРАЛ - Серийный номер', 'ЗАБРАЛ - Гос номер', 'ЗАБРАЛ - Номер айот',
                      'ОТДАЛ - Серийный номер', 'ОТДАЛ - Гос номер', 'ОТДАЛ - Номер айот']
            for i, h in enumerate(headers, start=1):
                report_sheet.update_cell(1, i, h)
        current_rows = report_sheet.get_all_values()
        new_row_idx = len(current_rows) + 1
        updates = [
            {'range': f'A{new_row_idx}', 'values': [[now]]},
            {'range': f'B{new_row_idx}', 'values': [[address]]},
            {'range': f'C{new_row_idx}', 'values': [[darks_number]]},
            {'range': f'D{new_row_idx}', 'values': [[old_data.get('serial', '')]]},
            {'range': f'E{new_row_idx}', 'values': [[old_data.get('gos', '')]]},
            {'range': f'F{new_row_idx}', 'values': [[old_data.get('iot', '')]]},
            {'range': f'G{new_row_idx}', 'values': [[new_data.get('serial', '')]]},
            {'range': f'H{new_row_idx}', 'values': [[new_data.get('gos', '')]]},
            {'range': f'I{new_row_idx}', 'values': [[new_data.get('iot', '')]]}
        ]
        report_sheet.batch_update(updates)
        export_removal_reason(old_data, evacuation_reason(find_cached_ticket(uid) or {'desc': 'Замена'}))
        logger.info(f"✅ Записана замена велосипеда для заявки {uid}")
    except Exception as e:
        logger.error(f"Ошибка записи замены велосипеда: {e}")
        raise

@app.route('/master/transit/replace', methods=['POST'])
@login_required
def transit_replace():
    data = request.json
    uid = data.get('uid')
    master_name = data.get('master')
    darks_number = data.get('darks_number')
    address = data.get('address', '')
    old_data = data.get('old_data', {})
    new_data = data.get('new_data', {})
    try:
        write_transit_replacement(uid, master_name, darks_number, address, old_data, new_data)
        update_ticket_in_admin_cache(uid, 'done', 'Заменен')
        add_to_queue({
            'uid': uid,
            'source': source_of(uid),
            'type': 'transit_replace',
            'data': {'master': master_name, 'darks_number': darks_number}
        })
        return jsonify({'success': True, 'message': 'Замена выполнена'})
    except Exception as e:
        logger.error(f"Ошибка замены велосипеда: {e}")
        return jsonify({'success': False, 'error': str(e)})

# ============================================================
# ФУНКЦИИ ДЛЯ ЭВАКУАЦИИ
# ============================================================
def find_bike_in_database(gos_number, darks_number):
    try:
        sheet_client = get_sheet_client()
        worksheet = sheet_client.worksheet("База данных вело")
        rows = worksheet.get_all_values()
        
        if len(rows) <= 1:
            return None
        
        gos_clean = re.sub(r'[^0-9A-Za-zА-Яа-я]', '', gos_number).upper() if gos_number else ''
        
        for row in rows[1:]:
            if len(row) >= 7:
                gos = row[1].strip() if len(row) > 1 else ''
                darks = row[4].strip() if len(row) > 4 else ''
                gos_digits = re.sub(r'[^0-9A-Za-zА-Яа-я]', '', gos).upper()
                
                if gos and darks:
                    if (gos == gos_number or (gos_clean and gos_digits == gos_clean)) and darks == darks_number:
                        return {
                            'serial': row[0].strip() if len(row) > 0 else '',
                            'gos': gos,
                            'iot': row[2].strip() if len(row) > 2 else '',
                            'address': row[3].strip() if len(row) > 3 else '',
                            'darks': darks,
                            'bike_type': row[5].strip() if len(row) > 5 else '',
                            'direction': row[6].strip() if len(row) > 6 else ''
                        }
        return None
    except Exception as e:
        logger.error(f"Ошибка поиска велосипеда в БД: {e}")
        return None

@app.route('/api/get_bike_data')
@login_required
def api_get_bike_data():
    try:
        gos = request.args.get('gos', '')
        darks = request.args.get('darks', '')
        
        if not gos:
            return jsonify({'success': False, 'error': 'Не указан госномер'})
        
        bike_data = find_bike_in_database(gos, darks)
        if bike_data:
            return jsonify({'success': True, 'data': bike_data})
        return jsonify({'success': False, 'error': 'Велосипед не найден'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

def export_removal_reason(old_data, reason):
    try:
        book = get_gspread_client().open_by_key('1BoZ7GFmL2q56bjqt1CFVV_PeqnKVaeb3anTJkf0s6xA')
        sheet = None
        for item in book.worksheets():
            title = (item.title or '').strip().lower()
            try:
                gid = int(item.id)
            except Exception:
                gid = item.id
            if title == 'общее' or gid == 1658833390:
                sheet = item
                break
        if sheet is None:
            logger.warning('Лист «Общее» не найден')
            return
        rows = sheet.get('A:B')
        iot = re.sub(r'\s+', '', str((old_data or {}).get('iot') or '')).lower()
        serial = re.sub(r'\s+', '', str((old_data or {}).get('serial') or '')).lower()
        iot_row = None
        serial_row = None
        for index, row in enumerate(rows, start=1):
            if index == 1:
                continue
            frame = re.sub(r'\s+', '', str(row[0] if row else '')).lower()
            module = re.sub(r'\s+', '', str(row[1] if len(row) > 1 else '')).lower()
            if iot and module == iot:
                iot_row = index
                break
            if serial and frame == serial and serial_row is None:
                serial_row = index
        target = iot_row or serial_row
        if not target:
            logger.warning(f'Велосипед не найден в «Общее»: IoT {iot or "-"}, рама {serial or "-"}')
            return
        sheet.update(f'J{target}', [[(reason or '').strip() or 'Замена']], value_input_option='USER_ENTERED')
        logger.info(f'Причина вывоза записана в «Общее», строка {target}')
    except Exception as e:
        logger.error(f'Не записал причину вывоза: {e}')

def write_evacuation_to_sheet(uid, master_name, darks_number, address, old_data, new_data, reason=''):
    try:
        append_bike_swap(address, darks_number, old_data, new_data, master_name, uid, 'Эвакуация')
        sheet_client = get_sheet_client()
        now = get_msk_now().strftime('%Y-%m-%d %H:%M:%S')
        
        try:
            report_sheet = sheet_client.worksheet("Эвакуация Транзит")
        except:
            report_sheet = sheet_client.add_worksheet("Эвакуация Транзит", 100, 20)
            headers = ['Отметка времени', 'Мастер', 'ID заявки', 'Адрес даркстора', 'Номер даркстора',
                      'ЗАБРАЛ - Серийный номер', 'ЗАБРАЛ - Гос номер', 'ЗАБРАЛ - Номер айот',
                      'ОТДАЛ - Серийный номер', 'ОТДАЛ - Гос номер', 'ОТДАЛ - Номер айот', 'Тип']
            for i, h in enumerate(headers, start=1):
                report_sheet.update_cell(1, i, h)
        
        current_rows = report_sheet.get_all_values()
        new_row_idx = len(current_rows) + 1
        
        updates = [
            {'range': f'A{new_row_idx}', 'values': [[now]]},
            {'range': f'B{new_row_idx}', 'values': [[master_name]]},
            {'range': f'C{new_row_idx}', 'values': [[uid]]},
            {'range': f'D{new_row_idx}', 'values': [[address]]},
            {'range': f'E{new_row_idx}', 'values': [[darks_number]]},
            {'range': f'F{new_row_idx}', 'values': [[old_data.get('serial', '')]]},
            {'range': f'G{new_row_idx}', 'values': [[old_data.get('gos', '')]]},
            {'range': f'H{new_row_idx}', 'values': [[old_data.get('iot', '')]]},
            {'range': f'I{new_row_idx}', 'values': [[new_data.get('serial', '')]]},
            {'range': f'J{new_row_idx}', 'values': [[new_data.get('gos', '')]]},
            {'range': f'K{new_row_idx}', 'values': [[new_data.get('iot', '')]]},
            {'range': f'L{new_row_idx}', 'values': [['Эвакуация']]}
        ]
        report_sheet.batch_update(updates)
        export_removal_reason(old_data, reason)
        logger.info(f"✅ Записана эвакуация для заявки {uid}")
        return True, ''
    except Exception as e:
        logger.error(f"Ошибка записи эвакуации: {e}")
        return False, str(e)

@app.route('/master/evacuation/replace', methods=['POST'])
@login_required
def master_evacuation_replace():
    try:
        data = request.json
        uid = data.get('uid')
        master_name = data.get('master')
        darks_number = data.get('darks_number')
        address = data.get('address', '')
        old_data = data.get('old_data', {})
        new_data = data.get('new_data', {})
        
        if not uid or not master_name or not darks_number:
            return jsonify({'success': False, 'error': 'Недостаточно данных'}), 400
        
        if not old_data.get('serial') or not old_data.get('gos') or not old_data.get('iot'):
            return jsonify({'success': False, 'error': 'Заполните все поля СТАРОГО велосипеда'}), 400
        
        if not new_data.get('serial') or not new_data.get('gos') or not new_data.get('iot'):
            return jsonify({'success': False, 'error': 'Заполните все поля НОВОГО велосипеда'}), 400
        
        ticket = find_cached_ticket(uid) or {}
        reason = evacuation_reason(ticket)
        saved, err = write_evacuation_to_sheet(uid, master_name, darks_number, address, old_data, new_data, reason)
        if not saved:
            return jsonify({'success': False, 'error': err or 'Не удалось записать замену в таблицу'})
        
        update_ticket_in_admin_cache(uid, 'done', 'Заменен при эвакуации', '✅ Выполнено')
        update_status_in_google_sheets(uid, '✅ Выполнено', 'Заменен при эвакуации')
        notify_dark_event(darks_number, f'✅ {old_data.get("gos") or ticket_label(ticket)} заменен и готов к эксплуатации', uid, 'replaced')
        maybe_ask_rating(darks_number, master_name)
        
        add_to_queue({
            'uid': uid,
            'source': source_of(uid),
            'type': 'evacuation_and_replace',
            'data': {
                'master': master_name,
                'darks_number': darks_number,
                'old_data': old_data,
                'new_data': new_data
            }
        })
        
        cache_data = get_master_cache(master_name)
        if cache_data:
            tickets = cache_data.get('tickets', [])
            tickets = [t for t in tickets if t.get('uid') != uid]
            save_master_cache(master_name, tickets, '')
        
        return jsonify({'success': True, 'message': 'Эвакуация выполнена'})
        
    except Exception as e:
        logger.error(f"Ошибка эвакуации: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500

# ============================================================
# ФУНКЦИИ ДЛЯ ТРАНЗИТА (ОБЪЕДИНЕНИЕ ЗАЯВОК)
# ============================================================
def get_transit_tickets_by_gos(master_name, gos_number):
    try:
        tickets = get_tickets_from_sheets()
        result = []
        for t in tickets:
            if (same_master(t.get('master'), master_name) and t.get('gos') == gos_number and is_active_status(t.get('status'))):
                result.append(t)
        return result
    except Exception as e:
        logger.error(f"Ошибка получения заявок Транзита: {e}")
        return []

def close_all_transit_tickets(uid, master_name, gos_number, parts=''):
    try:
        tickets = get_transit_tickets_by_gos(master_name, gos_number)
        
        if not tickets:
            return 0
        
        closed_count = 0
        
        for t in tickets:
            ticket_uid = t.get('uid')
            if not ticket_uid:
                continue
            
            update_ticket_in_admin_cache(ticket_uid, 'done', parts or 'Заменено Транзитом', '✅ Выполнено')
            update_status_in_google_sheets(ticket_uid, '✅ Выполнено', parts or 'Заменено Транзитом')
            
            add_to_queue({
                'uid': ticket_uid,
                'source': t.get('source') or 'Заявки',
                'type': 'transit_bulk_close',
                'data': {
                    'master': master_name,
                    'gos': gos_number,
                    'parts': parts,
                    'closed_with': uid
                }
            })
            closed_count += 1
        
        refresh_master_cache(master_name)
        return closed_count
        
    except Exception as e:
        logger.error(f"Ошибка закрытия заявок Транзита: {e}")
        return 0

@app.route('/master/transit/done/<uid>', methods=['POST'])
@login_required
def transit_done(uid):
    try:
        master_name = session.get('master_name')
        if not is_transit_master(master_name):
            return jsonify({'success': False, 'error': 'Только для Транзита'}), 403
        
        parts = request.form.get('parts', '')
        if not parts.strip():
            return jsonify({'success': False, 'error': 'Укажите запчасти'}), 400
        
        tickets = get_tickets_from_sheets()
        target_ticket = None
        for t in tickets:
            if t.get('uid') == uid:
                target_ticket = t
                break
        
        if not target_ticket:
            return jsonify({'success': False, 'error': 'Заявка не найдена'}), 404
        
        gos_number = target_ticket.get('gos')
        if not gos_number:
            return jsonify({'success': False, 'error': 'У заявки нет госномера'}), 400
        
        closed_count = close_all_transit_tickets(uid, master_name, gos_number, parts)
        
        if closed_count == 0:
            return jsonify({'success': False, 'error': 'Не найдено заявок для закрытия'}), 404
        maybe_ask_rating(target_ticket.get('darks'), master_name)
        return jsonify({
            'success': True, 
            'message': f'Закрыто {closed_count} заявок на велосипед {gos_number}',
            'closed_count': closed_count
        })
        
    except Exception as e:
        logger.error(f"Ошибка закрытия заявок Транзита: {e}")
        return jsonify({'success': False, 'error': str(e)}), 500

# ============================================================
# IOT-СИСТЕМА (ЗАМЕНА IOT-МОДУЛЕЙ)
# ============================================================

def load_iot_source(force_reload=False):
    """Загружает список IOT из листа 'Все IoT' (кэш в памяти + JSON)"""
    global _iot_source_memory
    
    # Если не force_reload и есть кэш в памяти — вернуть
    if not force_reload and _iot_source_memory is not None:
        return _iot_source_memory
    
    # Если не force_reload и есть JSON — прочитать
    if not force_reload:
        cache = read_cache(IOT_SOURCE_FILE)
        if cache and 'iot_list' in cache:
            _iot_source_memory = cache['iot_list']
            logger.info(f"✅ IOT Source загружен из кэша: {len(_iot_source_memory)} записей")
            return _iot_source_memory
    
    # Иначе — тянем из Google Sheets
    try:
        sheet = get_iot_sheet_by_id(IOT_SOURCE_SHEET_ID, "Все IoT")
        rows = sheet.get_all_values()
        
        iot_data = {}
        if len(rows) > 1:
            for row in rows[1:]:
                if len(row) >= 9:
                    iot = row[0].strip()
                    if not iot:
                        continue
                    status_velo = row[7].strip() if len(row) > 7 else ''
                    status_iot = row[8].strip() if len(row) > 8 else ''
                    replaced_text = row[9].strip() if len(row) > 9 else ''
                    firmware = row[5].strip() if len(row) > 5 else ''
                    iot_data[iot] = {
                        'status_velo': status_velo,
                        'status_iot': status_iot,
                        'replaced_text': replaced_text,
                        'firmware': firmware,
                    }
        
        cache_data = {
            'updated_at': get_msk_now().strftime('%Y-%m-%d %H:%M:%S'),
            'iot_list': iot_data
        }
        write_cache(IOT_SOURCE_FILE, cache_data)
        _iot_source_memory = iot_data
        logger.info(f"✅ IOT Source загружен из Google Sheets: {len(iot_data)} записей")
        return iot_data
    except Exception as e:
        logger.error(f"Ошибка загрузки IOT Source: {e}")
        return {}


def iot_needs_work(info):
    """В аренде, столбец J пустой, версия не 3.0.5 / 3.0.6 / 160."""
    if info.get('status_velo') != 'В аренде':
        return False
    if (info.get('replaced_text') or '').strip():
        return False
    version = (info.get('firmware') or '').strip()
    if version in IOT_DONE_VERSIONS:
        return False
    return True


def get_iot_source():
    """Возвращает кэш IOT Source (из памяти)"""
    return load_iot_source(force_reload=False)


def load_iot_vehicles(force_reload=False):
    """Загружает данные велосипедов из листа 'Учет вело ВВ' (кэш в памяти + JSON)"""
    global _iot_vehicles_memory
    
    # Если не force_reload и есть кэш в памяти — вернуть
    if not force_reload and _iot_vehicles_memory is not None:
        return _iot_vehicles_memory
    
    # Если не force_reload и есть JSON — прочитать
    if not force_reload:
        cache = read_cache(IOT_VEHICLES_FILE)
        if cache and 'vehicles' in cache:
            _iot_vehicles_memory = cache['vehicles']
            logger.info(f"✅ IOT Vehicles загружен из кэша: {len(_iot_vehicles_memory)} записей")
            return _iot_vehicles_memory
    
    # Иначе — тянем из Google Sheets
    try:
        sheet = get_iot_sheet_by_id(IOT_VEHICLES_SHEET_ID, "Учет вело ВВ")
        rows = sheet.get_all_values()
        
        vehicles = {}
        if len(rows) > 1:
            for row in rows[1:]:
                if len(row) >= 5:
                    frame_number = row[0].strip()
                    gos = row[1].strip()
                    iot = row[2].strip()
                    address = row[3].strip()
                    darks = row[4].strip()
                    
                    if iot:
                        vehicles[iot] = {
                            'frame_number': frame_number,
                            'gos': gos,
                            'iot': iot,
                            'address': address,
                            'darks': darks
                        }
        
        cache_data = {
            'updated_at': get_msk_now().strftime('%Y-%m-%d %H:%M:%S'),
            'vehicles': vehicles
        }
        write_cache(IOT_VEHICLES_FILE, cache_data)
        _iot_vehicles_memory = vehicles
        logger.info(f"✅ IOT Vehicles загружен из Google Sheets: {len(vehicles)} записей")
        return vehicles
    except Exception as e:
        logger.error(f"Ошибка загрузки IOT Vehicles: {e}")
        return {}


def get_iot_vehicles():
    """Возвращает кэш IOT Vehicles (из памяти)"""
    return load_iot_vehicles(force_reload=False)


def reset_iot_memory():
    """Сбрасывает кэш в памяти (при обновлении данных)"""
    global _iot_source_memory, _iot_vehicles_memory, _iot_curators_memory
    _iot_source_memory = None
    _iot_vehicles_memory = None
    _iot_curators_memory = None


def parse_curator_contacts(cell):
    found = []
    for part in re.split(r'[,;\s]+', cell or ''):
        token = part.strip()
        if token.startswith('@') and len(token) > 1 and token not in found:
            found.append(token)
    return found


def load_iot_curator_contacts(force_reload=False):
    """Контакты кураторов из листа Дарксторы, столбец G. Кэш в памяти и в файле."""
    global _iot_curators_memory
    if not force_reload and _iot_curators_memory is not None:
        return _iot_curators_memory
    if not force_reload:
        cache = read_cache(IOT_CURATORS_FILE)
        if cache and isinstance(cache.get('curators'), dict):
            _iot_curators_memory = cache['curators']
            logger.info(f"✅ Контакты кураторов из кэша: {len(_iot_curators_memory)} точек")
            return _iot_curators_memory
    try:
        sheet = get_iot_sheet_by_id(IOT_CURATORS_SHEET_ID, "Дарксторы")
        rows = sheet.get_all_values()
        curators = {}
        for row in rows[1:]:
            number = row[0].strip() if row else ''
            if not number:
                continue
            curators[number] = parse_curator_contacts(row[6] if len(row) > 6 else '')
        write_cache(IOT_CURATORS_FILE, {
            'updated_at': get_msk_now().strftime('%Y-%m-%d %H:%M:%S'),
            'curators': curators,
        })
        _iot_curators_memory = curators
        logger.info(f"✅ Контакты кураторов загружены: {len(curators)} точек")
        return curators
    except Exception as e:
        logger.error(f"Ошибка загрузки контактов кураторов: {e}")
        return _iot_curators_memory or {}


def get_iot_curator_contacts():
    return load_iot_curator_contacts(force_reload=False)


def get_iot_bag():
    cache = read_cache(IOT_BAG_FILE)
    if not cache:
        return {'bag': []}
    return cache


def save_iot_bag(bag_data):
    return write_cache(IOT_BAG_FILE, bag_data)


def get_iot_history():
    cache = read_cache(IOT_HISTORY_FILE)
    if not cache:
        return {'history': []}
    return cache


def save_iot_history(history_data):
    return write_cache(IOT_HISTORY_FILE, history_data)


def update_iot_source_column_j(old_iot, new_iot):
    """Записывает в столбец J листа 'Все IoT' информацию о замене (в строку СТАРОГО IOT)"""
    try:
        sheet = get_iot_sheet_by_id(IOT_SOURCE_SHEET_ID, "Все IoT")
        rows = sheet.get_all_values()
        
        for idx, row in enumerate(rows, start=1):
            if len(row) > 0 and row[0].strip() == old_iot:
                now = get_msk_now()
                date_str = now.strftime('%d.%m')
                text = f"{date_str} поменяли на прошитый (новый IOT: {new_iot})"
                sheet.update_cell(idx, 10, text)
                logger.info(f"✅ Запись в J для {old_iot}: {text}")
                return True
        
        logger.warning(f"⚠️ IOT {old_iot} не найден в листе 'Все IoT'")
        return False
    except Exception as e:
        logger.error(f"Ошибка записи в столбец J: {e}")
        return False


def write_iot_report(frame_number, new_iot, old_iot=''):
    """Записывает отчёт в лист 'КОРРЕКТИРОВКИ ВЕЛО'"""
    try:
        sheet = get_iot_sheet_by_id(IOT_REPORT_SHEET_ID, "КОРРЕКТИРОВКИ ВЕЛО")
        
        now = get_msk_now()
        date_str = now.strftime('%d.%m')
        
        all_values = sheet.get_all_values()
        new_row = len(all_values) + 1
        
        sheet.update(f'A{new_row}:F{new_row}', [[
            date_str,
            'Изменить IOT',
            frame_number,
            new_iot,
            '',
            'IOT с прошивкой'
        ]])
        
        logger.info(f"✅ Отчёт записан: {frame_number} → {new_iot}")
        return True
    except Exception as e:
        logger.error(f"Ошибка записи отчёта IOT: {e}")
        return False


@app.route('/iot')
@login_required
def iot_main():
    """Главная страница IOT-мастера (МГНОВЕННО из кэша)"""
    if session.get('role') != 'iot':
        return redirect(url_for('login_page'))
    
    bag_data = get_iot_bag()
    bag_items = [item.get('iot') for item in bag_data.get('bag', [])]
    
    history_data = get_iot_history()
    today = get_msk_now().strftime('%d.%m')
    replaced_today = 0
    for h in history_data.get('history', []):
        if h.get('date') == today:
            replaced_today += 1
    
    if not bag_items:
        return render_template('iot_empty.html', now=get_msk_now().strftime('%H:%M:%S'))
    
    # Читаем из кэша (мгновенно)
    vehicles_data = get_iot_vehicles()
    source_data = get_iot_source()
    
    darks_groups = {}
    for iot, info in source_data.items():
        if iot_needs_work(info):
            if iot in vehicles_data:
                v = vehicles_data[iot]
                darks = v.get('darks', 'без номера')
                if darks not in darks_groups:
                    darks_groups[darks] = {
                        'darks_number': darks,
                        'address': v.get('address', 'Адрес не указан'),
                        'vehicles_count': 0,
                        'preview': []
                    }
                darks_groups[darks]['vehicles_count'] += 1
                if len(darks_groups[darks]['preview']) < 3:
                    darks_groups[darks]['preview'].append({
                        'gos': v.get('gos', ''),
                        'old_iot': iot
                    })
    
    darks_list = list(darks_groups.values())
    darks_list.sort(key=lambda x: int(x['darks_number']) if x['darks_number'].isdigit() else 999999)
    curators = get_iot_curator_contacts()
    for group in darks_list:
        group['curators'] = curators.get(str(group['darks_number']), [])
    
    return render_template('iot_main.html',
                          bag_items=bag_items,
                          bag_count=len(bag_items),
                          replaced_today=replaced_today,
                          darks_groups=darks_list,
                          now=get_msk_now().strftime('%H:%M:%S'))


@app.route('/iot/darks/<darks_number>')
@login_required
def iot_darks(darks_number):
    """Страница даркстора для IOT-мастера (МГНОВЕННО из кэша)"""
    if session.get('role') != 'iot':
        return redirect(url_for('login_page'))
    
    bag_data = get_iot_bag()
    bag_items = [item.get('iot') for item in bag_data.get('bag', [])]
    
    # Читаем из кэша (мгновенно)
    vehicles_data = get_iot_vehicles()
    source_data = get_iot_source()
    
    vehicles = []
    address = ''
    for iot, info in source_data.items():
        if iot_needs_work(info):
            if iot in vehicles_data:
                v = vehicles_data[iot]
                if v.get('darks') == darks_number:
                    if not address:
                        address = v.get('address', '')
                    vehicles.append({
                        'old_iot': iot,
                        'frame_number': v.get('frame_number', ''),
                        'gos': v.get('gos', ''),
                        'address': v.get('address', ''),
                        'darks': v.get('darks', ''),
                        'firmware': info.get('firmware') or '',
                    })
    
    return render_template('iot_darks.html',
                          darks_number=darks_number,
                          address=address or 'Адрес не указан',
                          vehicles=vehicles,
                          bag_items=bag_items,
                          bag_count=len(bag_items),
                          curators=get_iot_curator_contacts().get(str(darks_number), []),
                          now=get_msk_now().strftime('%H:%M:%S'))


@app.route('/iot/history')
@login_required
def iot_history_page():
    """Страница истории замен"""
    if session.get('role') != 'iot':
        return redirect(url_for('login_page'))
    
    history_data = get_iot_history()
    history = history_data.get('history', [])
    
    history_by_darks = {}
    for h in reversed(history):
        darks = h.get('darks', 'без номера')
        if darks not in history_by_darks:
            history_by_darks[darks] = []
        history_by_darks[darks].append(h)
    
    return render_template('iot_history.html',
                          history_by_darks=history_by_darks,
                          now=get_msk_now().strftime('%H:%M:%S'))


@app.route('/api/iot/load_bag', methods=['POST'])
@login_required
def api_iot_load_bag():
    """Загружает IOT в багажник БЕЗ проверок"""
    if session.get('role') != 'iot':
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    
    try:
        data = request.json
        iots = data.get('iots', [])
        
        if not iots:
            return jsonify({'success': False, 'error': 'Пустой список'})
        
        bag_data = get_iot_bag()
        current_bag = [item.get('iot') for item in bag_data.get('bag', [])]
        
        added = []
        duplicates = []
        seen_in_input = set()
        
        for iot in iots:
            iot = iot.strip()
            if not iot:
                continue
            
            if iot in seen_in_input:
                continue
            seen_in_input.add(iot)
            
            if iot in current_bag:
                duplicates.append(iot)
                continue
            
            added.append(iot)
        
        for iot in added:
            bag_data['bag'].append({
                'iot': iot,
                'added': get_msk_now().strftime('%Y-%m-%d %H:%M:%S')
            })
        save_iot_bag(bag_data)
        
        return jsonify({
            'success': True,
            'added': added,
            'errors': [],
            'duplicates': duplicates,
            'total_input': len(seen_in_input)
        })
        
    except Exception as e:
        logger.error(f"Ошибка загрузки багажника: {e}")
        return jsonify({'success': False, 'error': str(e)})


@app.route('/api/iot/remove_from_bag', methods=['POST'])
@login_required
def api_iot_remove_from_bag():
    """Удаляет один IOT из багажника"""
    if session.get('role') != 'iot':
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    
    try:
        data = request.json
        iot = data.get('iot', '').strip()
        
        if not iot:
            return jsonify({'success': False, 'error': 'Не указан IOT'})
        
        bag_data = get_iot_bag()
        current_bag = bag_data.get('bag', [])
        
        if not any(item.get('iot') == iot for item in current_bag):
            return jsonify({'success': False, 'error': f'IOT {iot} не найден в багажнике'})
        
        bag_data['bag'] = [item for item in current_bag if item.get('iot') != iot]
        save_iot_bag(bag_data)
        
        logger.info(f"✅ IOT {iot} удалён из багажника")
        
        return jsonify({
            'success': True,
            'message': f'IOT {iot} удалён из багажника'
        })
        
    except Exception as e:
        logger.error(f"Ошибка удаления IOT из багажника: {e}")
        return jsonify({'success': False, 'error': str(e)})


@app.route('/api/iot/clear_bag', methods=['POST'])
@login_required
def api_iot_clear_bag():
    """Полностью очищает багажник"""
    if session.get('role') != 'iot':
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    
    try:
        bag_data = get_iot_bag()
        count = len(bag_data.get('bag', []))
        
        bag_data['bag'] = []
        save_iot_bag(bag_data)
        
        logger.info(f"✅ Багажник очищен ({count} IOT удалено)")
        
        return jsonify({
            'success': True,
            'message': f'Багажник очищен ({count} IOT удалено)',
            'count': count
        })
        
    except Exception as e:
        logger.error(f"Ошибка очистки багажника: {e}")
        return jsonify({'success': False, 'error': str(e)})


@app.route('/api/iot/sync_source', methods=['POST'])
@login_required
def api_iot_sync_source():
    """Обновляет ОБА кэша: iot_source и iot_vehicles"""
    if session.get('role') != 'iot':
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    
    try:
        # Сбрасываем кэш в памяти
        reset_iot_memory()
        
        # Принудительно тянем из Google Sheets
        source_data = load_iot_source(force_reload=True)
        vehicles_data = load_iot_vehicles(force_reload=True)
        try:
            curators = load_iot_curator_contacts(force_reload=True)
        except Exception as cur_error:
            logger.error(f"Контакты кураторов не обновились: {cur_error}")
            curators = get_iot_curator_contacts()
        
        logger.info(f"✅ Обновлено: IOT Source = {len(source_data)}, IOT Vehicles = {len(vehicles_data)}, кураторы = {len(curators)}")
        
        return jsonify({
            'success': True,
            'count': len(source_data),
            'source_count': len(source_data),
            'vehicles_count': len(vehicles_data),
            'curators_count': len(curators),
        })
    except Exception as e:
        logger.error(f"Ошибка синхронизации IOT: {e}")
        return jsonify({'success': False, 'error': str(e)})


@app.route('/api/iot/bag')
@login_required
def api_iot_bag():
    if session.get('role') != 'iot':
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    
    bag_data = get_iot_bag()
    return jsonify({'success': True, 'bag': bag_data})


@app.route('/api/iot/replace', methods=['POST'])
@login_required
def api_iot_replace():
    """Заменяет IOT: пишет в 'Все IoT' столбец J и в 'КОРРЕКТИРОВКИ ВЕЛО'"""
    if session.get('role') != 'iot':
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    
    try:
        data = request.json
        old_iot = data.get('old_iot', '').strip()
        new_iot = data.get('new_iot', '').strip()
        frame_number = data.get('frame_number', '').strip()
        gos = data.get('gos', '').strip()
        darks = data.get('darks', '').strip()
        address = data.get('address', '').strip()
        
        if not old_iot or not new_iot or not frame_number:
            return jsonify({'success': False, 'error': 'Недостаточно данных'})
        
        bag_data = get_iot_bag()
        current_bag = [item.get('iot') for item in bag_data.get('bag', [])]
        
        if new_iot not in current_bag:
            return jsonify({'success': False, 'error': f'IOT {new_iot} не в багажнике'})
        
        vehicles_data = load_iot_vehicles()
        if new_iot in vehicles_data:
            bag_data['bag'] = [item for item in bag_data.get('bag', []) if item.get('iot') != new_iot]
            save_iot_bag(bag_data)
            return jsonify({
                'success': False,
                'error': f'IOT {new_iot} уже установлен на велосипеде {vehicles_data[new_iot].get("frame_number", "")}! Верните его в цех.'
            })
        
        # 1. Запись в "КОРРЕКТИРОВКИ ВЕЛО"
        write_iot_report(frame_number, new_iot, old_iot)
        
        # 2. Запись в "Все IoT" столбец J (в строку СТАРОГО IOT)
        update_iot_source_column_j(old_iot, new_iot)
        
        # 3. Обновляем кэш в памяти: помечаем старый IOT как заменённый
        global _iot_source_memory
        if _iot_source_memory and old_iot in _iot_source_memory:
            now_str = get_msk_now().strftime('%d.%m')
            _iot_source_memory[old_iot]['replaced_text'] = f"{now_str} поменяли на прошитый (новый IOT: {new_iot})"
            # Сохраняем в JSON
            cache_data = {
                'updated_at': get_msk_now().strftime('%Y-%m-%d %H:%M:%S'),
                'iot_list': _iot_source_memory
            }
            write_cache(IOT_SOURCE_FILE, cache_data)
            logger.info(f"✅ Кэш в памяти обновлён: {old_iot} помечен как заменённый")
        
        # 4. Убираем новый IOT из багажника
        bag_data['bag'] = [item for item in bag_data.get('bag', []) if item.get('iot') != new_iot]
        save_iot_bag(bag_data)
        
        # 5. Сохраняем в локальную историю
        now = get_msk_now()
        history_data = get_iot_history()
        history_data['history'].append({
            'type': 'replace',
            'date': now.strftime('%d.%m'),
            'time': now.strftime('%H:%M'),
            'old_iot': old_iot,
            'new_iot': new_iot,
            'frame_number': frame_number,
            'gos': gos,
            'darks': darks,
            'address': address,
            'timestamp': now.strftime('%Y-%m-%d %H:%M:%S')
        })
        save_iot_history(history_data)
        
        logger.info(f"✅ IOT заменён: {old_iot} → {new_iot} (рама {frame_number})")
        
        return jsonify({
            'success': True,
            'message': f'IOT {old_iot} заменён на {new_iot}'
        })
        
    except Exception as e:
        logger.error(f"Ошибка замены IOT: {e}")
        return jsonify({'success': False, 'error': str(e)})


def mark_iot_flashed(old_iot):
    """F → 3.0.5, I → ОК, J → дата прошивки. Багажник и корректировки не трогает."""
    sheet = get_iot_sheet_by_id(IOT_SOURCE_SHEET_ID, "Все IoT")
    rows = sheet.get_all_values()
    now = get_msk_now()
    text = f"{now.strftime('%d.%m')} прошит (IOT остался: {old_iot})"
    for idx, row in enumerate(rows, start=1):
        if not row or row[0].strip() != old_iot:
            continue
        sheet.batch_update([
            {'range': f'F{idx}', 'values': [[IOT_FLASHED_VERSION]]},
            {'range': f'I{idx}', 'values': [['ОК']]},
            {'range': f'J{idx}', 'values': [[text]]},
        ])
        logger.info(f"✅ Прошивка записана для {old_iot}: {IOT_FLASHED_VERSION}, {text}")
        return text
    logger.warning(f"⚠️ IOT {old_iot} не найден в листе 'Все IoT'")
    return ''


@app.route('/api/iot/flash', methods=['POST'])
@login_required
def api_iot_flash():
    """Прошивка IOT на велосипеде: модуль не снимается."""
    if session.get('role') != 'iot':
        return jsonify({'success': False, 'error': 'Доступ запрещён'})
    try:
        data = request.json or {}
        old_iot = (data.get('old_iot') or '').strip()
        frame_number = (data.get('frame_number') or '').strip()
        gos = (data.get('gos') or '').strip()
        darks = (data.get('darks') or '').strip()
        address = (data.get('address') or '').strip()
        if not old_iot:
            return jsonify({'success': False, 'error': 'Не указан IOT'})
        text = mark_iot_flashed(old_iot)
        if not text:
            return jsonify({'success': False, 'error': f'IOT {old_iot} не найден в таблице'})
        global _iot_source_memory
        if _iot_source_memory is None:
            load_iot_source(force_reload=False)
        if _iot_source_memory is not None:
            current = _iot_source_memory.get(old_iot) or {}
            current['firmware'] = IOT_FLASHED_VERSION
            current['status_iot'] = 'ОК'
            current['replaced_text'] = text
            _iot_source_memory[old_iot] = current
            write_cache(IOT_SOURCE_FILE, {
                'updated_at': get_msk_now().strftime('%Y-%m-%d %H:%M:%S'),
                'iot_list': _iot_source_memory,
            })
        now = get_msk_now()
        history_data = get_iot_history()
        history_data['history'].append({
            'type': 'flash',
            'date': now.strftime('%d.%m'),
            'time': now.strftime('%H:%M'),
            'old_iot': old_iot,
            'new_iot': None,
            'frame_number': frame_number,
            'gos': gos,
            'darks': darks,
            'address': address,
            'timestamp': now.strftime('%Y-%m-%d %H:%M:%S'),
        })
        save_iot_history(history_data)
        logger.info(f"✅ IOT прошит: {old_iot} (рама {frame_number})")
        return jsonify({'success': True, 'message': 'IOT успешно прошит'})
    except Exception as e:
        logger.error(f"Ошибка прошивки IOT: {e}")
        return jsonify({'success': False, 'error': str(e)})


# ============================================================
# ЗАПУСК
# ============================================================
ensure_queue_thread()

if __name__ == "__main__":
    logger.info("🚀 Запуск приложения...")
    ensure_queue_thread()
    ensure_xdevice_loop()
    logger.info("🚀 Фоновый процесс обработки очереди запущен")
    
    try:
        logger.info("📂 Загрузка данных из Google Sheets...")
        tickets = get_tickets_from_sheets(force=True)
        publish_caches(tickets)
        logger.info(f"✅ Кэши созданы: {len(tickets)} заявок")
        
        # Создаём кэш IOT при старте
        logger.info("📂 Создание IOT кэшей...")
        try:
            load_iot_source(force_reload=True)
            load_iot_vehicles(force_reload=True)
            load_iot_curator_contacts(force_reload=True)
            logger.info("   ✅ IOT кэши созданы")
        except Exception as e:
            logger.error(f"   ❌ Ошибка создания IOT кэшей: {e}")
        
        queue_path = get_queue_path()
        if not os.path.exists(queue_path):
            write_queue({'tasks': [], 'last_sync': get_msk_now().strftime('%Y-%m-%d %H:%M:%S')})
            logger.info("✅ Создана пустая очередь")
        
        logger.info("✅ ВСЕ КЭШИ УСПЕШНО СОЗДАНЫ!")
        
    except Exception as e:
        logger.error(f"❌ КРИТИЧЕСКАЯ ОШИБКА ПРИ СОЗДАНИИ КЭШЕЙ: {e}")
        import traceback
        traceback.print_exc()
    
    port = int(os.environ.get("PORT", 5000))
    
    ssl_context = None
    try:
        if os.path.exists('/etc/ssl/certs/amvera.crt') and os.path.exists('/etc/ssl/private/amvera.key'):
            ssl_context = ('/etc/ssl/certs/amvera.crt', '/etc/ssl/private/amvera.key')
            logger.info("✅ SSL сертификаты найдены, запуск с HTTPS")
    except:
        pass
    
    if ssl_context:
        app.run(host="0.0.0.0", port=port, ssl_context=ssl_context)
    else:
        app.run(host="0.0.0.0", port=port)
